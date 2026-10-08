"""Prueba de humo de scripts/aplicar_tupa.py con un libro y un domain.yml FALSOS en una carpeta temporal (rótulo: SIMULADA). Nada se escribe en el repositorio.

Comprueba: el dry-run lista solo las filas Corregir + texto + confirmadas + sin alerta y NO escribe domain.yml; --aplicar se niega sin partición y evaluación; con ellas reemplaza solo esas
respuestas (el resto del archivo queda igual, sigue siendo YAML válido, copia previa) y una intención inexistente bloquea.
"""
import hashlib
import shutil
import subprocess
import sys
import tempfile
import warnings
from pathlib import Path

warnings.filterwarnings("ignore")
import yaml
from openpyxl import Workbook

ROOT = Path(__file__).resolve().parent.parent
RES = []


def check(n, c, d=""):
    RES.append(bool(c))
    print(f"  [{'PASS' if c else 'FAIL'}] {n}" + (f"  ({str(d)[:250]})" if d and not c else ""))


def run(*args):
    p = subprocess.run([sys.executable, str(ROOT / "scripts" / "aplicar_tupa.py"), *map(str, args)], capture_output=True, text=True, encoding="utf-8", errors="replace", cwd=ROOT,
                       env={"PYTHONIOENCODING": "utf-8", "PATH": __import__("os").environ["PATH"], "SYSTEMROOT": __import__("os").environ.get("SYSTEMROOT", "")})
    return p.returncode, (p.stdout or "") + (p.stderr or "")


def libro(ruta, filas):
    wb = Workbook(); ws = wb.active; ws.title = "Verificacion"
    ws["A1"] = "Verificación FALSA"
    for j, h in enumerate(["#", "Intención", "Categoría", "Prioridad", "x", "x", "x", "x", "x", "x", "x", "x", "x", "Resultado", "Texto corregido para domain.yml", "Fecha", "Obs", "Alerta", "Confirmado por el tesista"], 1):
        ws.cell(3, j, h)
    for i, (it, res, texto, alerta, conf) in enumerate(filas, 4):
        for j, v in ((1, i - 3), (2, it), (4, "Alta"), (14, res), (15, texto), (18, alerta), (19, conf)):
            ws.cell(i, j, v)
    rs = wb.create_sheet("Resumen"); rs["A5"] = "Confirmadas por el tesista"; rs["B5"] = 2
    wb.save(ruta)


def main():
    W = Path(tempfile.mkdtemp(prefix="aplicar_tupa_PRUEBA_SIMULADA_"))
    print(f"SIMULADA — carpeta temporal (datos FALSOS): {W}")
    dom = W / "domain.yml"
    dom.write_text("version: '3.1'\nintents:\n- a\n- b\n- c\nresponses:\n  utter_a:\n  - text: 'Respuesta vieja de a. [Verificar plazo exacto]'\n  utter_b:\n  - text: Respuesta vieja de b sin comillas.\n"
                   "  utter_c:\n  - text: Respuesta de c que no cambia\nactions: []\n", encoding="utf-8")
    lb = W / "Verificacion_TUPA_v1.xlsx"
    libro(lb, [("a", "Corregir", "Texto nuevo de a: cuesta S/ 50, requisitos A y B.", None, "Sí"), ("b", "Corregir", "Texto nuevo de b: 'con comillas' y dos puntos: así.", None, "Sí"),
               ("c", "Corregir", "Texto nuevo de c sin confirmar", None, None), ("x", "Pendiente", None, None, None), ("y", "No figura en la fuente", None, None, "Sí")])
    h0 = hashlib.sha256(dom.read_bytes()).hexdigest()
    inf = W / "informe.md"
    c, t = run("--libro", lb, "--domain", dom, "--informe", inf)
    check("dry-run: entran 2 de 5 (Corregir + texto + confirmadas); las otras 3 quedan fuera", c == 0 and "ENTRARÍAN: 2" in t and "Corregir sin confirmar 1" in t, t[-300:])
    check("dry-run: no escribe domain.yml y escribe el informe", hashlib.sha256(dom.read_bytes()).hexdigest() == h0 and inf.exists() and "utter_a" in inf.read_text(encoding="utf-8"))
    c, t = run("--libro", lb, "--domain", dom, "--informe", inf, "--aplicar", "--raiz", W)
    check("--aplicar sin partición ni evaluación se niega y no toca domain.yml", c != 0 and "ME NIEGO" in t and "partición" in t and hashlib.sha256(dom.read_bytes()).hexdigest() == h0, t[-300:])
    for r in ("corpus/v3_real/dataset_split_v3.csv", "logs/v3_real/eval_real_resumen.json", "logs/v3_real/test_registro.json"):
        (W / r).parent.mkdir(parents=True, exist_ok=True); (W / r).write_text("x", encoding="utf-8")
    c, t = run("--libro", lb, "--domain", dom, "--informe", inf, "--aplicar", "--raiz", W)
    d = yaml.safe_load(dom.read_text(encoding="utf-8"))
    check("--aplicar con partición y evaluación: reemplaza utter_a y utter_b y deja utter_c", c == 0 and d["responses"]["utter_a"][0]["text"].startswith("Texto nuevo de a") and d["responses"]["utter_b"][0]["text"].startswith("Texto nuevo de b")
          and d["responses"]["utter_c"][0]["text"] == "Respuesta de c que no cambia", t[-300:])
    check("conserva las comillas y los dos puntos del texto, el resto del archivo igual y la copia previa", "'con comillas' y dos puntos: así." in d["responses"]["utter_b"][0]["text"] and d["intents"] == ["a", "b", "c"] and d["actions"] == []
          and (W / "domain.ANTES_DE_TUPA.yml").exists() and hashlib.sha256((W / "domain.ANTES_DE_TUPA.yml").read_bytes()).hexdigest() == h0)
    dom2 = W / "d2.yml"; shutil.copy(W / "domain.ANTES_DE_TUPA.yml", dom2)
    lb2 = W / "Verificacion_TUPA_v2.xlsx"; libro(lb2, [("zzz", "Corregir", "Texto de una intención que no existe", None, "Sí")])
    c, t = run("--libro", lb2, "--domain", dom2, "--informe", inf, "--aplicar", "--raiz", W)
    check("una intención que no existe en domain.yml bloquea la aplicación", c != 0 and "ME NIEGO" in t and hashlib.sha256(dom2.read_bytes()).hexdigest() == h0, t[-300:])
    ok = sum(RES)
    print(f"\nRESULTADO (SIMULADA): {ok}/{len(RES)} comprobaciones PASS" + ("" if ok == len(RES) else "  <- HAY FALLAS"))
    shutil.rmtree(W, ignore_errors=True)
    sys.exit(0 if ok == len(RES) else 1)


if __name__ == "__main__":
    main()
