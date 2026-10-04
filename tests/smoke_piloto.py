"""Prueba de humo de scripts/congelar_modelo.py y scripts/analizar_piloto.py — rótulo: SIMULADA (datos falsos, predictor falso).

Crea en una carpeta temporal un modelo falso, un registro .xlsx FALSO (armado desde cero con openpyxl, nunca a partir de la plantilla real) y un
predictor falso inyectado (NO se carga ningún modelo Rasa). Nada se guarda en el repositorio: la salida de consola va a evidencias/piloto/ y al final se
verifica que incident_log.csv, logs/ y docs/piloto/ no cambiaron.

Comprueba: el caso normal; diferencias no normales (-> Wilcoxon); el alfa frente a numpy (y el aviso si difiere de Resumen!B24); hash del modelo
distinto (aborta); registro con «SIMULADO» (rechazado); segunda prueba final (se niega); fila EJ01 (excluida); n < 60 («exploratorio»); fórmulas sin
valor guardado; y el congelamiento (no sobrescribe sin --forzar --motivo; con ellos registra la incidencia, aquí en un archivo temporal).

Los encabezados del registro son SUPUESTOS (el registro real no estaba en el repositorio): esta prueba solo valida que el script los lee según ese
supuesto, no que coincidan con el registro real.

Uso:
    python tests/smoke_piloto.py [--conservar]
"""
import argparse
import contextlib
import hashlib
import io
import json
import shutil
import sys
import tempfile
import warnings
from pathlib import Path

warnings.filterwarnings("ignore")
import numpy as np  # noqa: E402
import pandas as pd  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from scipy import stats  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))
import analizar_piloto as ap  # noqa: E402
import congelar_modelo as cm  # noqa: E402

CATALOGO = ROOT / "docs" / "lote_real_1" / "situaciones_lote1_v1.csv"
ENC = (["Código de sesión", "Elegible y completa", "Minutos pre", "Minutos post", "P10"] + [f"Ítem {i}" for i in range(1, 10)] +
       [h for k in (1, 2, 3) for h in (f"Tarjeta {k}", f"Consulta {k}", f"¿Correcta? {k}", f"Obtuvo {k}")])
RES = []


def check(nombre, cond, detalle=""):
    RES.append(bool(cond))
    print(f"  [{'PASS' if cond else 'FAIL'}] {nombre}" + (f"  ({detalle})" if detalle and not cond else ""))


def alfa_numpy(m):
    c = np.cov(np.asarray(m, float), rowvar=False)  # independiente de la fórmula del script (matriz de covarianza)
    k = c.shape[0]
    return k / (k - 1) * (1 - np.trace(c) / c.sum())


def datos_falsos(n, normal=True, semilla=7):
    rng = np.random.default_rng(semilla)
    cat = pd.read_csv(CATALOGO, dtype=str)
    pre = np.round(rng.normal(14, 3, n).clip(4), 1)
    if normal:
        dif = rng.normal(8, 2, n)
    else:
        dif = rng.exponential(1.0, n) ** 3 + 0.3  # muy asimétrica
    post = np.round(np.maximum(pre - dif, 0.5), 1)
    p10 = rng.integers(1, 5, n)
    base = rng.integers(2, 5, n)
    items = np.clip(base[:, None] + rng.integers(-1, 2, (n, 8)), 1, 5)
    i9 = np.clip(p10 + 1, 1, 5)
    filas, pred = [], {}
    for s in range(n):
        tar = rng.choice(len(cat), 3, replace=False)
        fila = [f"SA{s + 1:02d}", "Sí", pre[s], post[s], int(p10[s])] + [int(x) for x in items[s]] + [int(i9[s])]
        for k, t in enumerate(tar, 1):
            q = f"consulta falsa {s + 1}-{k}"
            esperada = cat.loc[t, "intent_esperada"]
            ok = rng.random() < 0.75
            pred[q] = (esperada if ok else "intencion_equivocada", float(rng.uniform(0.6, 0.95) if ok else rng.uniform(0.2, 0.7)), 0.1)
            fila += [cat.loc[t, "scenario_id"], q, "Sí" if ok else "No", "Sí" if rng.random() < 0.8 else "No"]
        filas.append(fila)
    return filas, pred, np.column_stack([items, i9]), (pre, post, p10)


def escribir_registro(ruta, filas, b24, hojas_extra=(), formulas_sin_valor=False):
    wb = Workbook()
    ws = wb.active
    ws.title = "Sesiones"
    ws["A1"] = "REGISTRO FALSO DE PRUEBA"
    for j, h in enumerate(ENC, 1):
        ws.cell(3, j, h)
    ejemplo = list(filas[0]); ejemplo[0] = "EJ01"
    todas = [ejemplo] + [list(f) for f in filas] + [[None] * len(ENC)]
    todas.append(["SA99", "No"] + list(filas[0][2:]))
    for i, f in enumerate(todas, 4):
        for j, v in enumerate(f, 1):
            if formulas_sin_valor and j == 2:
                v = None
            ws.cell(i, j, v)
    rs = wb.create_sheet("Resumen")
    rs["B24"] = b24
    for h in hojas_extra:
        wb.create_sheet(h)
    wb.save(ruta)


def correr(argv, predictor=None):
    """Ejecuta analizar_piloto.main en el mismo proceso. Devuelve (código, texto, resultado)."""
    buf = io.StringIO()
    try:
        with contextlib.redirect_stdout(buf):
            R = ap.main(argv, predictor)
        return 0, buf.getvalue(), R
    except SystemExit as e:
        return (e.code if isinstance(e.code, int) else 1), buf.getvalue() + str(e.code), None


def congelar(argv):
    import os
    buf = io.StringIO()
    viejo = sys.argv
    sys.argv = ["congelar_modelo.py"] + argv
    try:
        with contextlib.redirect_stdout(buf):
            cm.main()
        return 0, buf.getvalue()
    except SystemExit as e:
        return (e.code if isinstance(e.code, int) else 1), buf.getvalue() + str(e.code)
    finally:
        sys.argv = viejo


def snapshot():
    out = {}
    for base in ("logs", "docs/piloto"):
        for p in sorted((ROOT / base).rglob("*")):
            if p.is_file() and "__pycache__" not in p.parts:
                out[p.relative_to(ROOT).as_posix()] = hashlib.sha256(p.read_bytes()).hexdigest()
    out["incident_log.csv"] = hashlib.sha256((ROOT / "incident_log.csv").read_bytes()).hexdigest()
    return out


def main():
    a = argparse.ArgumentParser()
    a.add_argument("--conservar", action="store_true")
    a = a.parse_args()
    W = Path(tempfile.mkdtemp(prefix="piloto_PRUEBA_SIMULADA_"))
    print(f"SIMULADA — carpeta temporal (datos FALSOS, no se commitean): {W}")
    antes = snapshot()

    modelo = W / "modelo_falso.tar.gz"
    modelo.write_bytes(b"modelo falso de prueba")
    cfg, dom = W / "config_falsa.yml", W / "domain_falso.yml"
    cfg.write_text("language: es\n", encoding="utf-8"); dom.write_text("version: '3.1'\n", encoding="utf-8")
    inc = W / "incidentes_falsos.csv"
    cm.INCIDENTES = inc  # la incidencia de --forzar va a un archivo temporal, no a incident_log.csv
    fz = W / "modelo_congelado.json"
    base = ["--modelo", str(modelo), "--config", str(cfg), "--domain", str(dom), "--corpus", str(W / "no_existe.csv"),
            "--umbral", str(W / "no_existe.json"), "--salida", str(fz)]

    print("\nCongelar el modelo")
    c, t = congelar(base)
    j = json.loads(fz.read_text(encoding="utf-8")) if fz.exists() else {}
    check("congela el modelo falso y guarda sha256, versiones, fecha UTC y commit",
          c == 0 and j.get("sha256", {}).get("modelo") == hashlib.sha256(modelo.read_bytes()).hexdigest() and "rasa" in j.get("versiones", {})
          and j.get("zona_horaria") == "UTC" and "python" in j and "commit" in j, t)
    c, t = congelar(base)
    check("no sobrescribe un congelamiento existente", c != 0 and "--forzar" in t and fz.exists(), t)
    c, t = congelar(base + ["--forzar"])
    check("--forzar sin --motivo también se niega", c != 0 and not inc.exists(), t)
    c, t = congelar(base + ["--forzar", "--motivo", "prueba falsa del mecanismo"])
    check("--forzar --motivo rehace el congelamiento, conserva el anterior y registra la incidencia (archivo temporal)",
          c == 0 and inc.exists() and "prueba falsa del mecanismo" in inc.read_text(encoding="utf-8") and len(list(W.glob("modelo_congelado_*.json"))) == 1, t)
    c, t = congelar(["--verificar", "--salida", str(fz)])
    check("--verificar acepta el modelo intacto", c == 0, t)

    # ------------------------------------------------------------------ caso normal
    print("\nCaso normal (n = 60)")
    filas, pred, items9, (pre, post, p10) = datos_falsos(60)
    alfa_ref = alfa_numpy(items9[:, :8])
    reg = W / "registro_falso.xlsx"
    escribir_registro(reg, filas, alfa_ref)
    predictor = lambda textos: [pred[x] for x in textos]  # noqa: E731
    out = W / "salida_normal"
    c, t, R = correr(["--registro", str(reg), "--modelo-congelado", str(fz), "--salida", str(out)], predictor)
    check("el caso normal termina con código 0 y n = 60", c == 0 and R and R["n"] == 60, t[-400:])
    if R:
        d = pre - post
        check("rotula REAL (n efectivo 60) y no marca «exploratorio» porque n = 60", R["rotulo"] == "REAL" and not R["exploratorio"] and "exploratorio, n" not in (out / "analisis_piloto.md").read_text(encoding="utf-8"))
        check("diferencias normales (Shapiro p > 0,05) -> t pareada", R["tiempo"]["contraste"]["normal"] and "t de Student" in R["tiempo"]["contraste"]["prueba"], str(R["tiempo"]["contraste"]))
        t_ref = stats.ttest_rel(pre, post, alternative="greater")
        check("la t pareada coincide con scipy", abs(R["tiempo"]["contraste"]["p"] - t_ref.pvalue) < 1e-12 and abs(R["tiempo"]["contraste"]["estadistico"] - t_ref.statistic) < 1e-9)
        check("reducción de medias y d_z coinciden con el cálculo directo", abs(R["tiempo"]["reduccion_medias"] - (pre.mean() - post.mean()) / pre.mean()) < 1e-12
              and abs(R["tiempo"]["contraste"]["efecto"] - d.mean() / d.std(ddof=1)) < 1e-9)
        s = R["satisfaccion"]
        check("P10 contra ítem 9: mismo procedimiento (H1: ítem 9 mayor)", s["contraste"]["n"] == 60 and abs(s["media_item9"] - items9[:, 8].mean()) < 1e-12 and s["contraste"]["p"] < 0.05)
        check("el alfa de Cronbach (ítems 1–8) coincide con numpy (matriz de covarianza)", abs(s["alfa_cronbach"] - alfa_ref) < 1e-9 and not s["alfa_difiere"], f"{s['alfa_cronbach']} vs {alfa_ref}")
        pf = R["prueba_final"]
        check("prueba final: 180 consultas de 60 sesiones, IC por bootstrap de sesiones y umbral ausente", pf["n_consultas"] == 180 and pf["n_sesiones"] == 60
              and pf["ic95_accuracy_bootstrap_por_sesion"][0] < pf["accuracy"] < pf["ic95_accuracy_bootstrap_por_sesion"][1] and pf["umbral"] is None)
        check("genera .json, .md, dos figuras y el detalle sin el texto de las consultas",
              all((out / f).exists() for f in ("analisis_piloto.json", "analisis_piloto.md", "tiempos_cajas.png", "tiempos_histograma_diferencias.png", "prueba_final_detalle.csv"))
              and "consulta" not in pd.read_csv(out / "prueba_final_detalle.csv").columns)
        det = pd.read_csv(out / "prueba_final_detalle.csv")
        check("la fila de ejemplo EJ01 y la fila no elegible SA99 quedan excluidas", R["n"] == 60 and "EJ01" not in set(det["sesion"]) and "SA99" not in set(det["sesion"]) and R["sesiones_excluidas_no_SA"] == 1)

    print("\nSegunda prueba final")
    c, t, _ = correr(["--registro", str(reg), "--modelo-congelado", str(fz), "--salida", str(out)], predictor)
    check("repetir la prueba final se niega sin --motivo", c != 0 and "una sola vez" in t, t[-300:])
    c, t, _ = correr(["--registro", str(reg), "--modelo-congelado", str(fz), "--salida", str(out), "--motivo", "repetición de prueba falsa"], predictor)
    check("con --motivo se permite y queda en el log", c == 0 and len((out / "evaluaciones_prueba_final.log").read_text(encoding="utf-8").splitlines()) == 2, t[-300:])

    # ------------------------------------------------------------------ no normal
    print("\nDiferencias no normales")
    filas2, pred2, items2, (pre2, post2, _) = datos_falsos(60, normal=False, semilla=11)
    sw = stats.shapiro(pre2 - post2).pvalue
    reg2 = W / "registro_no_normal.xlsx"
    escribir_registro(reg2, filas2, alfa_numpy(items2[:, :8]))
    c, t, R2 = correr(["--registro", str(reg2), "--modelo-congelado", str(fz), "--salida", str(W / "salida_nn")], lambda x: [pred2[q] for q in x])
    w_ref = stats.wilcoxon(pre2, post2, alternative="greater")
    check("los datos falsos son realmente no normales (Shapiro p < 0,05)", sw < 0.05, f"p={sw}")
    check("diferencias no normales -> Wilcoxon (igual a scipy)", c == 0 and R2 and "Wilcoxon" in R2["tiempo"]["contraste"]["prueba"] and abs(R2["tiempo"]["contraste"]["p"] - w_ref.pvalue) < 1e-12, t[-300:])

    # ------------------------------------------------------------------ alfa distinto
    print("\nAlfa distinto del de Resumen")
    reg3 = W / "registro_alfa.xlsx"
    escribir_registro(reg3, filas, alfa_ref + 0.05)
    c, t, R3 = correr(["--registro", str(reg3), "--modelo-congelado", str(fz), "--salida", str(W / "salida_alfa")], predictor)
    check("avisa si el alfa difiere de Resumen!B24 en más de 0,01", c == 0 and R3 and R3["satisfaccion"]["alfa_difiere"] and "AVISO" in t, t[:300])

    # ------------------------------------------------------------------ n < 60
    print("\nn < 60")
    filas4 = filas[:45]
    reg4 = W / "registro_45.xlsx"
    escribir_registro(reg4, filas4, alfa_numpy(items9[:45, :8]))
    c, t, R4 = correr(["--registro", str(reg4), "--modelo-congelado", str(fz), "--salida", str(W / "salida_45")], predictor)
    md = (W / "salida_45" / "analisis_piloto.md").read_text(encoding="utf-8") if c == 0 else ""
    check("con 45 sesiones queda rotulado «exploratorio, n = 45»", c == 0 and R4["exploratorio"] and "exploratorio, n = 45" in md and "exploratorio, n = 45" in t, t[-300:])

    # ------------------------------------------------------------------ rechazos
    print("\nRechazos")
    reg5 = W / "registro_SIMULADO.xlsx"
    escribir_registro(reg5, filas, alfa_ref, hojas_extra=("Datos SIMULADO",))
    c, t, _ = correr(["--registro", str(reg5), "--modelo-congelado", str(fz), "--salida", str(W / "salida_sim")], predictor)
    check("un registro con una hoja «SIMULADO» se rechaza", c != 0 and "SIMULADO" in t and not (W / "salida_sim").exists(), t[-300:])
    sim_out = W / "salida_sim2"
    c, t, R5 = correr(["--registro", str(reg5), "--modelo-congelado", str(fz), "--salida", str(sim_out), "--permitir-simulado"], predictor)
    check("con --permitir-simulado escribe solo en carpeta temporal y rotula _SIMULADO",
          c == 0 and R5 and R5["rotulo"] == "SIMULADO" and (sim_out / "analisis_piloto_SIMULADO.json").exists() and not (sim_out / "analisis_piloto.json").exists(), t[-300:])
    fuera = W.parent / "carpeta_fuera_de_temp_no_usar"
    c, t, _ = correr(["--registro", str(reg5), "--modelo-congelado", str(fz), "--salida", str(ROOT / "logs" / "piloto"), "--permitir-simulado"], predictor)
    check("--permitir-simulado no escribe en el repositorio aunque se lo pidan", c == 0 and "carpeta temporal" in t and not (ROOT / "logs" / "piloto").exists(), t[:300])
    reg6 = W / "registro_sin_filas.xlsx"
    escribir_registro(reg6, [["SA01", "No"] + list(filas[0][2:])], alfa_ref)
    c, t, _ = correr(["--registro", str(reg6), "--modelo-congelado", str(fz), "--salida", str(W / "salida_sf")], predictor)
    check("sin filas elegibles sale con error", c != 0 and "no hay filas elegibles" in t, t[-300:])
    reg7 = W / "registro_formulas.xlsx"
    escribir_registro(reg7, filas, alfa_ref, formulas_sin_valor=True)
    c, t, _ = correr(["--registro", str(reg7), "--modelo-congelado", str(fz), "--salida", str(W / "salida_fx")], predictor)
    check("fórmulas sin valor guardado: «abre y guarda el archivo en Excel»", c != 0 and "abre y guarda el archivo en Excel" in t, t[-300:])

    print("\nHash del modelo distinto")
    modelo.write_bytes(b"modelo falso MODIFICADO")
    c, t, _ = correr(["--registro", str(reg), "--modelo-congelado", str(fz), "--salida", str(W / "salida_hash")], predictor)
    check("si el hash del modelo cambia, la prueba final aborta", c != 0 and "hash del modelo" in t and not (W / "salida_hash" / "analisis_piloto.json").exists(), t[-300:])
    c, t = congelar(["--verificar", "--salida", str(fz)])
    check("--verificar detecta el cambio del modelo", c == 1 and "modelo" in t, t)

    # ------------------------------------------------------------------ integridad
    print("\nIntegridad del repositorio")
    despues = snapshot()
    cambios = sorted(k for k in set(antes) | set(despues) if antes.get(k) != despues.get(k))
    check("incident_log.csv, logs/ y docs/piloto/ no cambiaron durante la prueba", not cambios, f"cambiaron: {cambios}")
    ok = sum(RES)
    print(f"\nRESULTADO (SIMULADA): {ok}/{len(RES)} comprobaciones PASS" + ("" if ok == len(RES) else "  <- HAY FALLAS"))
    print("Recordatorio: datos FALSOS y predictor falso; nada de esto es un resultado del piloto.")
    if not a.conservar:
        shutil.rmtree(W, ignore_errors=True)
    sys.exit(0 if ok == len(RES) else 1)


if __name__ == "__main__":
    main()
