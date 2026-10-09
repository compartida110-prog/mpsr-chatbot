"""Prueba de humo del tablero de compuertas (scripts/estado_compuertas.py) y del modo --demo-simulada — rótulo: SIMULADA.

Usa datos falsos en una carpeta temporal (árboles de datos falsos para el tablero y --demo-raiz para el modo de demostración): nada se escribe en el repositorio, y al final se
verifica que corpus/real, logs/v3_real, logs/avance, evidencias/simulado_demostracion, docs/lote_real_1 y docs/piloto no cambiaron.

Comprueba:
  * el tablero sin evidencia: las siete compuertas Pendiente y «0 compuertas cumplidas con datos reales»;
  * una compuerta cumplida con datos simulados («Cumplida (Simulado)») que no suma a las reales, y una cumplida con datos reales que sí suma;
  * G3 con más de 2 ciclos de refinamiento, con el test evaluado dos veces o con F1 < 0,75: «No cumplida»;
  * G5 con el congelamiento anterior a G4 (o a G3, o con un archivo alterado): «No cumplida»; y bien ordenado: «Cumplida»;
  * G2 (marca de revisión de datos personales), G6 (alfa) y G7 (60 sesiones; cierre declarado con ≥ 30);
  * que el tablero solo escribe en logs/avance/ y se niega a escribir en otra carpeta;
  * --demo-simulada: escribe solo en evidencias/simulado_demostracion/, con el marcador SIMULADO en cada archivo, solo con las entradas permitidas, y se NIEGA a escribir
    en corpus/real, logs/v3_real y las carpetas privadas; no congela el modelo real;
  * que sin --demo-simulada el comportamiento anterior no cambia.

Uso:
    python tests/smoke_compuertas.py [--conservar]
"""
import argparse
import hashlib
import json
import os
import shutil
import subprocess
import sys
import tempfile
import warnings
from pathlib import Path

warnings.filterwarnings("ignore")
import numpy as np  # noqa: E402
from openpyxl import Workbook  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
SCRIPTS = ROOT / "scripts"
sys.path.insert(0, str(SCRIPTS))
sys.path.insert(0, str(Path(__file__).resolve().parent))
import deteccion_simulado as ds  # noqa: E402
import smoke_piloto as sp  # noqa: E402  (constructor de registros falsos)

LIBRO = ROOT / "docs" / "lote_real_1" / "ejemplos_simulados" / "Lote1_Transcripcion_SIMULADO_v2.xlsx"
REGISTRO = ROOT / "docs" / "piloto" / "ejemplos_simulados" / "Registro_Sesiones_Piloto_SIMULADO_v4.xlsx"
PLANTILLA_LOTE = ROOT / "docs" / "lote_real_1" / "Lote1_Transcripcion_v1.xlsx"
ENV = {**os.environ, "PYTHONIOENCODING": "utf-8", "TF_CPP_MIN_LOG_LEVEL": "2"}
RES = []
LISTO_SI = "¿Listo para la Parte B? SÍ (transcritos 25/15, intenciones con menos de 3 frases: 0)"
LISTO_NO = "¿Listo para la Parte B? NO (transcritos 9/15, intenciones con menos de 3 frases: 4)"


def check(nombre, cond, detalle=""):
    RES.append(bool(cond))
    print(f"  [{'PASS' if cond else 'FAIL'}] {nombre}" + (f"  ({str(detalle)[:300]})" if detalle and not cond else ""))


def run(script, *args):
    p = subprocess.run([sys.executable, str(SCRIPTS / script), *map(str, args)], capture_output=True, text=True, encoding="utf-8", errors="replace", env=ENV, cwd=ROOT)
    return p.returncode, (p.stdout or "") + (p.stderr or "")


def fz_real_hash():
    """Huella del congelamiento REAL (G5), si existe: las pruebas con datos falsos no deben crearlo ni modificarlo (puede existir de verdad desde el congelamiento del lote 2)."""
    f = ROOT / "logs" / "v3_real" / "modelo_congelado.json"
    return hashlib.sha256(f.read_bytes()).hexdigest() if f.exists() else None


FZ_REAL_0 = fz_real_hash()


def escribir(raiz, rel, texto):
    f = Path(raiz) / rel
    f.parent.mkdir(parents=True, exist_ok=True)
    f.write_text(texto, encoding="utf-8")
    return f


def tablero(raiz):
    c, t = run("estado_compuertas.py", "--raiz", raiz)
    j = json.loads((Path(raiz) / "logs" / "avance" / "estado_compuertas.json").read_text(encoding="utf-8")) if c == 0 else {}
    return c, t, {x["id"]: x for x in j.get("compuertas", [])}, j


def snapshot(raiz, excluir=("logs/avance",)):
    out = {}
    for p in sorted(Path(raiz).rglob("*")):
        if p.is_file() and not any(p.relative_to(raiz).as_posix().startswith(e) for e in excluir):
            out[p.relative_to(raiz).as_posix()] = hashlib.sha256(p.read_bytes()).hexdigest()
    return out


def snapshot_repo():
    out = {}
    for d in ("corpus/real", "logs/v3_real", "logs/avance", "evidencias/simulado_demostracion", "docs/lote_real_1", "docs/piloto"):
        base = ROOT / d
        if base.exists():
            for p in sorted(base.rglob("*")):
                if p.is_file() and "__pycache__" not in p.parts:
                    out[p.relative_to(ROOT).as_posix()] = hashlib.sha256(p.read_bytes()).hexdigest()
    return out


def tupa_falso(ruta, alta, alertas, sinconf, fecha, antiguo=False, confirmadas=0):
    wb = Workbook()
    ws = wb.active
    ws.title = "Verificacion"
    ws["A1"] = "Verificación FALSA de prueba"
    for j, h in enumerate(["#", "Intención", "Categoría", "Prioridad"] + ["x"] * 11 + ["Fecha de verificación"], 1):
        ws.cell(3, j, h)
    ws.cell(4, 1, 1); ws.cell(4, 16, fecha)
    rs = wb.create_sheet("Resumen")
    rs["A1"] = "Resumen de la verificación (fila de título de una sola celda: la v10 real las trae)"
    filas = [("Respuestas a verificar", 44), ("Verificadas (todo menos Pendiente)" if antiguo else "Con resultado (propuesto o confirmado)", 44 - alta)]
    if not antiguo:
        filas.append(("Confirmadas por el tesista", confirmadas))
    filas += [("Prioridad Alta pendientes", alta), ("Filas con alerta", alertas), ("Corregir o Coincide sin confirmar por el tesista", sinconf)]
    for i, (k, v) in enumerate(filas, 5):
        rs.cell(i, 1, k); rs.cell(i, 2, v)
    ruta.parent.mkdir(parents=True, exist_ok=True)
    wb.save(ruta)


def congelado_falso(raiz, fecha, alterar=False):
    d = Path(raiz) / "logs" / "v3_real"
    d.mkdir(parents=True, exist_ok=True)
    arch = {}
    for n, txt in (("modelo", "modelo falso"), ("config", "config falsa"), ("dominio", "dominio falso")):
        f = d / f"{n}_falso.txt"
        f.write_text(txt, encoding="utf-8")
        arch[n] = str(f)
    fz = {"fecha": fecha, "archivos": arch, "sha256": {n: hashlib.sha256(Path(p).read_bytes()).hexdigest() for n, p in arch.items()}}
    (d / "modelo_congelado.json").write_text(json.dumps(fz), encoding="utf-8")
    if alterar:
        Path(arch["modelo"]).write_text("modelo falso ALTERADO", encoding="utf-8")


def g3_falso(raiz, ciclos=2, evaluaciones=1, f1=0.80, fecha="2026-02-01T10:00:00"):
    escribir(raiz, "logs/avance/ciclos_refinamiento.csv", "ciclo,fecha,f1_macro_validacion\n" + "".join(f"{i},2026-01-0{i},{0.60 + 0.05 * i:.2f}\n" for i in range(1, ciclos + 1)))
    reg = {"evaluaciones": [{"fecha": fecha, "metodo": "rasa", "motivo": "evaluación final única"}] * evaluaciones}
    if evaluaciones:
        escribir(raiz, "logs/v3_real/test_registro.json", json.dumps(reg))
        escribir(raiz, "logs/v3_real/eval_real_resumen.json", json.dumps({"metodos": {"rasa": {"f1_macro": [f1, f1 - 0.05, f1 + 0.05]}}}))


def marcado_ok(f):
    suf = Path(f).suffix.lower()
    if "_SIMULADO" not in Path(f).name:
        return False
    if suf == ".csv":
        filas = [l.split(",") for l in Path(f).read_text(encoding="utf-8").splitlines() if l.strip()]
        return filas[0][-1] == "ESTADO" and all(r[-1].startswith("SIMULADO") for r in filas[1:])
    if suf in (".txt", ".md", ".log"):
        return Path(f).read_text(encoding="utf-8").splitlines()[0] == ds.MARCA_ESTADO
    if suf == ".json":
        return "ESTADO" in json.loads(Path(f).read_text(encoding="utf-8"))
    if suf == ".png":
        from PIL import Image
        with Image.open(f) as im:
            return str(im.text.get("ESTADO", "")).startswith("SIMULADO")
    return b"SIMULADO" in Path(f).read_bytes()


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--conservar", action="store_true")
    a = ap.parse_args()
    W = Path(tempfile.mkdtemp(prefix="compuertas_PRUEBA_SIMULADA_"))
    print(f"SIMULADA — carpeta temporal (datos FALSOS, no se commitean): {W}")
    antes = snapshot_repo()

    # ------------------------------------------------------------------ tablero sin evidencia
    print("\nTablero sin evidencia")
    r0 = W / "t0"
    r0.mkdir()
    c, t, g, j = tablero(r0)
    check("sin evidencia: las siete compuertas están «Pendiente (sin evidencia)»", c == 0 and len(g) == 7 and all(x["estado"] == "Pendiente (sin evidencia)" for x in g.values()), t[-300:])
    check("el encabezado dice «0 compuertas cumplidas con datos reales» y el JSON también", "Compuertas cumplidas con datos reales: 0 de 7" in t and j["compuertas_cumplidas_con_datos_reales"] == 0)
    check("el tablero solo escribe en logs/avance/ (estado_compuertas.md y .json)", sorted(snapshot(r0, excluir=()).keys()) == ["logs/avance/estado_compuertas.json", "logs/avance/estado_compuertas.md"])
    r_otro = W / "t_otro"
    r_otro.mkdir()
    c, t = run("estado_compuertas.py", "--raiz", r_otro, "--salida", W / "otra_carpeta")
    check("--salida fuera de logs/avance/ se rechaza y no escribe nada", c != 0 and "ME NIEGO" in t and not (W / "otra_carpeta").exists() and not snapshot(r_otro, excluir=()), t[-200:])

    # ------------------------------------------------------------------ simulado frente a real
    print("\nDatos simulados frente a reales")
    r1 = W / "t_sim"
    escribir(r1, "evidencias/simulado_demostracion/20260101_000000_ingest/ingesta_reporte_SIMULADO.txt", ds.MARCA_ESTADO + "\nINGESTA DEL LOTE 1 — REPORTE\n" + LISTO_SI + "\n")
    escribir(r1, "evidencias/simulado_demostracion/20260101_000100_piloto/analisis_piloto_SIMULADO.json", json.dumps({"n": 60}))
    c, t, g, j = tablero(r1)
    check("una compuerta cumplida con datos simulados aparece «Cumplida (Simulado)»", g["G1"]["estado"] == "Cumplida (Simulado)" and g["G1"]["datos"] == "Simulado" and g["G7"]["estado"] == "Cumplida (Simulado)", str(g["G1"]))
    check("y no suma a las reales: «0 compuertas cumplidas con datos reales»; las simuladas se cuentan aparte", j["compuertas_cumplidas_con_datos_reales"] == 0 and j["compuertas_cumplidas_con_datos_simulados"] == 2 and "reales: 0 de 7" in t)
    r2 = W / "t_real"
    escribir(r2, "logs/v3_real/ingesta_reporte.txt", "INGESTA DEL LOTE 1 — REPORTE\n" + LISTO_SI + "\n")
    c, t, g, j = tablero(r2)
    check("la misma evidencia en el lugar real (sin marca de simulación) es «Cumplida» con datos Reales y suma 1", g["G1"]["estado"] == "Cumplida" and g["G1"]["datos"] == "Real" and j["compuertas_cumplidas_con_datos_reales"] == 1, str(g["G1"]))
    r3 = W / "t_real_marcado"
    escribir(r3, "logs/v3_real/ingesta_reporte.txt", ds.MARCA_ESTADO + "\n" + LISTO_SI + "\n")
    c, t, g, j = tablero(r3)
    check("un reporte marcado SIMULADO que alguien deja en la carpeta real no cuenta como real", g["G1"]["estado"] == "Cumplida (Simulado)" and j["compuertas_cumplidas_con_datos_reales"] == 0, str(g["G1"]))
    r3b = W / "t_g1_no"
    escribir(r3b, "logs/v3_real/ingesta_reporte.txt", "INGESTA DEL LOTE 1 — REPORTE\n" + LISTO_NO + "\n")
    c, t, g, _ = tablero(r3b)
    check("G1 con «Listo para la Parte B? NO» queda «En curso»", g["G1"]["estado"] == "En curso", str(g["G1"]))

    # ------------------------------------------------------------------ G2
    print("\nG2: revisión de datos personales")
    r4 = W / "t_g2"
    escribir(r4, "logs/v3_real/ingesta_reporte.txt", "INGESTA DEL LOTE 1 — REPORTE\n" + LISTO_SI + "\n")
    c, t, g, _ = tablero(r4)
    check("ingesta sin bloqueos pero sin la marca del tesista: G2 «En curso»", g["G2"]["estado"] == "En curso" and "revision_pii.txt" in g["G2"]["nota"], str(g["G2"]))
    escribir(r4, "logs/avance/revision_pii.txt", "Revisé el reporte de datos personales. Luis Mario, 2026-10-05\n")
    c, t, g, _ = tablero(r4)
    check("con logs/avance/revision_pii.txt: G2 «Cumplida»", g["G2"]["estado"] == "Cumplida", str(g["G2"]))

    # ------------------------------------------------------------------ G3
    print("\nG3: ciclos y evaluación única del test")
    for nombre, kw, esperado in (("más de 2 ciclos", dict(ciclos=3, evaluaciones=0), "No cumplida"), ("test evaluado 2 veces", dict(ciclos=2, evaluaciones=2), "No cumplida"),
                                  ("F1 de la única evaluación = 0,70", dict(ciclos=2, evaluaciones=1, f1=0.70), "No cumplida"), ("2 ciclos y F1 = 0,80 en una sola evaluación", dict(ciclos=2, evaluaciones=1, f1=0.80), "Cumplida"),
                                  ("2 ciclos y el test aún sin evaluar", dict(ciclos=2, evaluaciones=0), "En curso")):
        rg = W / ("t_g3_" + hashlib.md5(nombre.encode()).hexdigest()[:6])
        g3_falso(rg, **kw)
        c, t, g, _ = tablero(rg)
        check(f"G3 con {nombre}: «{esperado}»", g["G3"]["estado"] == esperado, str(g["G3"]))

    print("\nG3 con lote 2 (V1.6): «Medida en lote 2» aparte de la medición del lote 1")

    def l2(nombre, f1=None, evaluaciones=1, congelado="2026-02-15T10:00:00", eval_fecha="2026-03-01T10:00:00"):
        r = W / ("t_l2_" + nombre)
        g3_falso(r, f1=0.70)  # lote 1: una sola evaluación con F1 = 0,70 -> «No cumplida, lote 1»
        if evaluaciones:
            escribir(r, "logs/v3_real/lote2/test_registro.json", json.dumps({"evaluaciones": [{"fecha": eval_fecha, "metodo": "rasa"}] * evaluaciones}))
            if f1 is not None:
                escribir(r, "logs/v3_real/lote2/eval_lote2_resumen.json", json.dumps({"metodos": {"rasa": {"f1_macro": [f1, f1 - 0.05, f1 + 0.05]}}}))
        if congelado:
            escribir(r, "logs/v3_real/lote2_congelado_previo.json", json.dumps({"fecha": congelado}))
        return tablero(r)[2]["G3"]
    x = l2("sin_lote2", evaluaciones=0, congelado="")
    check("G3 sin datos del lote 2: «No cumplida» con «No cumplida, lote 1», y el lote 2 «planificado, sin medición todavía»", x["estado"] == "No cumplida" and x["nota"].startswith("No cumplida, lote 1.") and "planificado, sin medición todavía" in x["nota"], str(x))
    x = l2("ok", f1=0.80)
    check("G3 con lote 2: F1 0,80 una sola vez y congelado antes: «Cumplida»; «Medida en lote 2» aparte y el lote 1 sigue como «No cumplida, lote 1»", x["estado"] == "Cumplida" and "Medida en lote 2: F1 macro = 0.8000" in x["nota"] and "No cumplida, lote 1" in x["nota"], str(x))
    x = l2("bajo", f1=0.70)
    check("G3 con lote 2: F1 0,70 una sola vez: «No cumplida» (no se baja el umbral ni se repite)", x["estado"] == "No cumplida" and "Medida en lote 2: F1 macro = 0.7000" in x["nota"] and "no se baja el umbral" in x["nota"], str(x))
    x = l2("dos", f1=0.80, evaluaciones=2)
    check("G3 con el test del lote 2 evaluado 2 veces: «No cumplida»", x["estado"] == "No cumplida" and "2 veces" in x["nota"], str(x))
    x = l2("sin_cong", f1=0.80, congelado="")
    check("G3 con lote 2 medido sin congelamiento previo: «No cumplida» (no cuenta)", x["estado"] == "No cumplida" and "no consta el congelamiento previo" in x["nota"], str(x))
    x = l2("cong_tarde", f1=0.80, congelado="2026-04-01T10:00:00")
    check("G3 con el congelamiento posterior a la evaluación del lote 2: «No cumplida» (no cuenta)", x["estado"] == "No cumplida" and "posterior" in x["nota"], str(x))
    r = W / "t_l2_lote1_ok"
    g3_falso(r, f1=0.80)
    escribir(r, "logs/v3_real/lote2/test_registro.json", json.dumps({"evaluaciones": [{"fecha": "2026-03-01T10:00:00", "metodo": "rasa"}]}))
    x = tablero(r)[2]["G3"]
    check("si el lote 1 ya cumplió (F1 ≥ 0,75), el lote 2 no cambia el resultado (reglas V1.2 intactas)", x["estado"] == "Cumplida" and "lote 2" not in x["nota"].lower(), str(x))

    # ------------------------------------------------------------------ G4 y G5
    print("\nG4 y G5: respuestas verificadas y orden del congelamiento")
    rg4 = W / "t_g4"
    tupa_falso(rg4 / "docs" / "tupa" / "Verificacion_TUPA_v5.xlsx", 0, 0, 0, "2026-03-01", confirmadas=10)
    tupa_falso(rg4 / "docs" / "tupa" / "Verificacion_TUPA_v4.xlsx", 5, 0, 0, "2026-03-01", antiguo=True)  # la versión más alta es la que cuenta
    c, t, g, _ = tablero(rg4)
    check("G4 con Alta pendientes = 0, alertas = 0 y sin confirmar = 0 (versión más alta): «Cumplida»", g["G4"]["estado"] == "Cumplida" and "v5" in g["G4"]["evidencia"] and "confirmadas por el tesista: 10" in g["G4"]["nota"], str(g["G4"]))
    rg4b = W / "t_g4b"
    tupa_falso(rg4b / "docs" / "tupa" / "Verificacion_TUPA_v5.xlsx", 23, 0, 10, "2026-03-01")
    c, t, g, _ = tablero(rg4b)
    check("G4 con 23 Alta pendientes y 10 sin confirmar: «En curso», y dice «con resultado» (propuesto), no «verificadas»", g["G4"]["estado"] == "En curso" and "con resultado, propuesto o confirmado: 21 de 44" in g["G4"]["nota"] and "verificadas" not in g["G4"]["nota"], str(g["G4"]))
    rg4c = W / "t_g4c"
    tupa_falso(rg4c / "docs" / "tupa" / "Verificacion_TUPA_v4.xlsx", 3, 0, 0, "2026-03-01", antiguo=True)
    c, t, g, _ = tablero(rg4c)
    check("una hoja con los rótulos de la v4 (respaldo) también se lee: «En curso» con 3 Alta pendientes", g["G4"]["estado"] == "En curso" and "con resultado, propuesto o confirmado: 41 de 44" in g["G4"]["nota"], str(g["G4"]))

    def arbol_g5(nombre, fecha_congelado, alterar=False, con_g3=True):
        r = W / nombre
        if con_g3:
            g3_falso(r, fecha="2026-02-01T10:00:00")
        tupa_falso(r / "docs" / "tupa" / "Verificacion_TUPA_v5.xlsx", 0, 0, 0, "2026-03-01")
        congelado_falso(r, fecha_congelado, alterar)
        return r
    c, t, g, _ = tablero(arbol_g5("t_g5_antes_g4", "2026-01-01T09:00:00"))
    check("G5 con el congelamiento ANTERIOR a G3 y a G4: «No cumplida»", g["G3"]["estado"] == "Cumplida" and g["G4"]["estado"] == "Cumplida" and g["G5"]["estado"] == "No cumplida", str(g["G5"]))
    c, t, g, _ = tablero(arbol_g5("t_g5_entre", "2026-02-15T09:00:00"))
    check("G5 con el congelamiento después de G3 pero ANTERIOR a G4: «No cumplida»", g["G5"]["estado"] == "No cumplida" and "anterior" in g["G5"]["nota"], str(g["G5"]))
    c, t, g, j = tablero(arbol_g5("t_g5_bien", "2026-04-01T09:00:00"))
    check("G5 congelado después de G3 y G4: «Cumplida», y G3, G4 y G5 suman 3 reales", g["G5"]["estado"] == "Cumplida" and j["compuertas_cumplidas_con_datos_reales"] == 3, str(g["G5"]))
    c, t, g, _ = tablero(arbol_g5("t_g5_alterado", "2026-04-01T09:00:00", alterar=True))
    check("G5 con un archivo del modelo alterado (huella distinta): «No cumplida»", g["G5"]["estado"] == "No cumplida" and "huella" in g["G5"]["nota"], str(g["G5"]))
    c, t, g, _ = tablero(arbol_g5("t_g5_sin_g3", "2026-04-01T09:00:00", con_g3=False))
    check("G5 congelado sin G3 cumplida: «No cumplida» (se congela después de G3 y G4)", g["G5"]["estado"] == "No cumplida", str(g["G5"]))

    # ------------------------------------------------------------------ G6 y G7 con registros falsos
    print("\nG6 y G7 con registros falsos")
    def registro(raiz, nombre, n, semilla):
        filas, _, items9, _ = sp.datos_falsos(n, semilla=semilla)
        ruta = Path(raiz) / "docs" / "piloto" / "privado" / nombre
        ruta.parent.mkdir(parents=True, exist_ok=True)
        sp.escribir_registro(ruta, filas, 0.0)
        return items9
    r6 = W / "t_g6"
    items = registro(r6, "Registro_Sesiones_Prepiloto_prueba.xlsx", 8, 3)
    alfa = sp.alfa_numpy(items[:, :8])
    c, t, g, _ = tablero(r6)
    esperado = "Cumplida" if alfa >= 0.70 else "No cumplida"
    check(f"G6 con 8 sesiones y alfa {alfa:.2f} (calculado aparte con numpy): «{esperado}»", g["G6"]["estado"] == esperado and f"{alfa:.3f}" in g["G6"]["nota"], str(g["G6"]))
    check("G6 avisa que con n pequeño el alfa es poco estable", "advertencia" in g["G6"]["nota"], str(g["G6"]))
    r6b = W / "t_g6b"
    registro(r6b, "Registro_Sesiones_Prepiloto_prueba.xlsx", 3, 3)
    c, t, g, _ = tablero(r6b)
    check("G6 con 3 sesiones (menos de 5): «En curso»", g["G6"]["estado"] == "En curso", str(g["G6"]))
    for n, cierre, esperado, nombre in ((60, False, "Cumplida", "60 sesiones elegibles"), (40, True, "Cumplida", "40 sesiones y cierre declarado"), (20, True, "No cumplida", "20 sesiones y cierre declarado (menos de 30)"),
                                        (40, False, "En curso", "40 sesiones sin cierre declarado")):
        r7 = W / f"t_g7_{n}_{int(cierre)}"
        registro(r7, "Registro_Sesiones_Piloto_prueba.xlsx", n, 5)
        if cierre:
            escribir(r7, "logs/avance/cierre_piloto.txt", "Declaro el cierre de la recolección. Luis Mario\n")
        c, t, g, j = tablero(r7)
        check(f"G7 con {nombre}: «{esperado}»", g["G7"]["estado"] == esperado and (esperado != "Cumplida" or j["compuertas_cumplidas_con_datos_reales"] == 1), str(g["G7"]))

    # ------------------------------------------------------------------ modo de demostración
    print("\nModo --demo-simulada")
    rd = W / "raiz_demo"
    rd.mkdir()
    libro = W / "Lote1_Transcripcion_SIMULADO_v2.xlsx"
    registro_demo = W / "Registro_Sesiones_Piloto_SIMULADO_v4.xlsx"
    shutil.copy(LIBRO, libro)
    shutil.copy(REGISTRO, registro_demo)
    base_demo = rd / "evidencias" / "simulado_demostracion"
    check("la carpeta de demostración por defecto es evidencias/simulado_demostracion/ del repositorio", ds.base_demo() == ROOT / "evidencias" / "simulado_demostracion")
    c, t = run("ingest_real_lote.py", "--libro", libro, "--demo-simulada", "--demo-raiz", rd, "--ejecucion", "ingesta")
    d1 = base_demo / "ingesta"
    check("ingesta con --demo-simulada: código 0 y escribe en evidencias/simulado_demostracion/<ejecución>/", c == 0 and d1.is_dir() and any(d1.glob("*.csv")), t[-400:])
    c, t = run("analizar_piloto.py", "--registro", registro_demo, "--demo-simulada", "--demo-raiz", rd, "--ejecucion", "piloto")
    d2 = base_demo / "piloto"
    check("análisis del piloto con --demo-simulada: código 0 y escribe en la carpeta de demostración", c == 0 and (d2 / "analisis_piloto_SIMULADO.json").exists(), t[-400:])
    c, t = run("congelar_modelo.py", "--demo-simulada", "--demo-raiz", rd, "--ejecucion", "congelado")
    d3 = base_demo / "congelado"
    check("congelar_modelo.py --demo-simulada escribe modelo_congelado_SIMULADO.json con un modelo de demostración de otro nombre",
          c == 0 and (d3 / "modelo_congelado_SIMULADO.json").exists() and (d3 / "modelo_demostracion_SIMULADO.tar.gz").exists(), t[-300:])
    archivos = [f for d in (d1, d2, d3) for f in d.rglob("*") if f.is_file()]
    sin_marca = [f.name for f in archivos if not marcado_ok(f)]
    check(f"los {len(archivos)} archivos generados llevan el sufijo _SIMULADO y el marcador «{ds.MARCA_ESTADO}» (línea inicial, campo ESTADO o metadatos)", archivos and not sin_marca, sin_marca)
    check("los CSV llevan el sufijo _SIMULADO y la columna ESTADO", all(f.name.endswith("_SIMULADO.csv") for f in archivos if f.suffix == ".csv") and any(f.suffix == ".csv" for f in archivos))
    check("el modo de demostración escribe solo bajo evidencias/simulado_demostracion/ (no crea nada más en la raíz)", {p.relative_to(rd).parts[:2] for p in rd.rglob("*") if p.is_file()} == {("evidencias", "simulado_demostracion")})
    fz = json.loads((d2 / "modelo_congelado_SIMULADO.json").read_text(encoding="utf-8"))
    check("el modelo de demostración no es el real: otro nombre de archivo y estado SIMULADO; no se crea ni se modifica logs/v3_real/modelo_congelado.json", "SIMULADO" in fz["estado"] and "demostracion" in fz["modelo_nombre"] and fz_real_hash() == FZ_REAL_0)
    c, t, g, j = tablero(rd)
    check("el tablero lee la carpeta de demostración y la muestra siempre como Simulado (G1 y G7 «Cumplida (Simulado)»; 0 reales)",
          g["G1"]["estado"] == "Cumplida (Simulado)" and g["G7"]["estado"] == "Cumplida (Simulado)" and j["compuertas_cumplidas_con_datos_reales"] == 0 and g["G1"]["datos"] == "Simulado", str({k: v["estado"] for k, v in g.items()}))

    print("\nNegativas del modo de demostración")
    for carpeta in ("corpus/real", "logs/v3_real", "docs/lote_real_1/privado", "docs/piloto/privado"):
        c, t = run("ingest_real_lote.py", "--libro", libro, "--demo-simulada", "--out-dir", ROOT / carpeta, "--log-dir", ROOT / carpeta, "--demo-raiz", rd)
        check(f"--demo-simulada se NIEGA a escribir en {carpeta}", c != 0 and "ME NIEGO" in t and "REALES" in t, t[-200:])
    c, t = run("ingest_real_lote.py", "--libro", libro, "--demo-simulada", "--out-dir", W / "fuera", "--demo-raiz", rd)
    check("--demo-simulada tampoco escribe fuera de evidencias/simulado_demostracion/ (aunque no sea una carpeta real)", c != 0 and "ME NIEGO" in t and not (W / "fuera").exists(), t[-200:])
    c, t = run("analizar_piloto.py", "--registro", registro_demo, "--demo-simulada", "--salida", ROOT / "logs" / "v3_real", "--demo-raiz", rd)
    check("analizar_piloto.py --demo-simulada se NIEGA a escribir en logs/v3_real", c != 0 and "ME NIEGO" in t, t[-200:])
    c, t = run("congelar_modelo.py", "--demo-simulada", "--salida", ROOT / "logs" / "v3_real" / "modelo_congelado_otro.json", "--demo-raiz", rd)
    check("congelar_modelo.py --demo-simulada se NIEGA a escribir en logs/v3_real", c != 0 and "ME NIEGO" in t and not (ROOT / "logs" / "v3_real" / "modelo_congelado_otro.json").exists(), t[-200:])
    c, t = run("congelar_modelo.py", "--demo-simulada", "--salida", ROOT / "logs" / "v3_real" / "modelo_congelado.json", "--demo-raiz", rd, "--ejecucion", "defecto")
    check("con la ruta real por defecto (indistinguible de no pasarla) escribe en la carpeta de demostración y no toca el modelo congelado real",
          c == 0 and (rd / "evidencias" / "simulado_demostracion" / "defecto" / "modelo_congelado_SIMULADO.json").exists() and fz_real_hash() == FZ_REAL_0, t[-200:])
    ajeno = W / "Lote1_Transcripcion_otro.xlsx"
    shutil.copy(LIBRO, ajeno)
    for nombre, ruta in (("un libro con otro nombre", ajeno), ("la plantilla vacía real", PLANTILLA_LOTE)):
        c, t = run("ingest_real_lote.py", "--libro", ruta, "--demo-simulada", "--demo-raiz", rd, "--ejecucion", "no_permitida")
        check(f"--demo-simulada solo acepta las entradas simuladas permitidas: rechaza {nombre}", c != 0 and "ME NIEGO" in t and not (base_demo / "no_permitida").exists(), t[-200:])
    c, t = run("ingest_real_lote.py", "--demo-simulada", "--demo-raiz", rd)
    check("--demo-simulada sin --libro se rechaza con un mensaje claro", c != 0 and "--libro" in t, t[-200:])

    print("\nSin --demo-simulada el comportamiento anterior no cambia")
    c, t = run("ingest_real_lote.py", "--libro", libro, "--out-dir", W / "sin_opcion", "--log-dir", W / "sin_opcion" / "log")
    check("el libro simulado sin ninguna opción se sigue rechazando (ME NIEGO) y no genera salidas", c != 0 and "ME NIEGO" in t and not (W / "sin_opcion").exists(), t[-200:])
    rd2 = W / "raiz_sin_demo"
    rd2.mkdir()
    c, t = run("ingest_real_lote.py", "--libro", libro, "--permitir-simulado", "--out-dir", W / "perm", "--log-dir", W / "perm" / "log", "--demo-raiz", rd2)
    check("con --permitir-simulado sigue escribiendo solo donde se le indica (carpeta temporal) y sin el marcador de demostración",
          c == 0 and (W / "perm" / "lote1_respuestas.csv").exists() and not any(rd2.rglob("*")) and "ESTADO" not in (W / "perm" / "lote1_respuestas.csv").read_text(encoding="utf-8").splitlines()[0], t[-300:])

    # ------------------------------------------------------------------ piezas de la demostración completa
    print("\nPiezas de la demostración completa (congelado con umbral, análisis, tablero de una ejecución)")
    bd = rd / "evidencias" / "simulado_demostracion"
    c4 = bd / "c4"
    c4.mkdir(parents=True)
    (c4 / "umbral_congelado.json").write_text(json.dumps({"t": 0.5, "ambiguity_threshold": 0.1}), encoding="utf-8")
    c, t = run("congelar_modelo.py", "--demo-simulada", "--demo-sin-marcar", "--demo-raiz", rd, "--salida", c4, "--umbral", c4 / "umbral_congelado.json")
    fz4 = json.loads((c4 / "modelo_congelado_SIMULADO.json").read_text(encoding="utf-8")) if c == 0 else {}
    check("congelar --demo-simulada acepta una carpeta de demostración y el umbral elegido allí (se copia adentro; no se referencia el archivo)",
          c == 0 and fz4.get("umbral_t") == 0.5 and "umbral" not in fz4.get("archivos", {}) and "advertencia" not in fz4, t[-300:])
    check("con la opción interna --demo-sin-marcar los archivos quedan sin marcar (el orquestador marca una sola vez al final)", "ESTADO" not in fz4 and not ds.MARCA_ESTADO in (c4 / "modelo_demostracion_SIMULADO.tar.gz").read_text(encoding="utf-8")[:0] and c == 0)
    c, t = run("congelar_modelo.py", "--demo-simulada", "--demo-raiz", rd, "--salida", c4, "--umbral", W / "umbral_de_afuera.json")
    check("un --umbral fuera de evidencias/simulado_demostracion/ se rechaza en la demostración", c != 0 and "ME NIEGO" in t, t[-200:])
    c, t = run("congelar_modelo.py", "--demo-simulada", "--demo-raiz", rd, "--salida", bd / "c5" / "modelo_congelado_SIMULADO.json")
    check("--salida con un archivo .json de la carpeta de demostración se interpreta como esa carpeta (no crea una carpeta con el nombre del archivo)",
          c == 0 and (bd / "c5" / "modelo_congelado_SIMULADO.json").is_file(), t[-200:])
    c, t = run("analizar_piloto.py", "--registro", registro_demo, "--demo-simulada", "--demo-sin-marcar", "--demo-raiz", rd, "--salida", bd / "c6", "--modelo-congelado", c4 / "modelo_congelado_SIMULADO.json")
    r6 = json.loads((bd / "c6" / "analisis_piloto_SIMULADO.json").read_text(encoding="utf-8")) if c == 0 else {}
    check("analizar_piloto.py --demo-simulada acepta el congelamiento de demostración de otra etapa y usa su umbral (t = 0,5)", c == 0 and (r6.get("prueba_final", {}).get("umbral") or {}).get("t") == 0.5, t[-300:])
    copia_fz = W / "modelo_congelado_SIMULADO.json"
    shutil.copy(c4 / "modelo_congelado_SIMULADO.json", copia_fz)
    c, t = run("analizar_piloto.py", "--registro", registro_demo, "--demo-simulada", "--demo-raiz", rd, "--salida", bd / "c7", "--modelo-congelado", copia_fz)
    check("un congelamiento de fuera de la carpeta de demostración se rechaza (nunca el real)", c != 0 and "ME NIEGO" in t and not (bd / "c7" / "analisis_piloto_SIMULADO.json").exists(), t[-200:])
    ds.marcar_directorio(c4)
    # tablero de una ejecución de demostración: lee solo los *_SIMULADO de esa carpeta; el mundo real no se evalúa
    rt = W / "raiz_tablero_demo"
    ej = rt / "evidencias" / "simulado_demostracion" / "ej_tablero"
    escribir(rt, "evidencias/simulado_demostracion/ej_tablero/01_ingesta/ingesta_reporte_SIMULADO.txt", ds.MARCA_ESTADO + "\nINGESTA DEL LOTE 1 — REPORTE\n" + LISTO_SI + "\n")
    escribir(rt, "logs/v3_real/ingesta_reporte.txt", "INGESTA DEL LOTE 1 — REPORTE\n" + LISTO_NO + "\n")  # evidencia «real» que no debe mezclarse
    c, t = run("estado_compuertas.py", "--raiz", rt, "--demo-simulada", "--ejecucion", "ej_tablero")
    jt = json.loads((ej / "06_tablero" / "estado_compuertas_SIMULADO.json").read_text(encoding="utf-8")) if c == 0 else {}
    gt = {x["id"]: x for x in jt.get("compuertas", [])}
    check("tablero --demo-simulada: escribe solo en <ejecución>/06_tablero/, con sufijo _SIMULADO y marcador, y G1 sale «Cumplida (Simulado)»",
          c == 0 and gt.get("G1", {}).get("estado") == "Cumplida (Simulado)" and (ej / "06_tablero" / "estado_compuertas_SIMULADO.md").read_text(encoding="utf-8").splitlines()[0] == ds.MARCA_ESTADO
          and "ESTADO" in jt and not (rt / "logs" / "avance").exists(), t[-300:])
    check("en ese tablero el mundo real no se evalúa (la evidencia real de G1 no se mezcla) y las reales siguen en 0",
          jt.get("compuertas_cumplidas_con_datos_reales") == 0 and "TABLERO DE LA DEMOSTRACIÓN SIMULADA" in (ej / "06_tablero" / "estado_compuertas_SIMULADO.md").read_text(encoding="utf-8")
          and all(x["datos"] in ("Simulado", "—") for x in gt.values()), str({k: (v["estado"], v["datos"]) for k, v in gt.items()}))
    c, t = run("estado_compuertas.py", "--raiz", rt, "--demo-simulada", "--ejecucion", "no_existe")
    check("el tablero de demostración con una ejecución inexistente se rechaza", c != 0 and "no existe" in t, t[-200:])
    c, t = run("estado_compuertas.py", "--raiz", rt, "--demo-simulada")
    check("el tablero de demostración exige --ejecucion", c != 0 and "--ejecucion" in t, t[-200:])
    c, t = run("estado_compuertas.py", "--raiz", rt, "--demo-simulada", "--ejecucion", "ej_tablero", "--salida", rt / "logs" / "avance")
    check("el tablero de demostración se niega a escribir en logs/avance/ (solo en la carpeta de la ejecución)", c != 0 and "ME NIEGO" in t, t[-200:])
    # marcado: YAML con comentario, el informe conserva su nombre y es idempotente
    mk = bd / "mk"
    mk.mkdir()
    (mk / "config.yml").write_text("language: es\n", encoding="utf-8")
    (mk / "INFORME_DEMOSTRACION_SIMULADA.md").write_text("# Informe\n", encoding="utf-8")
    ds.marcar_directorio(mk)
    primera = {f.name: f.read_text(encoding="utf-8") for f in mk.iterdir()}
    ds.marcar_directorio(mk)
    segunda = {f.name: f.read_text(encoding="utf-8") for f in mk.iterdir()}
    check("marcar_directorio: el YAML lleva el marcador como comentario, el INFORME_… conserva su nombre y marcar dos veces no cambia nada",
          primera == segunda and primera.get("config_SIMULADO.yml", "").startswith("# " + ds.MARCA_ESTADO) and "INFORME_DEMOSTRACION_SIMULADA.md" in primera
          and primera["INFORME_DEMOSTRACION_SIMULADA.md"].splitlines()[0] == ds.MARCA_ESTADO, str(sorted(primera)))
    c, t = run("demostracion_simulada.py")
    check("demostracion_simulada.py sin --demo-simulada se niega y no crea nada", c != 0 and "ME NIEGO" in t, t[-200:])

    # ------------------------------------------------------------------ integridad
    print("\nIntegridad del repositorio")
    despues = snapshot_repo()
    cambios = sorted(k for k in set(antes) | set(despues) if antes.get(k) != despues.get(k))
    check("corpus/real, logs/v3_real, logs/avance, evidencias/simulado_demostracion, docs/lote_real_1 y docs/piloto no cambiaron durante la prueba", not cambios, f"cambiaron: {cambios}")
    ok = sum(RES)
    print(f"\nRESULTADO (SIMULADA): {ok}/{len(RES)} comprobaciones PASS" + ("" if ok == len(RES) else "  <- HAY FALLAS"))
    print("Recordatorio: datos FALSOS de prueba; nada de esto es un avance real de las compuertas.")
    if not a.conservar:
        shutil.rmtree(W, ignore_errors=True)
    sys.exit(0 if ok == len(RES) else 1)


if __name__ == "__main__":
    main()
