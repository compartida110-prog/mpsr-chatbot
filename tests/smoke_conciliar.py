"""Prueba de humo de scripts/conciliar_seguimiento.py con DATOS FALSOS en una carpeta temporal.

Construye, a partir de la plantilla v3 vacía, seguimientos y respuestas ficticios ("prueba falsa") y comprueba que el
script: concilia un caso consistente; detecta una respuesta ausente que no figura como blanco, un blanco que sí tiene
respuesta, un Transcrito sin respuestas, respuestas de un participante no Transcrito y blancos inválidos; se NIEGA a
trabajar con los archivos simulados; y exporta los participantes SIN ocupación. Los datos no se guardan en el
repositorio, y al final se verifica que ningún archivo de docs/lote_real_1, corpus/real y logs/v3_real cambió.

Uso:
    python tests/smoke_conciliar.py [--conservar]
"""
import argparse
import csv
import datetime
import hashlib
import os
import shutil
import subprocess
import sys
import tempfile
import warnings
from pathlib import Path

warnings.filterwarnings("ignore")
import openpyxl  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
SCRIPT = ROOT / "scripts" / "conciliar_seguimiento.py"
PLANTILLA = ROOT / "docs" / "lote_real_1" / "seguimiento" / "Seguimiento_Lote1_Participantes_PLANTILLA_v3.xlsx"
SIM_V3 = ROOT / "docs" / "lote_real_1" / "seguimiento" / "ejemplos_simulados" / "Seguimiento_Lote1_Participantes_SIMULADO_v3.xlsx"
SIM_V2 = ROOT / "docs" / "lote_real_1" / "seguimiento" / "historico" / "Seguimiento_Lote1_Participantes_SIMULADO_v2.xlsx"
CATALOGO = ROOT / "docs" / "lote_real_1" / "situaciones_lote1_v1.csv"
ENV = {**os.environ, "PYTHONIOENCODING": "utf-8"}
RES = []


def check(nombre, cond, detalle=""):
    RES.append(bool(cond))
    print(f"  [{'PASS' if cond else 'FAIL'}] {nombre}" + (f"  ({detalle})" if detalle and not cond else ""))


def snapshot():
    out = {}
    for d in ("docs/lote_real_1", "corpus/real", "logs/v3_real"):
        for p in sorted((ROOT / d).rglob("*")):
            if p.is_file():
                out[str(p.relative_to(ROOT))] = hashlib.sha256(p.read_bytes()).hexdigest()
    return out


def read_csv(path):
    with open(path, encoding="utf-8") as f:
        return list(csv.DictReader(f))


def write_csv(path, header, rows):
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(header)
        w.writerows(rows)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--conservar", action="store_true")
    a = ap.parse_args()
    W = Path(tempfile.mkdtemp(prefix="conciliar_PRUEBA_"))
    print(f"Carpeta temporal (datos FALSOS, no se commitean): {W}")
    antes = snapshot()
    cat = read_csv(CATALOGO)
    sit_form = {r["scenario_id"]: r["form"] for r in cat}
    por_form = {f: sorted(s for s, g in sit_form.items() if g == f) for f in "ABCDE"}
    forma = lambda code: "ABCDE"[(int(code[1:]) - 1) % 5]
    codigos = [f"P{k:02d}" for k in range(1, 26)]
    transcritos = codigos[:20]
    blancos_base = [("P04", por_form[forma("P04")][0]), ("P07", por_form[forma("P07")][1])]  # P04: fila vacía; P07: sin fila

    def tracker(nombre, blancos=None, estados=None, obs="prueba falsa"):
        wb = openpyxl.load_workbook(PLANTILLA)
        ws = wb["Participantes"]
        for k, c in enumerate(codigos):
            r = 4 + k
            assert ws.cell(r, 1).value == c
            est = (estados or {}).get(c, "Transcrito" if c in transcritos else "Pendiente")
            ws.cell(r, 9).value = est
            if est != "Pendiente":
                ws.cell(r, 5).value = ["18–29", "30–44", "45–59", "60 a más"][k % 4]
                ws.cell(r, 6).value = "PRUEBA ocupación falsa"
                ws.cell(r, 7).value = "Sí"
                ws.cell(r, 8).value = "Sí" if k % 3 else "No"
                ws.cell(r, 10).value = datetime.date(2026, 1, 1)
                ws.cell(r, 14).value = obs
        wb_ = wb["Blancos"]
        for i, (c, s) in enumerate(blancos if blancos is not None else blancos_base):
            wb_.cell(4 + i, 1).value, wb_.cell(4 + i, 2).value = c, s
        p = W / nombre
        wb.save(p)
        return p

    def respuestas(nombre, quitar=(), extra=(), solo_fila_vacia=("P04",)):
        filas = []
        bl = set(blancos_base)
        for c in transcritos:
            for s in por_form[forma(c)]:
                if (c, s) in bl:
                    if c in solo_fila_vacia:
                        filas.append([c, forma(c), s, ""])  # blanco con fila vacía
                    continue  # blanco sin fila
                if (c, s) in set(quitar):
                    continue
                filas.append([c, forma(c), s, f"prueba falsa respuesta {c} {s}"])
        filas += [list(x) for x in extra]
        p = W / nombre
        write_csv(p, ["participant_code", "form", "scenario_id", "text"], filas)
        return p

    def correr(seg, resp, sub, cat_path=CATALOGO, *extra):
        p = subprocess.run([sys.executable, str(SCRIPT), "--seguimiento", str(seg), "--respuestas", str(resp), "--situaciones", str(cat_path),
                            "--out-participantes", str(W / sub / "participantes.csv"), "--log-dir", str(W / sub / "log"), *extra],
                           capture_output=True, text=True, encoding="utf-8", errors="replace", env=ENV, cwd=ROOT)
        return p.returncode, (p.stdout or "") + (p.stderr or "")

    # --------------------------------------------------------------- caso consistente
    print("\nCaso consistente")
    c, t = correr(tracker("seg_ok.xlsx"), respuestas("resp_ok.csv"), "ok")
    check("seguimiento y respuestas consistentes -> código 0", c == 0, t[-400:])
    check("[OK] en ausentes y en blancos con respuesta", "[OK ] Respuesta ausente que NO figura como blanco: 0" in t and "[OK ] Blanco que SÍ tiene respuesta: 0" in t)
    check("un blanco con fila vacía y otro sin fila se aceptan como consistentes", "Transcritos: 20" in t and "blancos registrados: 2" in t)
    check("cobertura real calculada y «Listo para la Parte B» = SÍ (20 transcritos, mínimo 3 por intención)", "¿Listo para la Parte B? SÍ" in t, t[-500:])
    exp = read_csv(W / "ok" / "participantes.csv")
    check("exporta 20 participantes Transcritos con las 5 columnas de ingest_real_lote.py", len(exp) == 20 and list(exp[0]) == ["participant_code", "form", "age_range", "vive_en_juliaca", "tramite_12m"], list(exp[0]) if exp else "")
    check("la exportación NO contiene ocupación, fecha ni observaciones", "ocupaci" not in (W / "ok" / "participantes.csv").read_text(encoding="utf-8").lower() and "prueba" not in (W / "ok" / "participantes.csv").read_text(encoding="utf-8").lower())
    check("el formulario exportado sigue la rotación A-E por código", all(r["form"] == forma(r["participant_code"]) for r in exp))
    check("no exporta a los participantes Pendientes", not any(r["participant_code"] in codigos[20:] for r in exp))

    # --------------------------------------------------------------- discrepancias
    print("\nDiscrepancias que debe detectar")
    s_aus = por_form[forma("P02")][2]
    c, t = correr(tracker("seg_a.xlsx"), respuestas("resp_aus.csv", quitar=[("P02", s_aus)]), "aus")
    check("respuesta ausente que NO figura como blanco -> código 1 y nombra (P02, situación)", c == 1 and f"P02 / {s_aus}: no hay respuesta y NO figura en Blancos" in t, t[-500:])
    s_bc = por_form[forma("P03")][1]
    c, t = correr(tracker("seg_b.xlsx", blancos=blancos_base + [("P03", s_bc)]), respuestas("resp_bc.csv"), "bc")
    check("blanco que SÍ tiene respuesta -> código 1 y nombra (P03, situación)", c == 1 and f"P03 / {s_bc}: figura como blanco pero SÍ tiene respuesta" in t, t[-500:])
    resp_sin5 = [r for r in read_csv(W / "resp_ok.csv") if r["participant_code"] != "P05"]
    write_csv(W / "resp_sin5.csv", ["participant_code", "form", "scenario_id", "text"], [list(r.values()) for r in resp_sin5])
    c, t = correr(tracker("seg_c.xlsx"), W / "resp_sin5.csv", "sin5")
    check("participante Transcrito sin respuestas -> código 1", c == 1 and "P05 (formulario" in t and "sin ninguna fila" in t, t[-500:])
    c, t = correr(tracker("seg_d.xlsx"), respuestas("resp_pend.csv", extra=[("P22", forma("P22"), por_form[forma("P22")][0], "prueba falsa respuesta")]), "pend")
    check("respuestas de un participante que NO está Transcrito -> código 1", c == 1 and "P22" in t and "el seguimiento dice 'Pendiente'" in t, t[-500:])
    c, t = correr(tracker("seg_e.xlsx"), respuestas("resp_desc.csv", extra=[("P88", "A", por_form["A"][0], "prueba falsa")]), "desc")
    check("respuestas de un código que no está en el seguimiento -> código 1", c == 1 and "P88" in t and "no está en el seguimiento" in t, t[-500:])
    otra = [s for s in sit_form if sit_form[s] != forma("P01")][0]
    c, t = correr(tracker("seg_f.xlsx", blancos=blancos_base + [("P01", otra)]), respuestas("resp_inv.csv"), "inv")
    check("blanco de una situación de otro formulario -> código 1 (Blancos inválidos)", c == 1 and "Blancos inválidos: 1" in t, t[-500:])
    c, t = correr(tracker("seg_g.xlsx", estados={"P01": "Terminado"}), respuestas("resp_est.csv"), "est")
    check("estado fuera de la lista -> código 1", c == 1 and "Estados inválidos: 1" in t, t[-400:])

    # --------------------------------------------------------------- simulados
    print("\nSeguimientos simulados: se niega")
    for nombre, ruta in (("SIMULADO_v3 (ejemplo del repositorio)", SIM_V3), ("SIMULADO_v2 (histórico del repositorio)", SIM_V2)):
        sub = "sim_" + nombre[:11].replace(" ", "")
        c, t = correr(ruta, W / "resp_ok.csv", sub)
        check(f"{nombre}: se niega (código 2) y no genera salidas", c == 2 and "ME NIEGO" in t and not (W / sub / "participantes.csv").exists() and not (W / sub / "log").exists(), t[-300:])
    c, t = correr(tracker("seg_h.xlsx", obs="SIMULADO · prueba"), W / "resp_ok.csv", "sim_obs")
    check("un seguimiento cuyas observaciones dicen «SIMULADO» también se rechaza", c == 2 and "ME NIEGO" in t and not (W / "sim_obs" / "participantes.csv").exists(), t[-300:])

    c, t = correr(tracker("seg_i.xlsx", obs="Dato sintético de prueba"), W / "resp_ok.csv", "sim_sint")
    check("título limpio pero Observaciones «Dato sintético de prueba»: se rechaza (código 2) y no genera salidas",
          c == 2 and "ME NIEGO" in t and not (W / "sim_sint" / "participantes.csv").exists() and not (W / "sim_sint" / "log").exists(), t[-300:])
    c, t = correr(tracker("seg_j.xlsx", obs="Es un DEMO"), W / "resp_ok.csv", "sim_demo")
    check("Observaciones con «DEMO» también se rechaza", c == 2 and "ME NIEGO" in t, t[-300:])
    c, t = correr(tracker("seg_k.xlsx", obs="Dato sintetico de prueba"), W / "resp_ok.csv", "sim_sin_tilde")
    check("«SINTETICO» sin tilde también se rechaza", c == 2 and "ME NIEGO" in t, t[-300:])
    c, t = correr(tracker("seg_l.xlsx", obs="Dato sintético de prueba"), respuestas("resp_ok2.csv"), "sim_perm", CATALOGO, "--permitir-simulado")
    check("con --permitir-simulado se acepta, avisa y escribe solo en la carpeta temporal indicada",
          c in (0, 1) and "--permitir-simulado" in t and (W / "sim_perm" / "participantes.csv").exists(), t[-300:])
    c, t = correr(tracker("seg_m.xlsx", obs="Dato sintético de prueba"), respuestas("resp_ok3.csv"), "sim_perm2", CATALOGO, "--permitir-simulado")
    check("--permitir-simulado con el seguimiento simulado no toca el repositorio (la integridad se verifica abajo)", c in (0, 1), t[-200:])
    c, t = correr(tracker("seg_n.xlsx", obs="Cliente con demora en el trámite"), respuestas("resp_ok4.csv"), "sin_marca")
    check("una observación corriente («demora») NO hace rechazar el seguimiento", c == 0 and "ME NIEGO" not in t, t[-300:])

    # --------------------------------------------------------------- plantilla vacía
    print("\nPlantilla vacía")
    write_csv(W / "resp_vacia.csv", ["participant_code", "form", "scenario_id", "text"], [])
    c, t = correr(PLANTILLA, W / "resp_vacia.csv", "plantilla")
    check("la plantilla vacía no es simulada: código 0, 0 transcritos, «Listo» = NO", c == 0 and "Transcritos: 0" in t and "¿Listo para la Parte B? NO" in t, t[-400:])
    check("la plantilla vacía exporta 0 participantes (solo encabezado)", len(read_csv(W / "plantilla" / "participantes.csv")) == 0)
    check("la fila de ejemplo de la plantilla no se cuenta como participante", "Participantes en el seguimiento: 25" in t, t[:300])

    # --------------------------------------------------------------- integridad
    print("\nIntegridad del repositorio")
    despues = snapshot()
    cambios = sorted(k for k in set(antes) | set(despues) if antes.get(k) != despues.get(k))
    check("ningún archivo de docs/lote_real_1, corpus/real ni logs/v3_real cambió durante la prueba", not cambios, f"cambiaron: {cambios}")
    ok = sum(RES)
    print(f"\nRESULTADO: {ok}/{len(RES)} comprobaciones PASS" + ("" if ok == len(RES) else "  <- HAY FALLAS"))
    print("Recordatorio: son datos FALSOS de prueba; nada de esto es un resultado.")
    if not a.conservar:
        shutil.rmtree(W, ignore_errors=True)
    sys.exit(0 if ok == len(RES) else 1)


if __name__ == "__main__":
    main()
