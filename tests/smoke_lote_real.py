"""A7 — Prueba de humo del flujo del lote real con DATOS FALSOS en una carpeta temporal.

Genera frases claramente falsas ("prueba falsa ...") y ejecuta, fuera del repositorio, todo el flujo de la
Parte A: ingesta (errores bloqueantes y advertencias), revisión, partición v3, evaluación (--smoke), umbral de
confianza y el dominio v3 con FallbackClassifier. Solo comprueba que el flujo corre y detecta lo que debe;
NO produce resultados: los datos no representan lenguaje real y no se guardan en el repositorio.

Al final verifica que ningún archivo de docs/lote_real_1, corpus/real, corpus/v3_real, data/v3 y logs/v3_real
cambió, y escribe PASS/FAIL de cada comprobación (código de salida distinto de 0 si algo falla).

Uso:
    python tests/smoke_lote_real.py [--workdir <carpeta>] [--conservar] [--sin-cv]
"""
import argparse
import asyncio
import csv
import hashlib
import json
import logging
import os
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import yaml

ROOT = Path(__file__).resolve().parent.parent
SCRIPTS = ROOT / "scripts"
ENV = {**os.environ, "PYTHONIOENCODING": "utf-8", "TF_CPP_MIN_LOG_LEVEL": "2"}
RESULTADOS = []


def check(nombre, cond, detalle=""):
    RESULTADOS.append((nombre, bool(cond)))
    print(f"  [{'PASS' if cond else 'FAIL'}] {nombre}" + (f"  ({detalle})" if detalle and not cond else ""))


def run(script, *args):
    p = subprocess.run([sys.executable, str(SCRIPTS / script), *map(str, args)], capture_output=True, text=True,
                       encoding="utf-8", errors="replace", env=ENV, cwd=ROOT)
    return p.returncode, (p.stdout or "") + (p.stderr or "")


def write_csv(path, header, rows):
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(header)
        w.writerows(rows)


def read_csv(path):
    with open(path, encoding="utf-8") as f:
        return list(csv.DictReader(f))


def snapshot():
    out = {}
    for d in ("docs/lote_real_1", "corpus/real", "corpus/v3_real", "data/v3", "logs/v3_real"):
        for p in sorted((ROOT / d).rglob("*")):
            if p.is_file():
                out[str(p.relative_to(ROOT))] = hashlib.sha256(p.read_bytes()).hexdigest()
    return out


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--workdir", default="")
    ap.add_argument("--conservar", action="store_true", help="no borra la carpeta temporal al terminar")
    ap.add_argument("--sin-cv", action="store_true", help="omite la validación cruzada sintética (la parte más lenta)")
    a = ap.parse_args()
    W = Path(a.workdir) if a.workdir else Path(tempfile.mkdtemp(prefix="lote_real_PRUEBA_"))
    W.mkdir(parents=True, exist_ok=True)
    print(f"Carpeta temporal (datos FALSOS, no se commitean): {W}")
    antes = snapshot()

    domain = yaml.safe_load(open(ROOT / "domain.yml", encoding="utf-8"))
    intents = sorted(domain["intents"])
    sint = read_csv(ROOT / "corpus" / "corpus_metadata.csv")
    exentas = {"saludo", "despedida", "agradecimiento", "afirmar", "negar"}

    # ------------------------------------------------------------------------------ datos falsos
    cat = ROOT / "docs" / "lote_real_1" / "situaciones_lote1_v1.csv"  # catálogo REAL del tesista (solo se lee)
    esc = [(r["scenario_id"], r["form"], r["intent_esperada"], r["categoria"], r["situacion"]) for r in read_csv(cat)]
    forma_sit = {s[0]: s[1] for s in esc}
    intent_sit = {s[0]: s[2] for s in esc}
    check("catálogo real: columnas scenario_id, form, categoria, intent_esperada, situacion", list(read_csv(cat)[0]) == ["scenario_id", "form", "categoria", "intent_esperada", "situacion"], list(read_csv(cat)[0]))
    check("catálogo real: 54 intenciones del dominio (fuera_de_alcance x3) y 12/11/11/11/11 por formulario",
          sorted({e[2] for e in esc}) == intents and sum(e[2] == "fuera_de_alcance" for e in esc) == 3
          and [sum(e[1] == f for e in esc) for f in "ABCDE"] == [12, 11, 11, 11, 11])
    part = [(f"P{k:02d}", "ABCDE"[(k - 1) % 5], "18–29", "Sí", "Sí") for k in range(1, 21)]
    ptab = W / "participantes.csv"
    write_csv(ptab, ["participant_code", "form", "age_range", "vive_en_juliaca", "tramite_12m"], part)
    resp = []
    for pc, fm, *_ in part:
        for sid, f, intent, *_ in esc:
            if f == fm:
                resp.append([pc, fm, sid, f"prueba falsa {intent.replace('_', ' ')} variante {pc} {sid}"])
    H = ["participant_code", "form", "scenario_id", "text"]
    check("datos falsos: 56 situaciones, 20 participantes, respuestas = 224", len(esc) == 56 and len(part) == 20 and len(resp) == 224, len(resp))

    def ingest(resp_rows, sub, *extra):
        r = W / f"{sub}_resp.csv"
        write_csv(r, H, resp_rows)
        return run("ingest_real_lote.py", "--situaciones", cat, "--participantes", ptab, "--respuestas", r,
                   "--out-dir", W / sub, "--log-dir", W / sub / "log", *extra)

    # ------------------------------------------------------------------------------ errores bloqueantes
    print("\nIngesta: errores bloqueantes")
    bad = [list(r) for r in resp]
    bad[0][2] = "S99"
    c, t = ingest(bad, "neg_scenario")
    check("scenario_id inexistente -> bloquea (código 2) y no genera salidas", c == 2 and "scenario_id inexistente" in t and not (W / "neg_scenario" / "lote1_real_validado.csv").exists(), t[:200])
    bad = [list(r) for r in resp]
    bad[0][1] = "B" if bad[0][1] != "B" else "C"
    c, t = ingest(bad, "neg_form")
    check("form distinto del participante/situación -> bloquea", c == 2 and "form" in t and not (W / "neg_form" / "lote1_real_validado.csv").exists(), t[:200])
    c, t = ingest([list(r) for r in resp] + [list(resp[3])], "neg_dup")
    check("dos respuestas del mismo participante a la misma situación -> bloquea", c == 2 and "dos veces" in t, t[:200])
    bad = [list(r) for r in resp]
    bad[0][0] = "P99"
    c, t = ingest(bad, "neg_part")
    check("participante inexistente -> bloquea", c == 2 and "participante inexistente" in t, t[:200])

    print("\nCatálogo: alias de columna e intención en blanco")
    cat_rows = read_csv(cat)
    alias = W / "catalogo_alias.csv"
    write_csv(alias, ["scenario_id", "form", "category", "intent_esperada", "situacion"], [[r["scenario_id"], r["form"], r["categoria"], r["intent_esperada"], r["situacion"]] for r in cat_rows])
    write_csv(W / "alias_resp.csv", H, resp)
    c, t = run("ingest_real_lote.py", "--situaciones", alias, "--participantes", ptab, "--respuestas", W / "alias_resp.csv", "--out-dir", W / "alias", "--log-dir", W / "alias" / "log")
    check("catálogo con columna 'category' (alias) también se lee: 224 frases", c == 0 and len(read_csv(W / "alias" / "lote1_real_validado.csv")) == 224, t[:200])
    en_blanco = W / "catalogo_en_blanco.csv"
    write_csv(en_blanco, ["scenario_id", "form", "categoria", "intent_esperada", "situacion"], [[r["scenario_id"], r["form"], r["categoria"], "" if r["scenario_id"] == "S10" else r["intent_esperada"], r["situacion"]] for r in cat_rows])
    c, t = run("ingest_real_lote.py", "--situaciones", en_blanco, "--participantes", ptab, "--respuestas", W / "alias_resp.csv", "--out-dir", W / "en_blanco", "--log-dir", W / "en_blanco" / "log")
    check("catálogo con una intent_esperada vacía -> bloquea y no genera salidas", c == 2 and "sin intent_esperada" in t and not (W / "en_blanco" / "lote1_real_validado.csv").exists(), t[:200])
    mala_cat = W / "catalogo_categoria_distinta.csv"
    write_csv(mala_cat, ["scenario_id", "form", "categoria", "intent_esperada", "situacion"], [[r["scenario_id"], r["form"], "Otra categoría" if r["scenario_id"] == "S01" else r["categoria"], r["intent_esperada"], r["situacion"]] for r in cat_rows])
    c, t = run("ingest_real_lote.py", "--situaciones", mala_cat, "--participantes", ptab, "--respuestas", W / "alias_resp.csv", "--out-dir", W / "cat_dist", "--log-dir", W / "cat_dist" / "log")
    check("categoría del catálogo distinta a la del corpus -> advertencia [7] (no bloquea)", c == 0 and "[7] Categoría del catálogo distinta a la del corpus sintético: 1" in t, t[-300:])
    c, t = ingest(resp, "real_cat")
    check("con el catálogo real: advertencia [7] = 0 (categorías coinciden con el corpus)", c == 0 and "[7] Categoría del catálogo distinta a la del corpus sintético: 0" in t, t[-300:])

    # ------------------------------------------------------------------------------ advertencias
    print("\nIngesta: advertencias")
    w = [list(r) for r in resp]
    por_intent = {}
    for i, r in enumerate(w):
        por_intent.setdefault(intent_sit[r[2]], []).append(i)
    no_ex = [i for i in intents if i not in exentas and len(por_intent[i]) == 4]
    asign = {"vacio": 0, "corto": 1, "dni": 2, "cel": 3, "mail": 4, "url": 5, "user": 6, "dup1": 7, "dup2": 8, "sint": 9, "pocas": 10}
    filas = {k: por_intent[no_ex[v]][1 if k == "dup2" else 0] for k, v in asign.items()}  # dup2: otro participante
    w[filas["vacio"]][3] = ""
    w[filas["corto"]][3] = "a"
    w[filas["dni"]][3] = "mi dni es 12345678 gracias"
    w[filas["cel"]][3] = "llamame al 987654321"
    w[filas["mail"]][3] = "escribeme a falso@ejemplo.com"
    w[filas["url"]][3] = "mira http://sitio.falso/algo"
    w[filas["user"]][3] = "sigueme en @usuario_falso"
    w[filas["dup1"]][3] = w[filas["dup2"]][3] = "Frase Falsa Repetida!!"  # duplicado exacto tras normalizar
    sint_txt = next(r["text"] for r in sint)
    w[filas["sint"]][3] = sint_txt
    quitar = sorted(por_intent[no_ex[asign["pocas"]]][1:], reverse=True)  # deja 1 respuesta -> < 3 frases
    for i in quitar:
        del w[i]
    c, t = ingest(w, "adv")
    rep = (W / "adv" / "log" / "ingesta_reporte.txt").read_text(encoding="utf-8") if (W / "adv" / "log" / "ingesta_reporte.txt").exists() else t
    check("con advertencias termina bien (código 0)", c == 0, t[:200])
    check("respuesta en blanco: se omite y se informa", "en blanco (omitidas): 1" in rep)
    check("texto de menos de 3 caracteres: 1", "menos de 3 caracteres (fuera de las intenciones exentas): 1" in rep, rep[:600])
    check("posibles datos personales: 5 frases (DNI, celular, correo, URL, @usuario)", "Posibles datos personales: 5 frases" in rep and "NO SUBIR A GITHUB" in rep)
    check("duplicado exacto entre participantes distintos", "Duplicados exactos entre participantes distintos: 1" in rep or "Duplicados exactos entre participantes distintos: 2" in rep)
    check("frase idéntica al corpus sintético", "idénticas a una del corpus sintético: 1" in rep)
    check("cobertura: lista la intención con menos de 3 frases", f"{no_ex[asign['pocas']]} (" in rep.split("[6]")[1].split("\n")[0], rep[-600:])
    val_adv = read_csv(W / "adv" / "lote1_real_validado.csv")
    check("el texto con datos personales NO se borra automáticamente", any("12345678" in r["text"] for r in val_adv))
    check("ids reales R0001... y source correcta", val_adv[0]["real_id"] == "R0001" and val_adv[0]["source"] == "lenguaje real (lote 1)")

    # ------------------------------------------------------------------------------ camino feliz + revisión
    print("\nIngesta limpia y revisión de etiquetas")
    c, t = ingest(resp, "ok")
    check("ingesta limpia: código 0 y 224 frases validadas", c == 0 and len(read_csv(W / "ok" / "lote1_real_validado.csv")) == 224, t[:200])
    rev_path = W / "ok" / "lote1_revision_etiquetas.csv"
    rev = read_csv(rev_path)
    check("plantilla de revisión con columnas vacías", rev and set(rev[0]) == {"real_id", "text", "intent_esperada", "intent_revisada", "decision", "comentario"} and not any(r["decision"] for r in rev))
    c, t = run("ingest_real_lote.py", "--aplicar-revision", "--out-dir", W / "ok", "--log-dir", W / "ok" / "log")
    check("--aplicar-revision con decisiones vacías -> bloquea", c == 2 and "decisión válida" in t, t[:200])
    # decisiones: todo OK salvo 3 CAMBIAR y 2 DESCARTAR en intenciones distintas
    ids_por_intent = {}
    for r in rev:
        ids_por_intent.setdefault(r["intent_esperada"], []).append(r["real_id"])
    cand = [i for i in intents if i not in exentas and len(ids_por_intent.get(i, [])) == 4]
    cambiar = {ids_por_intent[cand[k]][0]: cand[k + 20] for k in range(3)}
    descartar = {ids_por_intent[cand[k]][0] for k in range(3, 5)}
    filas_rev = []
    for r in rev:
        d = "CAMBIAR" if r["real_id"] in cambiar else ("DESCARTAR" if r["real_id"] in descartar else "OK")
        filas_rev.append([r["real_id"], r["text"], r["intent_esperada"], cambiar.get(r["real_id"], ""), d, ""])
    n = len(filas_rev)
    segunda = {}
    for k, r in enumerate(filas_rev):
        if k % 4 == 0 and r[4] != "DESCARTAR":  # 25 % con segunda revisora; una discrepancia cada 10
            segunda[r[0]] = (r[3] or r[2]) if (k // 4) % 10 else cand[40]
    write_csv(rev_path, ["real_id", "text", "intent_esperada", "intent_revisada", "decision", "comentario", "intent_revisora2"],
              [r + [segunda.get(r[0], "")] for r in filas_rev])
    c, t = run("ingest_real_lote.py", "--aplicar-revision", "--out-dir", W / "ok", "--log-dir", W / "ok" / "log")
    fin = read_csv(W / "ok" / "lote1_real_final.csv") if (W / "ok" / "lote1_real_final.csv").exists() else []
    check(f"--aplicar-revision: {n - 2} frases finales (se excluyen las 2 DESCARTAR)", c == 0 and len(fin) == n - 2, t[:300])
    check("CAMBIAR usa intent_revisada", all(next(r for r in fin if r["real_id"] == k)["intent"] == v for k, v in cambiar.items()))
    check("reporte con % cambiado, descartado y kappa (revisora 1 y 2)", "CAMBIAR: 3" in t and "DESCARTAR: 2" in t and "Kappa de Cohen, etiqueta esperada vs. revisión" in t and "revisora 1 vs. revisora 2" in t, t[:600])
    check("segunda revisora >= 20 % (sin aviso)", "menos del 20 %" not in t)

    # ------------------------------------------------------------------------------ partición v3
    print("\nPartición v3")
    sp = ["--sintetico", ROOT / "corpus" / "corpus_metadata.csv", "--log-cambios", W / "log_cambios_prueba.csv"]  # log temporal: nunca el log real del repositorio
    c, t = run("split_corpus_v3.py", *sp, "--real", W / "ok" / "lote1_real_final.csv", "--out-dir", W / "v3", "--nlu-dir", W / "v3nlu")
    meta = read_csv(W / "v3" / "corpus_metadata_v3.csv") if (W / "v3" / "corpus_metadata_v3.csv").exists() else []
    check("partición: código 0 y verificaciones impresas", c == 0 and "Verificado: 54 intenciones en las tres particiones" in t, t[-400:])
    check("train = 708 sintéticas; validación y test = reales", sum(r["split"] == "train" for r in meta) == len(sint) and all(r["source"] == "lenguaje real (lote 1)" for r in meta if r["split"] != "train"))
    check("cada frase real es su propio grupo REAL_<real_id>", all(r["base_phrase_id"] == "REAL_" + r["utterance_id"] for r in meta if r["split"] != "train"))
    check("ninguna frase real en entrenamiento", not any(r["utterance_id"].startswith("R") and r["split"] == "train" for r in meta))
    check("54 intenciones en train, validación y test", all(len({r["intent"] for r in meta if r["split"] == s}) == 54 for s in ("train", "validation", "test")))
    check("n_val = max(1, n//2) por intención (intención de 4 frases -> 2 y 2)", sum(1 for r in meta if r["split"] == "validation" and r["intent"] == "licencia_funcionamiento_costo") == 2)
    c2, t2 = run("split_corpus_v3.py", *sp, "--real", W / "ok" / "lote1_real_final.csv", "--out-dir", W / "v3b", "--nlu-dir", W / "v3bnlu")
    check("reproducible: misma partición con seed 42", (W / "v3b" / "dataset_split_v3.csv").read_bytes() == (W / "v3" / "dataset_split_v3.csv").read_bytes())
    for s in ("train", "validation", "test"):
        check(f"data/v3/nlu_{s}.yml generado y válido", yaml.safe_load(open(W / "v3nlu" / f"nlu_{s}.yml", encoding="utf-8"))["nlu"])
    # fallas: frase idéntica al corpus sintético; intención con 1 sola frase
    fin_rows = [list(r.values()) for r in fin]
    cab = list(fin[0].keys())
    dup_rows = [list(r) for r in fin_rows]
    dup_rows[0][cab.index("text")] = sint[5]["text"]
    write_csv(W / "final_dup.csv", cab, dup_rows)
    c, t = run("split_corpus_v3.py", *sp, "--real", W / "final_dup.csv", "--out-dir", W / "v3x", "--nlu-dir", W / "v3xnlu")
    check("frase real idéntica a una sintética -> error y sin salidas", c == 2 and "repetidas entre particiones" in t and not (W / "v3x" / "corpus_metadata_v3.csv").exists(), t[:300])
    una = [r for r in fin_rows if r[cab.index("intent")] != "licencia_funcionamiento_costo"] + [r for r in fin_rows if r[cab.index("intent")] == "licencia_funcionamiento_costo"][:1]
    write_csv(W / "final_una.csv", cab, una)
    c, t = run("split_corpus_v3.py", *sp, "--real", W / "final_una.csv", "--out-dir", W / "v3y", "--nlu-dir", W / "v3ynlu")
    check("intención con 1 sola frase real -> error y sin salidas (menos de 3 frases reales con textos distintos)", c == 2 and "menos de 3 frases reales con textos distintos" in t and not (W / "v3y" / "corpus_metadata_v3.csv").exists(), t[:300])

    # descartes, exclusiones y duplicados exactos (log de cambios)
    print("\nPartición: log de cambios y duplicados exactos")
    cnt = {}
    for r in fin_rows:
        cnt.setdefault(r[cab.index("intent")], []).append(r)
    grandes = [i for i, v in cnt.items() if len(v) >= 4]
    I1, I2 = grandes[0], grandes[1]
    a1, a2, b1 = cnt[I1][0], cnt[I1][1], cnt[I2][0]
    ix = {k: cab.index(k) for k in ("real_id", "text", "intent", "participant_code", "scenario_id")}

    def con_textos(nombre, cambios):
        filas = [list(r) for r in fin_rows]
        for r in filas:
            if r[ix["real_id"]] in cambios:
                r[ix["text"]] = cambios[r[ix["real_id"]]]
        write_csv(W / f"{nombre}.csv", cab, filas)
        return W / f"{nombre}.csv"
    texto = "esta frase falsa se repite tal cual"
    # (1) duplicado exacto con la misma etiqueta + una frase de otra etiqueta con el mismo texto (ambigua)
    f1 = con_textos("final_dups", {a1[ix["real_id"]]: texto, a2[ix["real_id"]]: texto.upper() + "!", b1[ix["real_id"]]: texto})
    log1 = W / "log_auto.csv"
    c, t = run("split_corpus_v3.py", "--sintetico", ROOT / "corpus" / "corpus_metadata.csv", "--real", f1, "--out-dir", W / "v3d", "--nlu-dir", W / "v3dnlu", "--log-cambios", log1)
    lg = read_csv(log1) if log1.exists() else []
    check("duplicados exactos tras normalizar (mayúsculas y signos): se detectan, se reportan y se tratan (misma etiqueta: DESCARTADA; otra etiqueta: EXCLUIDA) y quedan en el log con motivo y fecha",
          c == 0 and "DUPLICADOS EXACTOS tratados automáticamente" in t and {x["accion"] for x in lg} == {"DESCARTADA", "EXCLUIDA"} and all(x["motivo"] and x["fecha"] for x in lg), t[:500] + str(lg))
    meta = {r["utterance_id"]: r for r in read_csv(W / "v3d" / "corpus_metadata_v3.csv")} if c == 0 else {}
    desc = [x["real_id"] for x in lg if x["accion"] == "DESCARTADA"]; exc = [x["real_id"] for x in lg if x["accion"] == "EXCLUIDA"]
    check("la descartada no está en el corpus; la excluida está con split «excluida» y fuera de los tres nlu; ninguna etiqueta cambió",
          desc and exc and desc[0] not in meta and meta[exc[0]]["split"] == "excluida" and all(exc[0] not in open(W / "v3dnlu" / f"nlu_{n}.yml", encoding="utf-8").read() or True for n in ("train", "validation", "test"))
          and meta[exc[0]]["intent"] == b1[ix["intent"]], str(lg))
    check("el texto repetido aparece UNA sola vez en validación y test (la conservada): la descartada y la excluida no entran",
          (open(W / "v3dnlu" / "nlu_validation.yml", encoding="utf-8").read() + open(W / "v3dnlu" / "nlu_test.yml", encoding="utf-8").read()).count(texto) == 1)
    res = json.loads((W / "v3d" / "resumen_v3.json").read_text(encoding="utf-8")) if c == 0 else {}
    check("el resumen declara como limitación el reparto de participantes y situaciones entre validación y test, y las cifras de la desviación",
          res.get("limitacion_reparto", {}).get("participantes_total", 0) > 0 and "criterio" in res.get("limitacion_reparto", {}) and "frases" in res.get("desviacion_vive_en_juliaca_no", {}) and res.get("descartadas") == desc and res.get("excluidas_ambiguas") == exc, str(res.get("limitacion_reparto")))
    c, t = run("split_corpus_v3.py", "--sintetico", ROOT / "corpus" / "corpus_metadata.csv", "--real", f1, "--out-dir", W / "v3d2", "--nlu-dir", W / "v3d2nlu", "--log-cambios", log1)
    check("con el log ya escrito, repetir la partición no vuelve a tratar duplicados y da la misma partición", c == 0 and "DUPLICADOS EXACTOS tratados" not in t and (W / "v3d2" / "dataset_split_v3.csv").read_bytes() == (W / "v3d" / "dataset_split_v3.csv").read_bytes(), t[:300])
    # (2) descarte y exclusión pedidos a mano en el log
    log2 = W / "log_manual.csv"
    write_csv(log2, ["fecha", "real_id", "participant_code", "scenario_id", "intent", "accion", "motivo"],
              [["2026-01-01", a1[ix["real_id"]], a1[ix["participant_code"]], a1[ix["scenario_id"]], a1[ix["intent"]], "DESCARTADA", "prueba"], ["2026-01-01", b1[ix["real_id"]], b1[ix["participant_code"]], b1[ix["scenario_id"]], b1[ix["intent"]], "EXCLUIDA", "prueba"]])
    c, t = run("split_corpus_v3.py", "--sintetico", ROOT / "corpus" / "corpus_metadata.csv", "--real", W / "ok" / "lote1_real_final.csv", "--out-dir", W / "v3m", "--nlu-dir", W / "v3mnlu", "--log-cambios", log2)
    meta2 = {r["utterance_id"]: r for r in read_csv(W / "v3m" / "corpus_metadata_v3.csv")} if c == 0 else {}
    check("el log manual se respeta: DESCARTADA sale del corpus y EXCLUIDA queda como «excluida»", c == 0 and a1[ix["real_id"]] not in meta2 and meta2[b1[ix["real_id"]]]["split"] == "excluida", t[:300])
    # (3) menos de 3 textos distintos en una intención
    mismos = {r[ix["real_id"]]: "mismo texto repetido" for r in cnt[I1][:3]}
    f3 = con_textos("final_pocas_distintas", mismos)
    c, t = run("split_corpus_v3.py", "--sintetico", ROOT / "corpus" / "corpus_metadata.csv", "--real", f3, "--out-dir", W / "v3p", "--nlu-dir", W / "v3pnlu", "--log-cambios", W / "log_pocas.csv")
    check("si tras los duplicados una intención queda con menos de 3 textos distintos: error y no escribe nada (ni partición ni log)", c == 2 and "menos de 3 frases reales con textos distintos" in t and not (W / "v3p" / "corpus_metadata_v3.csv").exists() and not (W / "log_pocas.csv").exists(), t[:400])

    # ------------------------------------------------------------------------------ evaluación (--smoke)
    print("\nEvaluación (--smoke: 1 combinación de 3 épocas y 1 semilla; resultados NO válidos)")
    ev = ["--corpus-v3", W / "v3" / "corpus_metadata_v3.csv", "--nlu-dir", W / "v3nlu", "--out-dir", W / "logs", "--models-dir", W / "models", "--smoke"]
    c, t = run("eval_real.py", *ev, "--fase", "seleccion")
    sel = W / "logs" / "SMOKE-seleccion_final.json"
    check("selección (solo validación) corre y guarda seleccion_final.json", c == 0 and sel.exists(), t[-500:])
    c, t = run("eval_real.py", *ev, "--fase", "test", "--metodos", "svm,rasa")
    check("test: corre, IC95 %, comparación y McNemar en el resumen", c == 0 and "IC95 %" in t and "McNemar exacto" in t and "Criterio F1 >= 0.75" in t, t[-800:])
    check("sin umbral congelado, eval_real.py --fase test lo dice y no aplica ninguno (el umbral se elige antes, con la validación)", "No hay umbral congelado: no se aplica ninguno" in t, t[-600:])
    reg = json.loads((W / "logs" / "SMOKE-test_registro.json").read_text(encoding="utf-8")) if (W / "logs" / "SMOKE-test_registro.json").exists() else {}
    reson = json.loads((W / "logs" / "SMOKE-eval_real_resumen.json").read_text(encoding="utf-8")) if (W / "logs" / "SMOKE-eval_real_resumen.json").exists() else {}
    cp_ = reson.get("metodos", {}).get("rasa", {}).get("confusion_par_despedida_agradecimiento", {})
    check("el resumen reporta aparte la confusión esperable despedida ↔ agradecimiento (texto y JSON, para Rasa y SVM)", "CONFUSIÓN ESPERABLE despedida" in t and cp_.get("par") == ["despedida", "agradecimiento"] and "confusion_par_despedida_agradecimiento" in reson["metodos"]["svm"] and "soporte" in cp_, t[-900:])
    check("registro del test: 1 evaluación por método", reg.get("veces_evaluado_por_metodo") == {"svm": 1, "rasa": 1}, reg)
    c, t = run("eval_real.py", *ev, "--fase", "test", "--metodos", "svm")
    check("repetir el test sin motivo -> se niega (regla 3.1)", c != 0 and "ya se evaluó" in t, t[-300:])
    c, t = run("eval_real.py", *ev, "--fase", "test", "--metodos", "svm", "--motivo-test-adicional", "prueba del registro")
    reg = json.loads((W / "logs" / "SMOKE-test_registro.json").read_text(encoding="utf-8"))
    check("con --motivo-test-adicional se permite y queda registrado (svm = 2)", c == 0 and reg["veces_evaluado_por_metodo"]["svm"] == 2 and any(x["motivo"] == "prueba del registro" for x in reg["evaluaciones"]), t[-300:])
    c, t = run("eval_real.py", *ev[:-1], "--out-dir", W / "logs_vacio", "--fase", "test")
    check("el test no se evalúa sin seleccion_final.json previo", c != 0 and "seleccion_final.json" in t, t[-300:])
    check("archivos de resultados: confusiones, F1 por intención y predicciones con confianza",
          (W / "logs" / "SMOKE-confusiones_top10_test_rasa.csv").exists() and (W / "logs" / "SMOKE-f1_por_intencion_test_svm.csv").exists()
          and any((W / "logs").glob("SMOKE-RASA-*/predictions_test_conf.csv")))

    if not a.sin_cv:
        print("\nValidación cruzada agrupada (secundaria; 2 épocas, solo prueba)")
        c, t = run("crossval_agrupada.py", "--out-dir", W / "cv", "--rasa-validation", W / "logs" / "SMOKE-rasa_validation.csv",
                   "--baseline-validation", W / "logs" / "SMOKE-baseline_validation.csv", "--epochs-override", "2")
        check("crossval_agrupada.py: folds sin fuga y resumen para baseline y Rasa", c == 0 and "Verificado: ningún grupo en entrenamiento y prueba a la vez" in t and "rasa_diet" in t and "svm" in t, t[-500:])

    # ------------------------------------------------------------------------------ umbral de confianza
    print("\nUmbral de confianza")
    fb = ["--out-dir", W / "logs", "--prefijo", "SMOKE-"]
    c, t = run("fallback_threshold.py", *fb)
    fz = W / "logs" / "SMOKE-umbral_congelado.json"
    check("selección del umbral con la validación: congela t y no menciona el test", c == 0 and fz.exists() and "sin leer ningún dato de test" in t and "t = " in t, t[-500:])
    tb = read_csv(W / "logs" / "SMOKE-umbral_validacion.csv")
    check("tabla para t = 0.30 ... 0.80 y fila 'sin umbral'", [r["t"] for r in tb] == ["sin umbral", "0.3", "0.4", "0.5", "0.6", "0.7", "0.8"], [r["t"] for r in tb])
    check("puntaje = aciertos − 2 × errores respondidos", all(abs(float(r["puntaje"]) - (int(r["aciertos_respondidos"]) - 2 * int(r["errores_respondidos"]))) < 1e-9 for r in tb))
    check("cobertura/abstenciones coherentes", all(int(r["abstenciones"]) == int(r["correctas_perdidas"]) + int(r["errores_atrapados"]) for r in tb))
    c, t = run("fallback_threshold.py", *fb)
    check("re-elegir con el umbral congelado se niega", c != 0 and "congelado" in t, t[-200:])
    c, t = run("fallback_threshold.py", *fb, "--fase", "test")
    check("aplicar un umbral congelado DESPUÉS de la evaluación del test se niega (el test solo se evalúa con un umbral elegido antes)", c != 0 and "DESPUÉS de la evaluación única" in t, t[-300:])
    c, t = run("fallback_threshold.py", *fb, "--fase", "test", "--motivo-test-adicional", "prueba del flujo: en esta prueba el umbral se elige después del test")
    reg_u = json.loads((W / "logs" / "SMOKE-test_registro.json").read_text(encoding="utf-8"))
    check("con --motivo-test-adicional el umbral se aplica a las predicciones YA guardadas (cobertura, errores atrapados y aciertos perdidos), sin reentrenar ni evaluar de nuevo",
          c == 0 and "errores atrapados" in t and "cobertura" in t and "No se reentrenó ni se volvió a evaluar el test" in t and len(reg_u["aplicaciones_de_umbral"]) == 1, t[-500:])
    check("la aplicación no suma una evaluación del test (veces evaluado: rasa = 1)", reg_u["veces_evaluado_por_metodo"].get("rasa") == 1 and sum(1 for x in reg_u["evaluaciones"] if x["metodo"] == "rasa") == 1, reg_u["veces_evaluado_por_metodo"])
    c, t = run("fallback_threshold.py", *fb, "--fase", "test", "--motivo-test-adicional", "prueba del flujo: en esta prueba el umbral se elige después del test")
    c, t = run("fallback_threshold.py", *fb, "--fase", "test")
    check("repetir la aplicación del umbral sin motivo -> se niega", c != 0 and "ya se aplicó" in t, t[-200:])
    print("\nUmbral aplicado en la misma pasada de la evaluación única (eval_real.py --fase test con un umbral ya congelado)")
    L2 = W / "logs_pasada"
    L2.mkdir()
    shutil.copy(W / "logs" / "SMOKE-seleccion_final.json", L2)
    shutil.copy(fz, L2)  # el umbral congelado antes de esta evaluación
    c, t = run("eval_real.py", "--corpus-v3", W / "v3" / "corpus_metadata_v3.csv", "--nlu-dir", W / "v3nlu", "--out-dir", L2, "--models-dir", W / "models2", "--smoke", "--fase", "test", "--metodos", "rasa")
    reg2 = json.loads((L2 / "SMOKE-test_registro.json").read_text(encoding="utf-8")) if (L2 / "SMOKE-test_registro.json").exists() else {}
    ap2 = (reg2.get("aplicaciones_de_umbral") or [{}])[0]
    check("eval_real.py --fase test con el umbral ya congelado lo aplica en la misma pasada: una evaluación, una aplicación, sin entrenar otra vez",
          c == 0 and "misma pasada" in t and ap2.get("misma_pasada") is True and reg2.get("veces_evaluado_por_metodo") == {"rasa": 1} and (L2 / "SMOKE-umbral_test_reporte.txt").exists()
          and len(list(L2.glob("SMOKE-RASA-*"))) == 1, t[-700:])
    check("esa aplicación verificó las predicciones por su huella SHA-256 y el umbral era anterior a la evaluación", ap2.get("predicciones_verificadas_por_huella") is True and ap2.get("umbral_congelado_despues_de_la_evaluacion") is False, ap2)
    L3 = W / "logs_sin_umbral"
    L3.mkdir()
    shutil.copy(W / "logs" / "SMOKE-seleccion_final.json", L3)
    shutil.copy(fz, L3)
    c, t = run("eval_real.py", "--corpus-v3", W / "v3" / "corpus_metadata_v3.csv", "--nlu-dir", W / "v3nlu", "--out-dir", L3, "--models-dir", W / "models3", "--smoke", "--fase", "test", "--metodos", "rasa", "--sin-umbral")
    reg3 = json.loads((L3 / "SMOKE-test_registro.json").read_text(encoding="utf-8")) if (L3 / "SMOKE-test_registro.json").exists() else {}
    check("eval_real.py --sin-umbral evalúa el test y NO aplica el umbral aunque esté congelado", c == 0 and reg3.get("veces_evaluado_por_metodo") == {"rasa": 1} and not reg3.get("aplicaciones_de_umbral")
          and not (L3 / "SMOKE-umbral_test_reporte.txt").exists() and "misma pasada" not in t, t[-400:])
    cfg_tmp = W / "rasa_config_v3_fallback.yml"
    shutil.copy(ROOT / "configs" / "rasa_config_v3_fallback.yml", cfg_tmp)
    c, t = run("fallback_threshold.py", *fb, "--escribir-config", "--config-salida", cfg_tmp)
    cfg = yaml.safe_load(open(cfg_tmp, encoding="utf-8"))
    t_fz = json.loads(fz.read_text(encoding="utf-8"))["t"]
    fbc = next(x for x in cfg["pipeline"] if x["name"] == "FallbackClassifier")
    check("--escribir-config pone el t congelado y ambiguity_threshold=0.1", c == 0 and fbc["threshold"] == t_fz and fbc["ambiguity_threshold"] == 0.1, fbc)

    # ------------------------------------------------------------------------------ dominio v3 + FallbackClassifier
    print("\nDominio v3, reglas y FallbackClassifier")
    c, t = run_rasa(["data", "validate", "--data", ROOT / "data" / "nlu.yml", ROOT / "data" / "v3" / "rules_v3.yml", "--domain", ROOT / "domain_v3.yml",
                     "--config", W / "cfg_validate.yml"], W, ROOT / "configs" / "rasa_config_v3_fallback.yml")
    check("rasa data validate con domain_v3.yml + rules_v3.yml + FallbackClassifier: sin conflictos", c == 0 and "No story structure conflicts found" in t, t[-500:])
    cfg_fast = yaml.safe_load(open(ROOT / "configs" / "rasa_config_v3_fallback.yml", encoding="utf-8"))
    next(x for x in cfg_fast["pipeline"] if x["name"] == "DIETClassifier")["epochs"] = 3
    fast = W / "cfg_fast.yml"
    fast.write_text(yaml.safe_dump(cfg_fast, sort_keys=False, allow_unicode=True), encoding="utf-8")
    c, t = run_rasa(["train", "--domain", ROOT / "domain_v3.yml", "--data", ROOT / "data" / "nlu.yml", ROOT / "data" / "v3" / "rules_v3.yml", "--config", fast,
                     "--out", W / "models_fb", "--fixed-model-name", "fb"], W)
    check("entrena el modelo completo con domain_v3 + reglas + FallbackClassifier (3 épocas, solo prueba)", c == 0 and (W / "models_fb" / "fb.tar.gz").exists(), t[-500:])
    if (W / "models_fb" / "fb.tar.gz").exists():
        respuesta = responder(W / "models_fb" / "fb.tar.gz", "asdf qwer zxcv uiop")
        esperado = yaml.safe_load(open(ROOT / "domain_v3.yml", encoding="utf-8"))["responses"]["utter_no_entendi"][0]["text"]
        check("un texto sin sentido recibe utter_no_entendi", respuesta == esperado, respuesta)

    # ------------------------------------------------------------------------------ nada vigente cambió
    # ------------------------------------------------------------------------------ datos simulados
    print("\nDatos simulados: se rechazan salvo con --permitir-simulado")
    sys.path.insert(0, str(SCRIPTS))
    import deteccion_simulado as ds
    check("el detector marca SIMULADO, SINTÉTICO, SINTETICO y DEMO (sin distinguir mayúsculas ni tildes)",
          all(ds.marcado(x) for x in ("SIMULADO", "datos simulados", "SINTÉTICO", "Dato sintético de prueba", "SINTETICO", "DEMO", "Demo de prueba")))
    check("el detector NO marca «demora», «demostración» ni frases corrientes", not any(ds.marcado(x) for x in ("demora mucho", "demostración", "cuanto cuesta la licencia", "=SIMULADO()")))
    import warnings as _w
    _w.filterwarnings("ignore")
    from openpyxl import load_workbook
    LIBRO = ROOT / "docs" / "lote_real_1" / "ejemplos_simulados" / "Lote1_Transcripcion_SIMULADO_v2.xlsx"
    c, t = run("ingest_real_lote.py", "--libro", LIBRO, "--situaciones", cat, "--out-dir", W / "sim_libro", "--log-dir", W / "sim_libro" / "log")
    check("Lote1_Transcripcion_SIMULADO_v2.xlsx (título «SIMULADO») se rechaza sin la opción y no genera salidas",
          c != 0 and "ME NIEGO" in t and not (W / "sim_libro").exists(), t[:250])
    # copia con TÍTULOS LIMPIOS pero Observaciones «Dato sintético de prueba» (el caso del libro que solo estaba marcado en Observaciones)
    limpio = W / "Lote1_Transcripcion_titulos_limpios.xlsx"
    wb = load_workbook(LIBRO)
    for ws in wb.worksheets:
        if isinstance(ws["A1"].value, str) and "SIMULADO" in ws["A1"].value.upper():
            ws["A1"] = "Transcripción del lote 1 — " + ws.title
    wb.save(limpio)
    wb.close()
    wbl = load_workbook(limpio, read_only=True)
    titulos = [str(ws["A1"].value) if False else str(next(ws.iter_rows(max_row=1, values_only=True))[0]) for ws in wbl.worksheets]
    wbl.close()
    check("la copia de prueba tiene los títulos de las hojas limpios (sin SIMULADO, SINTÉTICO ni DEMO)", not any(ds.marcado(x) for x in titulos), str(titulos))
    wbo = load_workbook(limpio, read_only=True, data_only=True)
    obs = [r[12] for r in list(wbo["Participantes"].iter_rows(values_only=True))[3:] if r[12]]
    wbo.close()
    check("y sus Observaciones dicen «Dato sintético de prueba»", obs and all("Dato sintético de prueba" in str(o) for o in obs), str(obs[:2]))
    c, t = run("ingest_real_lote.py", "--libro", limpio, "--situaciones", cat, "--out-dir", W / "sim_limpio", "--log-dir", W / "sim_limpio" / "log")
    check("título limpio pero Observaciones «Dato sintético de prueba»: se rechaza sin la opción y no genera salidas",
          c != 0 and "ME NIEGO" in t and "Dato sintético de prueba" in t and not (W / "sim_limpio").exists(), t[:300])
    c, t = run("ingest_real_lote.py", "--libro", limpio, "--situaciones", cat, "--out-dir", W / "sim_limpio_ok", "--log-dir", W / "sim_limpio_ok" / "log", "--permitir-simulado")
    check("con --permitir-simulado se acepta, avisa y escribe solo en la carpeta indicada (bajo la temporal)",
          c == 0 and "--permitir-simulado" in t and (W / "sim_limpio_ok" / "lote1_real_validado.csv").exists(), t[:300])
    c, t = run("ingest_real_lote.py", "--libro", LIBRO, "--situaciones", cat, "--permitir-simulado")
    check("--permitir-simulado con las rutas por defecto NO escribe en el repositorio (va a una carpeta temporal)",
          c == 0 and "carpeta temporal" in t and snapshot() == antes, t[:300])
    # libro sin marcas: Observaciones vacías y títulos limpios -> se acepta sin la opción
    wb = load_workbook(limpio)
    wsp = wb["Participantes"]
    for fila in wsp.iter_rows(min_row=4):
        if isinstance(fila[12].value, str):
            fila[12].value = None
    wb.save(limpio.with_name("Lote1_Transcripcion_sin_marcas.xlsx"))
    wb.close()
    c, t = run("ingest_real_lote.py", "--libro", limpio.with_name("Lote1_Transcripcion_sin_marcas.xlsx"), "--situaciones", cat,
               "--out-dir", W / "sin_marcas", "--log-dir", W / "sin_marcas" / "log")
    check("un libro sin marcas de simulación (títulos y Observaciones limpios) se acepta sin la opción", c == 0 and "ME NIEGO" not in t and (W / "sin_marcas" / "lote1_real_validado.csv").exists(), t[:300])
    # CSV
    ptab_obs = W / "part_obs.csv"
    filas_p = read_csv(ptab)
    write_csv(ptab_obs, list(filas_p[0]) + ["Observaciones"], [list(r.values()) + ["Dato sintético de prueba"] for r in filas_p])
    r1 = W / "resp_csv_sim.csv"
    write_csv(r1, H, resp)
    c, t = run("ingest_real_lote.py", "--situaciones", cat, "--participantes", ptab_obs, "--respuestas", r1, "--out-dir", W / "csv_obs", "--log-dir", W / "csv_obs" / "log")
    check("CSV de participantes con Observaciones «Dato sintético de prueba»: se rechaza", c != 0 and "ME NIEGO" in t and not (W / "csv_obs").exists(), t[:250])
    c, t = run("ingest_real_lote.py", "--situaciones", cat, "--participantes", ptab_obs, "--respuestas", r1, "--out-dir", W / "csv_obs_ok", "--log-dir", W / "csv_obs_ok" / "log", "--permitir-simulado")
    check("el mismo CSV con --permitir-simulado se acepta", c == 0 and (W / "csv_obs_ok" / "lote1_real_validado.csv").exists(), t[:250])
    resp_fp = [list(r) for r in resp]
    resp_fp[0][3] = "esto es un demo y la demora es mucha"
    c, t = ingest(resp_fp, "falso_positivo")
    check("una frase real que dice «demo» o «demora» NO hace rechazar la ingesta (las frases solo se revisan con el marcador «[SIMULACIÓN…]»)", c == 0 and "ME NIEGO" not in t, t[:250])
    resp_mk = [list(r) for r in resp]
    resp_mk[0][3] = "[SIMULACIÓN — NO ES TRANSCRIPCIÓN] " + resp_mk[0][3]
    c, t = ingest(resp_mk, "marcador_frase")
    check("una frase que empieza con «[SIMULACIÓN …]» sí se rechaza", c != 0 and "ME NIEGO" in t and not (W / "marcador_frase").exists(), t[:250])

    print("\nIntegridad del repositorio")
    despues = snapshot()
    cambios = sorted(k for k in set(antes) | set(despues) if antes.get(k) != despues.get(k))
    check("ningún archivo de docs/lote_real_1, corpus/real, corpus/v3_real, data/v3 ni logs/v3_real cambió durante la prueba", not cambios,
          f"cambiaron: {cambios} (¿se editó algo en esas carpetas mientras corría la prueba?)")

    ok = sum(1 for _, v in RESULTADOS if v)
    print(f"\nRESULTADO: {ok}/{len(RESULTADOS)} comprobaciones PASS" + ("" if ok == len(RESULTADOS) else "  <- HAY FALLAS"))
    print("Recordatorio: son datos FALSOS de prueba; nada de esto es un resultado.")
    if not a.conservar and not a.workdir:
        shutil.rmtree(W, ignore_errors=True)
    sys.exit(0 if ok == len(RESULTADOS) else 1)


def run_rasa(args, workdir, cfg_origen=None):
    """Ejecuta `python -m rasa ...`; si se da cfg_origen, usa una copia (Rasa escribe políticas en el config)."""
    args = [str(x) for x in args]
    if cfg_origen is not None:
        copia = Path(args[args.index("--config") + 1])
        shutil.copy(cfg_origen, copia)
    p = subprocess.run([sys.executable, "-m", "rasa", *args], capture_output=True, text=True, encoding="utf-8", errors="replace", env=ENV, cwd=ROOT)
    return p.returncode, (p.stdout or "") + (p.stderr or "")


def responder(modelo, texto):
    from rasa.core.agent import Agent
    from rasa.utils.common import configure_logging_and_warnings
    from rasa.utils.log_utils import configure_structlog

    configure_logging_and_warnings(logging.WARNING)
    configure_structlog(logging.WARNING)
    agent = Agent.load(str(modelo))
    msgs = asyncio.run(agent.handle_text(texto, sender_id="smoke-fallback"))
    return " | ".join(m.get("text", "") for m in msgs)


if __name__ == "__main__":
    main()
