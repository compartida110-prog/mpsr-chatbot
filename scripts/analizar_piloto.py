"""Análisis del piloto exploratorio (protocolo V1.3: una sesión asistida, n = 60 sesiones elegibles).

Solo LEE el registro (openpyxl, data_only=True, sin guardar nunca el .xlsx). Usa las filas con «Elegible y completa = Sí» y código que empiece por SA
(así quedan fuera EJ01 y las vacías).

  OE3  tiempo pre/post por persona (minutos): Shapiro-Wilk sobre las diferencias -> t pareada o Wilcoxon (unilateral, H1: el post es menor),
       tamaño de efecto, reducción de medias, media de reducciones individuales e IC95 % bootstrap (1000 remuestreos, semilla 42).
  OE4  P10 contra el ítem 9 (mismo procedimiento, H1: el ítem 9 es mayor), media de los ítems 1–8 con su IC y alfa de Cronbach (1–8),
       comparada con Resumen!B24 del registro.
  Prueba final del modelo congelado: exige que el hash del modelo coincida con modelo_congelado.json; predice las 3 consultas de cada sesión,
       las compara con la intención esperada de su tarjeta (catálogo), accuracy y F1 macro con IC por bootstrap por conglomerados de sesión,
       umbral congelado, acuerdo con el juicio del aplicador y de la persona. SE EVALÚA UNA SOLA VEZ (logs/piloto/evaluaciones_prueba_final.log);
       repetirla exige --motivo.

Rechazos: título de hoja con «SIMULADO» (salvo --permitir-simulado, que escribe solo en una carpeta temporal y rotula _SIMULADO), sin filas
elegibles, fórmulas sin valor guardado («abre y guarda el archivo en Excel»), hash distinto del modelo.

SUPUESTO A REVISAR: el registro real no estaba en el repositorio al escribir este script. Los encabezados se reconocen por patrones razonables
(ver COLUMNAS); si el registro usa otros nombres, indícalos con --mapa-columnas mapa.json  ({"min_pre": "Minutos antes", ...}). Al terminar, el
script muestra qué encabezado usó para cada campo para que lo confirmes.

Todo resultado sale rotulado REAL (o SIMULADO) con el n efectivo. Es exploratorio, nunca confirmatorio.

Uso:
    python scripts/analizar_piloto.py --registro docs/piloto/privado/Registro_lleno.xlsx --modelo-congelado logs/v3_real/modelo_congelado.json \\
        --catalogo docs/lote_real_1/situaciones_lote1_v1.csv --salida logs/piloto/
"""
import argparse
import json
import re
import sys
import tempfile
import unicodedata
from datetime import datetime
from pathlib import Path

import numpy as np
import pandas as pd
from scipy import stats

from common import ROOT, file_sha256

ALPHA = 0.05
N_BOOT = 1000
SEMILLA = 42
N_OBJETIVO = 60
FILA_ENCABEZADO = 3
HOJA = "Sesiones"
LIMITACIONES = [
    "La línea base (minutos «antes») es recordada por la persona: sesgo de recuerdo.",
    "Sesgo de novedad: la persona prueba un asistente nuevo con el aplicador presente.",
    "El tiempo pre mide un trámite y el post una consulta al asistente: no son exactamente lo mismo (OE3 habla de consultas).",
    "Muestra de conveniencia y piloto exploratorio: no confirmatorio, no generalizable.",
]

# campo lógico -> patrones (regex sobre el encabezado normalizado: minúsculas, sin tildes, solo letras y números separados por espacio)
COLUMNAS = {
    "codigo": [r"^codigo( de sesion)?$", r"^sesion$", r"^id( de sesion)?$"],
    "elegible": [r"^elegible( y completa)?$"],
    "min_pre": [r"^minutos? pre$", r"^pre minutos?$", r"^minutos? antes$", r"^tiempo pre( min)?$", r"^pre$"],
    "min_post": [r"^minutos? post$", r"^post minutos?$", r"^minutos? despues$", r"^tiempo post( min)?$", r"^post$"],
    "p10": [r"^p ?10$", r"^p10 .*", r"^satisfaccion (previa|pre|actual)$"],
}
for _i in range(1, 10):
    COLUMNAS[f"item{_i}"] = [rf"^(item|i|p) ?{_i}$", rf"^(item|i|p) ?{_i} .*", rf"^p ?{_i}$"]
for _k in range(1, 4):
    COLUMNAS[f"t{_k}_tarjeta"] = [rf"^(tarjeta|escenario|situacion) ?{_k}$", rf"^t ?{_k} (tarjeta|escenario|id)$", rf"^tarjeta {_k} .*"]
    COLUMNAS[f"t{_k}_consulta"] = [rf"^consulta ?{_k}$", rf"^t ?{_k} consulta$", rf"^texto consulta ?{_k}$", rf"^consulta {_k} .*"]
    COLUMNAS[f"t{_k}_correcta"] = [rf"^(correcta|aplicador) ?{_k}$", rf"^t ?{_k} correcta$", rf"^correcta {_k} .*", rf"^t ?{_k} aplicador$"]
    COLUMNAS[f"t{_k}_obtuvo"] = [rf"^obtuvo ?{_k}$", rf"^obtuvo lo que necesitaba ?{_k}$",rf"^t ?{_k} obtuvo$", rf"^obtuvo {_k} .*", rf"^t ?{_k} persona$"]
NO_OBLIGATORIAS = {f"t{k}_{c}" for k in range(1, 4) for c in ("correcta", "obtuvo")}


def norm(t):
    t = unicodedata.normalize("NFD", str(t)).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]+", " ", t).strip()


def si(v):
    return norm(v) in ("si", "s", "1", "true", "yes", "sí")


# --------------------------------------------------------------------------- lectura
def leer_registro(ruta, mapa_usuario=None, permitir_simulado=False):
    from openpyxl import load_workbook
    wb = load_workbook(ruta, read_only=True, data_only=True)  # solo lectura: nunca se guarda
    try:
        sim = [h for h in wb.sheetnames if "simulado" in h.lower()]
        if sim and not permitir_simulado:
            sys.exit(f"ERROR: el registro tiene hojas marcadas SIMULADO ({sim}). No se analiza como piloto real (usa --permitir-simulado solo para pruebas).")
        if HOJA not in wb.sheetnames:
            sys.exit(f"ERROR: no hay hoja «{HOJA}» en el registro (hojas: {wb.sheetnames}).")
        filas = list(wb[HOJA].iter_rows(values_only=True))
        resumen_b24 = None
        if "Resumen" in wb.sheetnames:
            r = list(wb["Resumen"].iter_rows(min_row=24, max_row=24, min_col=2, max_col=2, values_only=True))
            resumen_b24 = r[0][0] if r else None
    finally:
        wb.close()
    if len(filas) < FILA_ENCABEZADO:
        sys.exit("ERROR: la hoja de sesiones no tiene la fila de encabezados (fila 3).")
    enc = [("" if c is None else str(c).strip()) for c in filas[FILA_ENCABEZADO - 1]]
    encn = [norm(c) for c in enc]
    mapa, usados = {}, {}
    for campo, pats in COLUMNAS.items():
        if mapa_usuario and campo in mapa_usuario:
            deseado = norm(mapa_usuario[campo])
            if deseado not in encn:
                sys.exit(f"ERROR: --mapa-columnas pide «{mapa_usuario[campo]}» para «{campo}», pero ese encabezado no existe. Encabezados: {enc}")
            mapa[campo] = encn.index(deseado)
            continue
        for j, h in enumerate(encn):
            if h and any(re.match(p, h) for p in pats):
                mapa[campo] = j
                break
    faltan = [c for c in COLUMNAS if c not in mapa and c not in NO_OBLIGATORIAS]
    if faltan:
        sys.exit("ERROR: no reconozco estas columnas del registro: " + ", ".join(faltan) +
                 f"\n  Encabezados de la fila {FILA_ENCABEZADO}: {enc}\n  Indica los nombres reales con --mapa-columnas mapa.json "
                 '(por ejemplo {"min_pre": "Minutos antes"}).')
    datos = []
    for fila in filas[FILA_ENCABEZADO:]:
        cod = fila[mapa["codigo"]] if mapa["codigo"] < len(fila) else None
        if cod is None or not str(cod).strip():
            continue
        datos.append({c: (fila[j] if j < len(fila) else None) for c, j in mapa.items()})
    df = pd.DataFrame(datos)
    if df.empty:
        sys.exit("ERROR: el registro no tiene filas con código de sesión.")
    df["codigo"] = df["codigo"].astype(str).str.strip()
    cand = df[df["codigo"].str.upper().str.startswith("SA")].copy()
    if cand["elegible"].isna().all():
        sys.exit("ERROR: «Elegible y completa» no tiene valores guardados (las fórmulas no se calcularon): abre y guarda el archivo en Excel.")
    el = cand[cand["elegible"].map(si)].copy()
    if el.empty:
        sys.exit("ERROR: no hay filas elegibles (Elegible y completa = Sí con código SA…).")
    cabeceras = {c: enc[j] for c, j in mapa.items()}
    return el.reset_index(drop=True), cabeceras, resumen_b24, len(df) - len(cand)


def numerico(df, cols, nombre, lo=None, hi=None):
    out = df[cols].apply(pd.to_numeric, errors="coerce")
    if out.isna().any().any():
        malas = df.loc[out.isna().any(axis=1), "codigo"].tolist()
        sys.exit(f"ERROR: valores faltantes o no numéricos en {nombre} para las sesiones {malas}.")
    if lo is not None and ((out < lo) | (out > hi)).any().any():
        sys.exit(f"ERROR: valores fuera de la escala {lo}–{hi} en {nombre}.")
    return out


# --------------------------------------------------------------------------- estadística
def alfa_cronbach(items):
    items = np.asarray(items, dtype=float)
    k = items.shape[1]
    return k / (k - 1) * (1 - items.var(axis=0, ddof=1).sum() / items.sum(axis=1).var(ddof=1))


def ic_bootstrap(f, n, rng, n_boot=N_BOOT):
    vals = np.array([f(rng.integers(0, n, n)) for _ in range(n_boot)])
    return [float(np.nanpercentile(vals, 2.5)), float(np.nanpercentile(vals, 97.5))]


def contraste_pareado(mayor, menor):
    """H1: media(mayor − menor) > 0. Shapiro-Wilk sobre las diferencias -> t pareada o Wilcoxon."""
    mayor, menor = np.asarray(mayor, float), np.asarray(menor, float)
    d = mayor - menor
    n = len(d)
    if n < 3:
        sys.exit("ERROR: se necesitan al menos 3 pares.")
    if np.allclose(d, d[0]):
        sw_w, sw_p, normal = float("nan"), float("nan"), False
    else:
        sw_w, sw_p = stats.shapiro(d)
        normal = bool(sw_p > ALPHA)
    res = {"n": n, "shapiro_W": None if np.isnan(sw_w) else float(sw_w), "shapiro_p": None if np.isnan(sw_p) else float(sw_p), "normal": normal}
    if normal:
        t, p = stats.ttest_rel(mayor, menor, alternative="greater")
        res.update(prueba="t de Student pareada (unilateral)", estadistico=float(t), p=float(p), efecto_nombre="d_z", efecto=float(d.mean() / d.std(ddof=1)))
    else:
        nz = d[d != 0]
        try:
            w, p = stats.wilcoxon(mayor, menor, alternative="greater")
            wp, wm = stats.rankdata(np.abs(nz))[nz > 0].sum(), stats.rankdata(np.abs(nz))[nz < 0].sum()
            res.update(prueba="Wilcoxon de rangos con signo (unilateral)", estadistico=float(w), p=float(p), efecto_nombre="r (correlación biserial por rangos)",
                       efecto=float((wp - wm) / (wp + wm)))
        except ValueError:
            res.update(prueba="Wilcoxon: no calculable (todas las diferencias son 0)", estadistico=None, p=1.0, efecto_nombre="r", efecto=0.0)
    res["decision"] = "se rechaza H0" if res["p"] < ALPHA else "no se rechaza H0"
    return res


def f1_macro(y, p, labels):
    from sklearn.metrics import f1_score
    return float(f1_score(y, p, labels=labels, average="macro", zero_division=0))


# --------------------------------------------------------------------------- prueba final
def predictor_modelo(ruta_modelo):
    def _pred(textos):
        sys.path.insert(0, str(Path(__file__).parent))
        import eval_real
        return eval_real.predict_conf(ruta_modelo, textos)
    return _pred


def cargar_congelado(ruta):
    ruta = Path(ruta)
    if not ruta.exists():
        sys.exit(f"ERROR: no existe {ruta}: congela el modelo antes con scripts/congelar_modelo.py")
    fz = json.loads(ruta.read_text(encoding="utf-8"))
    mp = Path(fz["archivos"]["modelo"])
    mp = mp if mp.is_absolute() else ROOT / mp
    if not mp.exists():
        sys.exit(f"ERROR: falta el modelo {fz['archivos']['modelo']}.")
    if file_sha256(mp) != fz["sha256"]["modelo"]:
        sys.exit("ERROR: el hash del modelo NO coincide con modelo_congelado.json. El modelo cambió desde el congelamiento: se aborta la prueba final.")
    avisos = []
    for clave, esperado in fz["sha256"].items():
        if clave == "modelo":
            continue
        p = Path(fz["archivos"][clave])
        p = p if p.is_absolute() else ROOT / p
        if not p.exists() or file_sha256(p) != esperado:
            avisos.append(f"{clave} difiere del congelamiento ({fz['archivos'][clave]})")
    return fz, mp, avisos


def prueba_final(df, catalogo, fz, predictor, rng):
    cat = pd.read_csv(catalogo, dtype=str).set_index("scenario_id")["intent_esperada"].to_dict()
    filas = []
    for _, r in df.iterrows():
        for k in (1, 2, 3):
            q = r.get(f"t{k}_consulta")
            if q is None or (isinstance(q, float) and np.isnan(q)) or not str(q).strip():
                continue
            tid = str(r[f"t{k}_tarjeta"]).strip().upper()
            if re.fullmatch(r"\d+(\.0)?", tid):
                tid = f"S{int(float(tid)):02d}"
            if tid not in cat:
                sys.exit(f"ERROR: la tarjeta «{tid}» de la sesión {r['codigo']} no está en el catálogo.")
            filas.append({"sesion": r["codigo"], "tarjeta": tid, "consulta": str(q).strip(), "esperada": cat[tid],
                          "correcta_aplicador": r.get(f"t{k}_correcta"), "obtuvo_persona": r.get(f"t{k}_obtuvo")})
    if not filas:
        sys.exit("ERROR: no hay consultas registradas en las sesiones elegibles.")
    q = pd.DataFrame(filas)
    pred = predictor(q["consulta"].tolist())
    q["predicha"] = [p[0] for p in pred]
    q["confianza"] = [float(p[1]) for p in pred]
    q["confianza_2"] = [float(p[2]) for p in pred]
    q["acierto"] = q["predicha"] == q["esperada"]
    labels = sorted(q["esperada"].unique())
    ses = q["sesion"].unique()
    por_ses = {s: q[q["sesion"] == s] for s in ses}

    def metricas(idx):
        sub = pd.concat([por_ses[ses[i]] for i in idx]) if len(idx) else q
        return float(sub["acierto"].mean()), f1_macro(sub["esperada"], sub["predicha"], labels)

    acc, f1 = metricas(range(len(ses)))
    boots = np.array([metricas(rng.integers(0, len(ses), len(ses))) for _ in range(N_BOOT)])
    out = {"n_sesiones": int(len(ses)), "n_consultas": int(len(q)), "accuracy": acc, "f1_macro_intenciones_presentes": f1,
           "ic95_accuracy_bootstrap_por_sesion": [float(np.percentile(boots[:, 0], 2.5)), float(np.percentile(boots[:, 0], 97.5))],
           "ic95_f1_bootstrap_por_sesion": [float(np.percentile(boots[:, 1], 2.5)), float(np.percentile(boots[:, 1], 97.5))],
           "n_por_intencion": {k: int(v) for k, v in q["esperada"].value_counts().items()},
           "confianza": {"media_aciertos": float(q.loc[q["acierto"], "confianza"].mean()) if q["acierto"].any() else None,
                         "media_errores": float(q.loc[~q["acierto"], "confianza"].mean()) if (~q["acierto"]).any() else None,
                         "mediana": float(q["confianza"].median())}}
    t = fz.get("umbral_t")
    if t is not None:
        amb = (fz.get("umbral_detalle") or {}).get("ambiguity_threshold", 0.1)
        ab = (q["confianza"] < t) | ((q["confianza"] - q["confianza_2"]) < amb)
        resp = ~ab
        out["umbral"] = {"t": t, "ambiguedad": amb, "cobertura": float(resp.mean()),
                         "precision_respondida": float(q.loc[resp, "acierto"].mean()) if resp.any() else None,
                         "errores_atrapados": int((ab & ~q["acierto"]).sum()), "aciertos_perdidos": int((ab & q["acierto"]).sum())}
    else:
        out["umbral"] = None
    acuerdos = {}
    for col, nom in (("correcta_aplicador", "juicio_del_aplicador"), ("obtuvo_persona", "lo_que_necesitaba_la_persona")):
        v = q[col].dropna()
        if len(v):
            acuerdos[nom] = {"n": int(len(v)), "acuerdo_con_modelo_pct": float((v.map(si).values == q.loc[v.index, "acierto"].values).mean() * 100),
                             "respondieron_si_pct": float(v.map(si).mean() * 100)}
    out["acuerdos"] = acuerdos
    return out, q


# --------------------------------------------------------------------------- figuras e informe
def figuras(pre, post, salida, sufijo):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    rot = "SIMULADO" if sufijo else "REAL"
    fig, ax = plt.subplots(figsize=(4.5, 4))
    ax.boxplot([pre, post], labels=["pre (recordado)", "post (asistente)"])
    ax.set_ylabel("minutos"); ax.set_title(f"Tiempo por persona — {rot}, n = {len(pre)}")
    fig.tight_layout(); fig.savefig(salida / f"tiempos_cajas{sufijo}.png", dpi=130); plt.close(fig)
    fig, ax = plt.subplots(figsize=(4.5, 4))
    ax.hist(np.asarray(post) - np.asarray(pre), bins=max(5, min(15, len(pre) // 3)), color="#4c78a8", edgecolor="white")
    ax.axvline(0, color="k", lw=1); ax.set_xlabel("post − pre (min)"); ax.set_title(f"Diferencias — {rot}, n = {len(pre)}")
    fig.tight_layout(); fig.savefig(salida / f"tiempos_histograma_diferencias{sufijo}.png", dpi=130); plt.close(fig)


def fmt(x, d=3):
    return "—" if x is None else f"{x:.{d}f}"


def informe_md(R):
    rot = R["rotulo"]
    L = [f"# Análisis del piloto exploratorio (V1.3) — {rot}", "",
         f"**{rot}** · generado {R['fecha']} · sesiones elegibles analizadas: **n = {R['n']}**" +
         (f" · **exploratorio, n = {R['n']}** (menos que las {N_OBJETIVO} previstas)" if R["n"] < N_OBJETIVO else f" · piloto exploratorio (objetivo {N_OBJETIVO})"),
         "", "Exploratorio: estos resultados no son confirmatorios.", ""]
    t = R["tiempo"]
    L += ["## OE3 — Tiempo de respuesta (minutos)", "",
          f"- Pre: media {fmt(t['media_pre'], 2)}, mediana {fmt(t['mediana_pre'], 2)} · Post: media {fmt(t['media_post'], 2)}, mediana {fmt(t['mediana_post'], 2)}",
          f"- Diferencia post − pre: media {fmt(t['media_dif_post_menos_pre'], 2)}",
          f"- Shapiro-Wilk sobre las diferencias: W = {fmt(t['contraste']['shapiro_W'], 4)}, p = {fmt(t['contraste']['shapiro_p'], 4)} -> "
          f"{'normales' if t['contraste']['normal'] else 'no normales'}",
          f"- Prueba: {t['contraste']['prueba']}; estadístico = {fmt(t['contraste']['estadistico'], 4)}, p = {t['contraste']['p']:.6f} ({t['contraste']['decision']}, α = {ALPHA})",
          f"- Tamaño de efecto {t['contraste']['efecto_nombre']}: {fmt(t['contraste']['efecto'], 3)}",
          f"- Reducción de las medias: {t['reduccion_medias'] * 100:.1f} % (IC95 % bootstrap {t['ic95_reduccion_medias'][0] * 100:.1f} % a {t['ic95_reduccion_medias'][1] * 100:.1f} %)",
          f"- Media de las reducciones individuales: {t['media_reducciones_individuales'] * 100:.1f} %",
          f"- IC95 % bootstrap de la diferencia media pre − post: [{t['ic95_dif_pre_menos_post'][0]:.2f}, {t['ic95_dif_pre_menos_post'][1]:.2f}] min "
          f"({N_BOOT} remuestreos, semilla {SEMILLA})", ""]
    s = R["satisfaccion"]
    L += ["## OE4 — Satisfacción", "",
          f"- P10 frente al ítem 9: P10 media {fmt(s['media_p10'], 2)}, ítem 9 media {fmt(s['media_item9'], 2)}; {s['contraste']['prueba']}, p = {s['contraste']['p']:.6f} "
          f"({s['contraste']['decision']}); efecto {s['contraste']['efecto_nombre']} = {fmt(s['contraste']['efecto'], 3)}; Shapiro p = {fmt(s['contraste']['shapiro_p'], 4)}",
          f"- Media de los ítems 1–8: {fmt(s['media_items_1_8'], 3)} (IC95 % [{s['ic95_media_items_1_8'][0]:.3f}, {s['ic95_media_items_1_8'][1]:.3f}])",
          f"- Alfa de Cronbach (ítems 1–8): {fmt(s['alfa_cronbach'], 3)}" +
          (f" · Resumen!B24 del registro: {fmt(s['alfa_resumen_b24'], 3)}" if s["alfa_resumen_b24"] is not None else " · Resumen!B24 no disponible") +
          (" · **AVISO: difieren más de 0,01**" if s["alfa_difiere"] else ""), ""]
    p = R["prueba_final"]
    L += ["## Prueba final del modelo congelado (evaluada una sola vez)", "",
          f"- Modelo: {R['modelo_congelado']['modelo_nombre']} (sha256 {R['modelo_congelado']['sha256'][:12]}…, congelado {R['modelo_congelado']['fecha']})",
          f"- {p['n_consultas']} consultas de {p['n_sesiones']} sesiones · accuracy {p['accuracy']:.3f} (IC95 % por sesión [{p['ic95_accuracy_bootstrap_por_sesion'][0]:.3f}, {p['ic95_accuracy_bootstrap_por_sesion'][1]:.3f}])",
          f"- F1 macro (intenciones presentes) {p['f1_macro_intenciones_presentes']:.3f} (IC95 % por sesión [{p['ic95_f1_bootstrap_por_sesion'][0]:.3f}, {p['ic95_f1_bootstrap_por_sesion'][1]:.3f}])",
          f"- Confianza: media en aciertos {fmt(p['confianza']['media_aciertos'])}, en errores {fmt(p['confianza']['media_errores'])}, mediana {fmt(p['confianza']['mediana'])}"]
    if p["umbral"]:
        u = p["umbral"]
        L.append(f"- Umbral congelado t = {u['t']}: cobertura {u['cobertura']:.3f}, precisión de lo respondido {fmt(u['precision_respondida'])}, "
                 f"errores atrapados {u['errores_atrapados']}, aciertos perdidos {u['aciertos_perdidos']}")
    else:
        L.append("- Sin umbral congelado.")
    for k, v in p["acuerdos"].items():
        L.append(f"- Acuerdo del modelo con {k.replace('_', ' ')}: {v['acuerdo_con_modelo_pct']:.1f} % (n = {v['n']}; respondieron «Sí»: {v['respondieron_si_pct']:.1f} %)")
    L += ["", "n por intención: " + ", ".join(f"{k} {v}" for k, v in sorted(p["n_por_intencion"].items())), "",
          "## Limitaciones", ""] + [f"- {x}" for x in LIMITACIONES] + ["",
          "## Columnas del registro usadas", "", "| Campo | Encabezado |", "|---|---|"] + [f"| {k} | {v} |" for k, v in R["columnas"].items()]
    return "\n".join(L) + "\n"


# --------------------------------------------------------------------------- principal
def main(argv=None, predictor=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--registro", required=True)
    ap.add_argument("--modelo-congelado", default=str(ROOT / "logs" / "v3_real" / "modelo_congelado.json"))
    ap.add_argument("--catalogo", default=str(ROOT / "docs" / "lote_real_1" / "situaciones_lote1_v1.csv"))
    ap.add_argument("--salida", default=str(ROOT / "logs" / "piloto"))
    ap.add_argument("--mapa-columnas", default="", help="JSON {campo: encabezado} para los nombres que no se reconocen solos")
    ap.add_argument("--permitir-simulado", action="store_true", help="solo para pruebas: escribe en una carpeta temporal y rotula _SIMULADO")
    ap.add_argument("--motivo", default="", help="obligatorio para repetir la prueba final")
    a = ap.parse_args(argv)

    mapa_usr = json.loads(Path(a.mapa_columnas).read_text(encoding="utf-8")) if a.mapa_columnas else None
    df, cab, b24, excl = leer_registro(a.registro, mapa_usr, a.permitir_simulado)
    simulado = bool(a.permitir_simulado)
    sufijo = "_SIMULADO" if simulado else ""
    salida = Path(a.salida)
    if simulado:
        tmp = Path(tempfile.gettempdir()).resolve()
        if tmp not in salida.resolve().parents and salida.resolve() != tmp:
            salida = Path(tempfile.mkdtemp(prefix="piloto_simulado_"))
            print(f"Modo simulado: las salidas van a la carpeta temporal {salida}")
    salida.mkdir(parents=True, exist_ok=True)
    n = len(df)
    print(f"Sesiones elegibles: {n} (filas con código que no empieza por SA, p. ej. EJ01: {excl}) — {'exploratorio, n = ' + str(n) if n < N_OBJETIVO else 'n completo'}")

    fz, ruta_modelo, avisos = cargar_congelado(a.modelo_congelado)
    for av in avisos:
        print("AVISO:", av)

    log = salida / "evaluaciones_prueba_final.log"
    previas = [l for l in log.read_text(encoding="utf-8").splitlines() if l.strip()] if log.exists() else []
    if previas and not a.motivo.strip():
        sys.exit(f"ERROR: la prueba final ya se evaluó ({len(previas)} vez/veces, ver {log}). El test se evalúa una sola vez; repetirlo exige --motivo \"<texto>\".")

    rng = np.random.default_rng(SEMILLA)
    pre = numerico(df, ["min_pre"], "minutos pre")["min_pre"].values
    post = numerico(df, ["min_post"], "minutos post")["min_post"].values
    c_t = contraste_pareado(pre, post)
    red_ind = np.where(pre > 0, (pre - post) / np.where(pre > 0, pre, 1), np.nan)
    tiempo = {"media_pre": float(pre.mean()), "mediana_pre": float(np.median(pre)), "media_post": float(post.mean()), "mediana_post": float(np.median(post)),
              "media_dif_post_menos_pre": float((post - pre).mean()), "contraste": c_t, "reduccion_medias": float((pre.mean() - post.mean()) / pre.mean()),
              "media_reducciones_individuales": float(np.nanmean(red_ind)),
              "ic95_reduccion_medias": ic_bootstrap(lambda i: (pre[i].mean() - post[i].mean()) / pre[i].mean(), n, rng),
              "ic95_dif_pre_menos_post": ic_bootstrap(lambda i: (pre[i] - post[i]).mean(), n, rng)}

    p10 = numerico(df, ["p10"], "P10", 1, 5)["p10"].values
    it = numerico(df, [f"item{i}" for i in range(1, 10)], "ítems 1–9", 1, 5)
    i9 = it["item9"].values
    i18 = it[[f"item{i}" for i in range(1, 9)]].values
    punt = i18.mean(axis=1)
    ci = stats.t.interval(0.95, n - 1, loc=punt.mean(), scale=punt.std(ddof=1) / np.sqrt(n))
    alfa = float(alfa_cronbach(i18))
    b24n = float(b24) if isinstance(b24, (int, float)) else None
    satis = {"media_p10": float(p10.mean()), "media_item9": float(i9.mean()), "contraste": contraste_pareado(i9, p10),
             "media_items_1_8": float(punt.mean()), "ic95_media_items_1_8": [float(ci[0]), float(ci[1])], "alfa_cronbach": alfa,
             "alfa_resumen_b24": b24n, "alfa_difiere": b24n is not None and abs(alfa - b24n) > 0.01}
    if satis["alfa_difiere"]:
        print(f"AVISO: el alfa calculado ({alfa:.3f}) difiere de Resumen!B24 ({b24n:.3f}) en más de 0,01.")

    predictor = predictor or predictor_modelo(ruta_modelo)
    pf, detalle = prueba_final(df, a.catalogo, fz, predictor, rng)

    R = {"rotulo": "SIMULADO" if simulado else "REAL", "fecha": datetime.now().isoformat(timespec="seconds"), "n": n,
         "exploratorio": n < N_OBJETIVO, "sesiones_excluidas_no_SA": excl, "tiempo": tiempo, "satisfaccion": satis, "prueba_final": pf,
         "modelo_congelado": {"modelo_nombre": fz.get("modelo_nombre", Path(ruta_modelo).name), "sha256": fz["sha256"]["modelo"], "fecha": fz["fecha"]},
         "limitaciones": LIMITACIONES, "columnas": cab}
    (salida / f"analisis_piloto{sufijo}.json").write_text(json.dumps(R, indent=2, ensure_ascii=False), encoding="utf-8")
    (salida / f"analisis_piloto{sufijo}.md").write_text(informe_md(R), encoding="utf-8")
    detalle.drop(columns=["consulta"]).to_csv(salida / f"prueba_final_detalle{sufijo}.csv", index=False, encoding="utf-8")  # sin el texto de las consultas
    figuras(pre, post, salida, sufijo)
    with open(log, "a", encoding="utf-8") as f:
        f.write(json.dumps({"fecha": R["fecha"], "modelo_sha256": fz["sha256"]["modelo"], "n_sesiones": n, "n_consultas": pf["n_consultas"],
                            "motivo": a.motivo.strip(), "rotulo": R["rotulo"]}, ensure_ascii=False) + "\n")
    print(informe_md(R))
    print(f"Guardado en {salida}")
    return R


if __name__ == "__main__":
    main()
