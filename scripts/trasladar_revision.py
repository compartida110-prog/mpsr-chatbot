"""Traslada las decisiones de la revisión HUMANA de etiquetas (libro Excel del tesista, hoja «Revision») a corpus/real/lote1_revision_etiquetas.csv, y informa la cobertura resultante.

Empareja por participant_code + scenario_id (a través de lote1_real_validado.csv, que trae real_id). ANTES de escribir comprueba, fila por fila, que la frase y la intención esperada del libro
coinciden con las del CSV; si algo no coincide, se detiene sin escribir. Copia la decisión (OK / CAMBIAR / DESCARTAR) y, en CAMBIAR, la intención correcta, tal como las anotó la revisora; no cambia
ninguna decisión ni las rotula como automáticas (revisión humana de una sola revisora: sin segunda revisora no hay kappa y se declara). Solo LEE el libro (nunca lo guarda).

Informa cuántas frases quedan por intención con esas decisiones (CAMBIAR mueve la frase a la intención correcta; DESCARTAR la quita) y qué intenciones quedan bajo el mínimo (3).
NO ejecuta --aplicar-revision ni ninguna etapa posterior: con intenciones bajo el mínimo hay que reunir más frases antes.

Uso:
    python scripts/trasladar_revision.py --libro docs/lote_real_1/privado/Revision_Etiquetas_Lote1_v1_1.xlsx [--solo-verificar]
"""
import argparse
import shutil
import sys
import unicodedata
from collections import Counter
from pathlib import Path

import pandas as pd
import yaml

from common import ROOT

MIN_FRASES = 3
DECISIONES = {"OK", "CAMBIAR", "DESCARTAR"}


def norm(t):
    return " ".join(unicodedata.normalize("NFC", str(t or "")).split())


def leer_libro(ruta):
    from openpyxl import load_workbook
    wb = load_workbook(ruta, read_only=True, data_only=True)
    try:
        filas = list(wb["Revision"].iter_rows(values_only=True))
    finally:
        wb.close()
    fe = next((i for i, f in enumerate(filas[:10]) if f and "participant_code" in [norm(c) for c in f]), None)
    if fe is None:
        sys.exit("ERROR: no encuentro el encabezado «participant_code» en la hoja Revision.")
    enc = [norm(c) for c in filas[fe]]

    def col(prefijo):
        j = next((i for i, c in enumerate(enc) if c.startswith(prefijo)), None)
        if j is None:
            sys.exit(f"ERROR: falta la columna «{prefijo}…» en la hoja Revision.")
        return j
    j = {k: col(p) for k, p in (("p", "participant_code"), ("s", "scenario_id"), ("esp", "Intención esperada"), ("txt", "Frase escrita"), ("dec", "Decisión"), ("int", "Intención correcta"), ("com", "Comentario"))}
    out = []
    for f in filas[fe + 1:]:
        if f and f[j["p"]]:
            out.append({k: norm(f[i]) for k, i in j.items()})
    return pd.DataFrame(out)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--libro", required=True)
    ap.add_argument("--validado", default=str(ROOT / "corpus" / "real" / "lote1_real_validado.csv"))
    ap.add_argument("--revision", default=str(ROOT / "corpus" / "real" / "lote1_revision_etiquetas.csv"))
    ap.add_argument("--domain", default=str(ROOT / "domain.yml"))
    ap.add_argument("--base-revision", default="", help="revisión ya trasladada (p. ej. las 169 primeras): se suma al informe de cobertura y, con --unir, se le agregan estas filas")
    ap.add_argument("--unir", action="store_true", help="agrega las filas trasladadas a --base-revision (con copia previa) para que --aplicar-revision vea todas las frases")
    ap.add_argument("--solo-verificar", action="store_true", help="comprueba y informa sin escribir")
    a = ap.parse_args()
    lib = leer_libro(a.libro)
    val = pd.read_csv(a.validado, dtype=str, keep_default_na=False, encoding="utf-8-sig")
    rev = pd.read_csv(a.revision, dtype=str, keep_default_na=False, encoding="utf-8-sig")
    val = val[val["real_id"].isin(rev["real_id"])]  # solo las frases de la hoja de revisión que se está trasladando (puede ser un subconjunto)
    intents = set(yaml.safe_load(open(a.domain, encoding="utf-8"))["intents"])
    problemas = []
    if lib.duplicated(["p", "s"]).any():
        problemas.append("el libro tiene pares participant_code + scenario_id repetidos")
    v = val.assign(clave=val["participant_code"] + "|" + val["scenario_id"]).set_index("clave")
    lib["clave"] = lib["p"] + "|" + lib["s"]
    faltan, sobran = sorted(set(v.index) - set(lib["clave"])), sorted(set(lib["clave"]) - set(v.index))
    if faltan or sobran:
        problemas.append(f"pares sin emparejar: {len(faltan)} del CSV que no están en el libro y {len(sobran)} del libro que no están en el CSV")
    rev_texto = dict(zip(rev["real_id"], rev["text"]))
    rev_esp = dict(zip(rev["real_id"], rev["intent_esperada"]))
    nuevas = {}
    for r in lib.itertuples():
        if r.clave not in v.index:
            continue
        rid = v.loc[r.clave, "real_id"]
        if norm(r.txt) != norm(v.loc[r.clave, "text"]) or norm(r.txt) != norm(rev_texto.get(rid)):
            problemas.append(f"{r.p} {r.s}: la frase del libro no coincide con la del CSV")
        if r.esp != norm(v.loc[r.clave, "intent_esperada"]) or r.esp != norm(rev_esp.get(rid)):
            problemas.append(f"{r.p} {r.s}: la intención esperada del libro no coincide con la del CSV")
        d = r.dec.upper()
        if d not in DECISIONES:
            problemas.append(f"{r.p} {r.s}: decisión «{r.dec}» no válida (OK / CAMBIAR / DESCARTAR)")
        if d == "CAMBIAR" and r.int not in intents:
            problemas.append(f"{r.p} {r.s}: CAMBIAR sin una intención correcta válida («{r.int}»)")
        if d == "OK" and r.int:
            problemas.append(f"{r.p} {r.s}: OK pero trae una intención correcta anotada")
        nuevas[rid] = (d, r.int if d == "CAMBIAR" else "", r.com)
    if problemas:
        print("SE DETIENE (no se escribió nada):\n" + "\n".join(f"  - {p}" for p in problemas[:40]))
        sys.exit(2)
    cuenta = Counter(d for d, _, _ in nuevas.values())
    print(f"Emparejadas: {len(nuevas)} de {len(val)} | OK {cuenta['OK']} | CAMBIAR {cuenta['CAMBIAR']} | DESCARTAR {cuenta['DESCARTAR']}")
    # ---- cobertura con esas decisiones
    esperada = dict(zip(rev["real_id"], rev["intent_esperada"]))
    destino = Counter()
    if a.base_revision:  # lo ya trasladado antes también cuenta en la cobertura
        base = pd.read_csv(a.base_revision, dtype=str, keep_default_na=False, encoding="utf-8-sig")
        if set(base["real_id"]) & set(rev["real_id"]):
            sys.exit("ERROR: --base-revision comparte real_id con la hoja que se traslada.")
        if (base["decision"] == "").any():
            sys.exit("ERROR: --base-revision tiene filas sin decisión.")
        for r in base.itertuples():
            esperada[r.real_id] = r.intent_esperada
            if r.decision != "DESCARTAR":
                destino[r.intent_revisada if r.decision == "CAMBIAR" else r.intent_esperada] += 1
        print(f"Cobertura junto con {len(base)} frases ya trasladadas de {Path(a.base_revision).name}.")
    for rid, (d, nueva, _) in nuevas.items():
        if d != "DESCARTAR":
            destino[nueva if d == "CAMBIAR" else esperada[rid]] += 1
    todas = sorted(intents - {"nlu_fallback", "out_of_scope"}) if False else sorted(set(esperada.values()) | set(destino))
    bajo = {i: destino.get(i, 0) for i in todas if destino.get(i, 0) < MIN_FRASES}
    print(f"Intenciones: {len(todas)} | por debajo del mínimo ({MIN_FRASES}): {len(bajo)} -> {bajo}")
    cambios = Counter((esperada[r], n) for r, (d, n, _) in nuevas.items() if d == "CAMBIAR")
    print(f"Cambios de etiqueta: {sum(cambios.values())} (de→a, en {len(cambios)} pares distintos)")
    print("Revisión humana de UNA revisora: sin segunda revisora, el kappa de Cohen NO se calculó.")
    if a.solo_verificar:
        print("(--solo-verificar: no se escribió nada)")
        return
    copia = Path(a.revision).with_name(Path(a.revision).stem + ".ANTES_DE_TRASLADAR.csv")
    if not copia.exists():
        shutil.copyfile(a.revision, copia)
    rev["decision"] = [nuevas[r][0] for r in rev["real_id"]]
    rev["intent_revisada"] = [nuevas[r][1] for r in rev["real_id"]]
    rev["comentario"] = [nuevas[r][2] for r in rev["real_id"]]
    rev.to_csv(a.revision, index=False, encoding="utf-8")
    print(f"Escrito {a.revision} (copia previa: {copia.name}).")
    if a.unir:
        if not a.base_revision:
            sys.exit("ERROR: --unir necesita --base-revision.")
        cb = Path(a.base_revision).with_name(Path(a.base_revision).stem + ".ANTES_DE_UNIR.csv")
        shutil.copyfile(a.base_revision, cb)
        pd.concat([base, rev], ignore_index=True).to_csv(a.base_revision, index=False, encoding="utf-8")
        print(f"Unidas: {a.base_revision} ahora tiene {len(base) + len(rev)} filas (copia previa: {cb.name}).")
    print("NO se ejecutó --aplicar-revision.")


if __name__ == "__main__":
    main()
