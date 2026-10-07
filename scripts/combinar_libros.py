"""Combina en un libro de transcripción (BASE) los datos de otro libro (FUENTE): solo las celdas de ENTRADA de Participantes y Respuestas, editando el XML de la base (no se guarda con openpyxl: se
conservan validaciones, formato y fórmulas). Ni la base ni la fuente se modifican: el resultado va a --salida (en la carpeta privada, fuera de Git).

  * Participantes: se emparejan por participant_code; se copian Estado, rango de edad, vive en Juliaca, trámite, consentimiento, fecha y observaciones (columnas C–H y M) a las filas de la base.
  * Respuestas: se emparejan por participant_code + scenario_id; se copia la frase (columna D).
  * Antes de escribir comprueba que el formulario y la situación coinciden y que las celdas de destino de la base están VACÍAS (o con el «Pendiente» por defecto del Estado): nunca sobrescribe datos
    existentes; si algo no cuadra, se detiene sin escribir. Las fórmulas no se tocan; las de las filas copiadas se verifican (que apunten a su propia fila).
  * Al final borra los valores guardados de las fórmulas y marca fullCalcOnLoad: Excel recalcula todo al abrir. Agrega «Formulario F» al Resumen si falta.

Uso:
    python scripts/combinar_libros.py --base <libro con P26–P28> --fuente <libro con P01–P25> --salida docs/lote_real_1/privado/<nombre>.xlsx
"""
import argparse
import re
import sys
import warnings
import zipfile
from datetime import datetime
from pathlib import Path

warnings.filterwarnings("ignore")
from openpyxl import load_workbook  # noqa: E402
from openpyxl.utils.datetime import to_excel  # noqa: E402

import libro_xml as lx  # noqa: E402

COLS_PART = ["C", "D", "E", "F", "G", "H", "M"]


def rutas_hojas(arch):
    wb = arch["xl/workbook.xml"].decode("utf-8")
    rels = arch["xl/_rels/workbook.xml.rels"].decode("utf-8")
    ids = re.findall(r'<sheet\s[^>]*?name="([^"]+)"[^>]*?r:id="(rId\d+)"', wb)
    dest = {}
    for tag in re.findall(r"<Relationship\s[^>]*?/>", rels):
        i, t = re.search(r'Id="([^"]+)"', tag), re.search(r'Target="([^"]+)"', tag)
        dest[i.group(1)] = "xl/" + t.group(1).replace("/xl/", "").lstrip("/").replace("xl/", "")
    return {n: dest[r] for n, r in ids}


def leer_entrada(ruta):
    """{codigo: {col: valor}} de Participantes y {(codigo, situacion): frase} de Respuestas (solo lectura)."""
    wb = load_workbook(ruta, read_only=True, data_only=True)
    try:
        part, resp, forma, sit_forma = {}, {}, {}, {}
        for f in wb["Participantes"].iter_rows(min_row=4):
            if f[0].value:
                cod = str(f[0].value).strip()
                part[cod] = {c: f[ord(c) - 65].value for c in COLS_PART}
                forma[cod] = f[1].value
        for f in wb["Respuestas"].iter_rows(min_row=4, max_col=4):
            if f[0].value:
                resp[(str(f[0].value).strip(), str(f[2].value).strip())] = f[3].value
                sit_forma[(str(f[0].value).strip(), str(f[2].value).strip())] = f[1].value
        return part, resp, forma, sit_forma
    finally:
        wb.close()


def valor(v):
    if isinstance(v, datetime):
        return float(to_excel(v))
    if isinstance(v, str):
        return v
    return v


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--base", required=True)
    ap.add_argument("--fuente", required=True)
    ap.add_argument("--salida", required=True)
    a = ap.parse_args()
    if Path(a.salida).resolve() in (Path(a.base).resolve(), Path(a.fuente).resolve()):
        sys.exit("ERROR: --salida no puede ser la base ni la fuente (no se modifican).")
    pb, rb, fb, sb = leer_entrada(a.base)
    pf, rf, ff, sf = leer_entrada(a.fuente)
    problemas = []
    # -------- Participantes: solo los de la fuente que están en la base; vacíos en destino
    comunes = [c for c in pf if c in pb]
    for c in pf:
        if c not in pb:
            problemas.append(f"{c} está en la fuente pero no en la base")
    for c in comunes:
        if fb[c] != ff[c]:
            problemas.append(f"{c}: el formulario de la base ({fb[c]}) y el de la fuente ({ff[c]}) difieren")
        for col in COLS_PART:
            if pb[c][col] not in (None, "") and not (col == "C" and pb[c][col] == "Pendiente"):
                problemas.append(f"{c}: la celda {col} de la base ya tiene datos ({'estado' if col == 'C' else 'dato'}): no se sobrescribe")
    # -------- Respuestas
    pares = [k for k, v in rf.items() if v not in (None, "") and str(v).strip()]
    for k in pares:
        if k not in rb:
            problemas.append(f"{k[0]} {k[1]}: la fila de Respuestas no existe en la base")
        elif sb[k] != sf[k]:
            problemas.append(f"{k[0]} {k[1]}: formulario distinto entre base y fuente")
        elif rb[k] not in (None, "") and str(rb[k]).strip():
            problemas.append(f"{k[0]} {k[1]}: la base ya tiene una frase: no se sobrescribe")
    if problemas:
        print("SE DETIENE (no se escribió nada):\n" + "\n".join(f"  - {p}" for p in problemas[:40]))
        sys.exit(2)
    # -------- editar el XML de la base
    zin = zipfile.ZipFile(a.base)
    arch = {n: zin.read(n) for n in zin.namelist()}
    rutas = rutas_hojas(arch)
    hoja = {n: arch[r].decode("utf-8") for n, r in rutas.items()}
    # fila de cada participante / de cada (participante, situación) en la base, leída del propio libro
    wbb = load_workbook(a.base, read_only=True, data_only=True)
    fila_p = {str(f[0].value).strip(): f[0].row for f in wbb["Participantes"].iter_rows(min_row=4, max_col=1) if f[0].value}
    fila_r = {(str(f[0].value).strip(), str(f[2].value).strip()): f[0].row for f in wbb["Respuestas"].iter_rows(min_row=4, max_col=3) if f[0].value}
    wbb.close()
    n_part = n_resp = 0
    for c in comunes:
        r = fila_p[c]
        for col in COLS_PART:
            v = valor(pf[c][col])
            if v in (None, ""):
                continue
            hoja["Participantes"] = lx.set_celda(hoja["Participantes"], f"{col}{r}", v)
        n_part += 1
    for k in pares:
        hoja["Respuestas"] = lx.set_celda(hoja["Respuestas"], f"D{fila_r[k]}", str(rf[k]))
        n_resp += 1
    # -------- las fórmulas de todas las filas deben apuntar a su propia fila
    ultimo_p = max(fila_p.values())
    ultimo_r = max(fila_r.values())
    malos = lx.verificar_filas(hoja["Participantes"], range(4, ultimo_p + 1)) + lx.verificar_filas(hoja["Respuestas"], range(4, ultimo_r + 1))
    if malos:
        sys.exit("ERROR: referencias relativas mal desplazadas en el libro base (no se escribió nada): " + "; ".join(malos[:8]))
    hoja["Resumen"] = lx.agregar_fila_f(hoja["Resumen"])
    for nombre in ("Participantes", "Respuestas", "Resumen", "Cobertura", "Situaciones"):
        hoja[nombre] = lx.quitar_valores_guardados(hoja[nombre])
    wbx = arch["xl/workbook.xml"].decode("utf-8")
    if "fullCalcOnLoad" not in wbx:
        wbx = re.sub(r"<calcPr([^>]*?)/>", r'<calcPr\1 fullCalcOnLoad="1"/>', wbx)
    arch["xl/workbook.xml"] = wbx.encode("utf-8")
    for nombre, r in rutas.items():
        arch[r] = hoja[nombre].encode("utf-8")
    salida = Path(a.salida)
    salida.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(salida, "w", zipfile.ZIP_DEFLATED) as zo:
        for info in zin.infolist():
            zo.writestr(info.filename, arch[info.filename])
    print(f"Copiados {n_part} participantes y {n_resp} frases de la fuente a {salida.name}. Base y fuente no se modificaron. Abre el libro en Excel y guárdalo una vez para recalcular los totales.")


if __name__ == "__main__":
    main()
