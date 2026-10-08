"""Agrega a un libro de transcripción (vacío o ya lleno) las filas de un complemento: participantes nuevos con su formulario y una fila de Respuestas por cada situación.

Ejemplo (lote 1c): P29, P30 y P31 con el formulario G y las situaciones S46 y S47 (6 filas de respuestas).

Edita el XML (no se guarda con openpyxl: se conservan validaciones, formato y fórmulas) y NO modifica el libro base: escribe --salida. Las filas nuevas salen SIN datos (Estado «Pendiente»), aunque el libro
base esté lleno; los códigos no pueden existir ya en el libro (no se reutilizan ids). Extiende los rangos de las fórmulas, las validaciones y los formatos condicionales, agrega «Formulario <letra>» al Resumen,
verifica que las fórmulas de las filas nuevas apunten a su propia fila, borra los valores guardados de las fórmulas y marca el recálculo completo al abrir (hay que abrir y guardar el libro una vez en Excel).

Uso:
    python scripts/ampliar_libro.py --base <libro> --salida <libro nuevo> --codigos P29,P30,P31 --forma G --situaciones S46,S47
"""
import argparse
import html
import re
import sys
import warnings
import zipfile
from pathlib import Path

warnings.filterwarnings("ignore")
from openpyxl import load_workbook  # noqa: E402

import libro_xml as lx  # noqa: E402
from combinar_libros import rutas_hojas  # noqa: E402

COLS_PART_ENTRADA = ["D", "E", "F", "G", "H", "M"]  # C (Estado) vuelve a «Pendiente»


def cadenas(z):
    x = z.read("xl/sharedStrings.xml").decode("utf-8")
    return [html.unescape("".join(re.findall(r"<t[^>]*>([^<]*)</t>", si))) for si in re.findall(r"<si>.*?</si>", x, re.S)]


def vaciar(fila_xml, ref):
    """Deja la celda `ref` vacía conservando su estilo."""
    m = re.search(r'<c r="%s"((?:\s+[A-Za-z:]+="[^"]*")*)\s*(?:/>|>.*?</c>)' % ref, fila_xml, re.S)
    if not m:
        return fila_xml
    estilo = re.search(r's="(\d+)"', m.group(1))
    nuevo = f'<c r="{ref}"' + (f' s="{estilo.group(1)}"' if estilo else "") + "/>"
    return fila_xml[:m.start()] + nuevo + fila_xml[m.end():]


def poner_inline(fila_xml, ref, texto):
    m = re.search(r'<c r="%s"((?:\s+[A-Za-z:]+="[^"]*")*)\s*(?:/>|>.*?</c>)' % ref, fila_xml, re.S)
    estilo = re.search(r's="(\d+)"', m.group(1)).group(1) if re.search(r's="(\d+)"', m.group(1)) else ""
    return fila_xml[:m.start()] + lx.inline(ref, estilo, texto) + fila_xml[m.end():]


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--base", required=True)
    ap.add_argument("--salida", required=True)
    ap.add_argument("--codigos", required=True, help="p. ej. P29,P30,P31")
    ap.add_argument("--forma", required=True, help="letra del formulario, p. ej. G")
    ap.add_argument("--situaciones", required=True, help="p. ej. S46,S47")
    a = ap.parse_args()
    codigos, sits = [c.strip() for c in a.codigos.split(",")], [s.strip() for s in a.situaciones.split(",")]
    if Path(a.salida).resolve() == Path(a.base).resolve():
        sys.exit("ERROR: --salida no puede ser la base (no se modifica).")
    wb = load_workbook(a.base, read_only=True, data_only=True)
    try:
        ult_p = max(f[0].row for f in wb["Participantes"].iter_rows(min_row=4, max_col=1) if f[0].value)
        ult_r = max(f[0].row for f in wb["Respuestas"].iter_rows(min_row=4, max_col=1) if f[0].value)
        existentes = {str(f[0].value).strip() for f in wb["Participantes"].iter_rows(min_row=4, max_col=1) if f[0].value}
    finally:
        wb.close()
    repetidos = sorted(set(codigos) & existentes)
    if repetidos:
        sys.exit(f"ERROR: los códigos {repetidos} ya existen en el libro: no se reutilizan ids.")
    zin = zipfile.ZipFile(a.base)
    ss = cadenas(zin)
    arch = {n: zin.read(n) for n in zin.namelist()}
    rutas = rutas_hojas(arch)
    hoja = {n: arch[r].decode("utf-8") for n, r in rutas.items()}
    # ------------------------------------------------ Participantes
    xp = hoja["Participantes"]
    modelo_p = re.search(r'<row r="%d"[^>]*>.*?</row>' % ult_p, xp, re.S).group(0)
    filas_p = []
    for k, cod in enumerate(codigos):
        n = ult_p + 1 + k
        f = lx.desplazar_fila(modelo_p, ult_p, n)
        f = poner_inline(f, f"A{n}", cod)
        f = poner_inline(f, f"B{n}", a.forma)
        f = poner_inline(f, f"C{n}", "Pendiente")
        for c in COLS_PART_ENTRADA:
            f = vaciar(f, f"{c}{n}")
        f = re.sub(r"(<f\b[^>]*?(?:/>|>[^<]*</f>))<v[^>]*?(?:/>|>[^<]*</v>)", r"\1", f)
        filas_p.append(f)
    hoja["Participantes"] = xp.replace(modelo_p, modelo_p + "".join(filas_p), 1)
    # ------------------------------------------------ Respuestas
    xr = hoja["Respuestas"]
    modelo_r = re.search(r'<row r="%d"[^>]*>.*?</row>' % ult_r, xr, re.S).group(0)
    ref = {}
    for m in re.finditer(r'<row r="\d+"[^>]*>.*?</row>', xr, re.S):
        c = re.search(r'<c r="C\d+"[^>]*t="s"[^>]*><v>(\d+)</v>', m.group(0))
        if c and ss[int(c.group(1))] in sits and ss[int(c.group(1))] not in ref:
            ref[ss[int(c.group(1))]] = m.group(0)
    if set(ref) != set(sits):
        sys.exit(f"ERROR: no encuentro en Respuestas las situaciones {sorted(set(sits) - set(ref))}.")
    filas_r, n = [], ult_r
    for cod in codigos:
        for s in sits:
            n += 1
            f = lx.desplazar_fila(modelo_r, ult_r, n)
            f = poner_inline(f, f"A{n}", cod)
            f = poner_inline(f, f"B{n}", a.forma)
            f = poner_inline(f, f"C{n}", s)
            f = vaciar(f, f"D{n}")
            for col in ("E", "F"):  # texto de referencia de la situación y su intención: de una fila existente de esa situación
                m = re.search(r'<c r="%s\d+"((?:\s+[A-Za-z:]+="[^"]*")*)\s*>.*?</c>' % col, ref[s], re.S)
                celda = re.sub(r'r="%s\d+"' % col, f'r="{col}{n}"', m.group(0), count=1)
                f = re.sub(r'<c r="%s%d"(?:\s+[A-Za-z:]+="[^"]*")*\s*(?:/>|>.*?</c>)' % (col, n), lambda _m: celda, f, count=1, flags=re.S)
            f = re.sub(r"(<f\b[^>]*?(?:/>|>[^<]*</f>))<v[^>]*?(?:/>|>[^<]*</v>)", r"\1", f)
            filas_r.append(f)
    hoja["Respuestas"] = xr.replace(modelo_r, modelo_r + "".join(filas_r), 1)
    nuevo_p, nuevo_r = ult_p + len(codigos), ult_r + len(filas_r)
    malos = lx.verificar_filas(hoja["Participantes"], range(ult_p + 1, nuevo_p + 1)) + lx.verificar_filas(hoja["Respuestas"], range(ult_r + 1, nuevo_r + 1))
    if malos:
        sys.exit("ERROR: referencias relativas mal desplazadas (no se escribió nada): " + "; ".join(malos[:6]))
    # ------------------------------------------------ rangos, validaciones, formatos, dimensiones
    for nombre in list(hoja):
        x = re.sub(r"(Participantes!\$[A-Z]\$4:\$[A-Z]\$)%d\b" % ult_p, r"\g<1>%d" % nuevo_p, hoja[nombre])
        x = re.sub(r"(Respuestas!\$[A-Z]\$4:\$[A-Z]\$)%d\b" % ult_r, r"\g<1>%d" % nuevo_r, x)
        if nombre == "Participantes":
            x = re.sub(r'sqref="([A-Z]+)4:([A-Z]+)%d"' % ult_p, r'sqref="\g<1>4:\g<2>%d"' % nuevo_p, x)
            x = x.replace(f'<dimension ref="A1:M{ult_p}"/>', f'<dimension ref="A1:M{nuevo_p}"/>')
        if nombre == "Respuestas":
            x = re.sub(r'sqref="([A-Z]+)4:([A-Z]+)%d"' % ult_r, r'sqref="\g<1>4:\g<2>%d"' % nuevo_r, x)
            x = x.replace(f'<dimension ref="A1:H{ult_r}"/>', f'<dimension ref="A1:H{nuevo_r}"/>')
        hoja[nombre] = x
    letras = sorted(set(re.findall(r">Formulario ([A-Z])<", hoja["Resumen"])) | set(re.findall(r"Formulario ([A-Z])", " ".join(ss))))
    previa = max(l for l in letras if l < a.forma) if any(l < a.forma for l in letras) else None
    if previa:
        hoja["Resumen"] = lx.agregar_fila_formulario(hoja["Resumen"], a.forma, previa)
    for nombre in ("Participantes", "Respuestas", "Resumen", "Cobertura", "Situaciones"):
        hoja[nombre] = lx.quitar_valores_guardados(hoja[nombre])
    wbx = arch["xl/workbook.xml"].decode("utf-8")
    if "fullCalcOnLoad" not in wbx:
        wbx = re.sub(r"<calcPr([^>]*?)/>", r'<calcPr\1 fullCalcOnLoad="1"/>', wbx)
    arch["xl/workbook.xml"] = wbx.encode("utf-8")
    for nombre, r in rutas.items():
        arch[r] = hoja[nombre].encode("utf-8")
    salida = Path(a.salida)
    salida.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(salida, "w", zipfile.ZIP_DEFLATED) as zo:
        for info in zin.infolist():
            zo.writestr(info.filename, arch[info.filename])
    # ------------------------------------------------ comprobación: las filas nuevas no traen datos
    wc = load_workbook(salida, read_only=True, data_only=True)
    try:
        for f in wc["Participantes"].iter_rows(min_row=ult_p + 1, max_row=nuevo_p):
            v = [c.value for c in f[:13]]
            if any(v[3:8]) or v[12] or v[2] != "Pendiente" or v[0] not in codigos or v[1] != a.forma:
                sys.exit(f"ERROR: la fila {f[0].row} de Participantes no quedó vacía: {v}")
        for f in wc["Respuestas"].iter_rows(min_row=ult_r + 1, max_row=nuevo_r, max_col=8):
            v = [c.value for c in f]
            if v[3] or v[0] not in codigos or v[1] != a.forma or v[2] not in sits or not v[4] or not v[5]:
                sys.exit(f"ERROR: la fila {f[0].row} de Respuestas no quedó como se esperaba")
    finally:
        wc.close()
    print(f"Escrito {salida.name}: {len(codigos)} participantes ({codigos[0]}–{codigos[-1]}, formulario {a.forma}) y {len(filas_r)} filas de respuestas ({', '.join(sits)}); "
          f"Participantes 4:{nuevo_p}, Respuestas 4:{nuevo_r}. Filas nuevas sin datos. Base intacta. Abre y guarda el libro una vez en Excel para recalcular.")


if __name__ == "__main__":
    main()
