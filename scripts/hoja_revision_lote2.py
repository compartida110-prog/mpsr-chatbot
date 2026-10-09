"""Lote 2 — hoja de revisión de etiquetas para que el tesista la llene a mano (Excel nuevo en docs/lote_real_2/privado/; trae frases reales: no se versiona ni se imprime).

Mismo formato que la del lote 1 (hoja «Revision» con Decisión OK / CAMBIAR / DESCARTAR e Intención correcta si CAMBIAR; hoja «Intenciones» con la lista de las 54), pero SIN la columna de sugerencia de Claude
y SIN ninguna predicción del modelo: la revisión se hace sin ver lo que predice el modelo. «Nota del sistema» solo trae avisos de procedimiento por identificador (idéntica a entrenamiento, posible otro idioma).
Crea un libro nuevo (no edita ninguno existente). Para trasladar las decisiones después: scripts/trasladar_revision.py (lee las columnas por nombre).

Uso:
    python scripts/hoja_revision_lote2.py
"""
import argparse
from pathlib import Path

import pandas as pd
import yaml
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from common import ROOT, load_jerga, normalize

NOTAS = [
    "Revisión de etiquetas — lote 2 (frases REALES, solo test)",
    "",
    "En la hoja Revision, columna H (amarilla), elige OK, CAMBIAR o DESCARTAR para cada frase.",
    "• OK: la intención esperada corresponde a lo que la persona escribió.",
    "• CAMBIAR: la frase pertenece a otra intención; elígela en la columna I.",
    "• DESCARTAR: solo si es inválida (sin relación con ningún trámite, ilegible, etc.). Los duplicados exactos y las frases idénticas a una de entrenamiento los trata la regla fijada de antemano (split_lote2.py); no hace falta marcarlos aquí.",
    "Criterio: el mismo del lote 1 — la etiqueta se decide por lo que dice la frase, no por la situación que se le pidió.",
    "Nota: «gracias» se usa a la vez para agradecer y para despedirse. Las frases que solo agradecen son agradecimiento; las que cierran la conversación (aunque lleven «gracias») son despedida.",
    "Esta hoja se generó SIN ver ninguna predicción del modelo y no trae sugerencias: decide tú, sin consultar lo que predice el modelo sobre estas frases.",
    "Es solo test: estas decisiones no se usan para ajustar el modelo, el umbral ni el entrenamiento.",
    "Después de decidir, la hoja se pasa a Claude Code solo para trasladar las decisiones (trasladar_revision.py); el archivo no se sube a Git.",
]


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--validado", default=str(ROOT / "corpus" / "real_lote2" / "lote2_real_validado.csv"))
    ap.add_argument("--catalogo", default=str(ROOT / "docs" / "lote_real_2" / "situaciones_lote2_v1.csv"))
    ap.add_argument("--alertas", default=str(ROOT / "corpus" / "real_lote2" / "lote2_alertas_frases.csv"))
    ap.add_argument("--entrenamiento", default=str(ROOT / "corpus" / "v3_lote2" / "entrenamiento_lote2.csv"))
    ap.add_argument("--domain", default=str(ROOT / "domain.yml"))
    ap.add_argument("--salida", default=str(ROOT / "docs" / "lote_real_2" / "privado" / "Revision_Etiquetas_Lote2_v1.xlsx"))
    a = ap.parse_args()
    salida = Path(a.salida)
    if salida.exists():
        raise SystemExit(f"ERROR: ya existe {salida}; no se sobrescribe (puede tener decisiones del tesista).")
    val = pd.read_csv(a.validado, dtype=str, keep_default_na=False, encoding="utf-8")
    cat = pd.read_csv(a.catalogo, dtype=str, keep_default_na=False, encoding="utf-8-sig").set_index("scenario_id")
    ent = pd.read_csv(a.entrenamiento, dtype=str, keep_default_na=False, encoding="utf-8")
    alert = pd.read_csv(a.alertas, dtype=str, keep_default_na=False, encoding="utf-8") if Path(a.alertas).exists() else pd.DataFrame(columns=["real_id", "alerta"])
    intents = sorted(yaml.safe_load(open(a.domain, encoding="utf-8"))["intents"])
    jerga = load_jerga()
    por = ent.assign(_n=ent["text"].map(lambda t: normalize(t, jerga))).groupby("_n")["utterance_id"].first()
    otro = dict(zip(alert["real_id"], alert["alerta"]))

    wb = Workbook()
    ws0 = wb.active
    ws0.title = "Instrucciones"
    for i, t in enumerate(NOTAS, 1):
        ws0.cell(i, 1, t)
    ws0["A1"].font = Font(bold=True, size=13)
    ws0.column_dimensions["A"].width = 130
    for r in ws0.iter_rows():
        for c in r:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws = wb.create_sheet("Revision")
    enc = ["Nº", "Código", "participant_code", "scenario_id", "Situación (referencia)", "Intención esperada", "Frase escrita por la persona", "Decisión", "Intención correcta (si CAMBIAR)", "Comentario (opcional)", "Nota del sistema"]
    ws.append(enc)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="D9D9D9")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    amarillo = PatternFill("solid", fgColor="FFF2CC")
    for i, r in enumerate(val.sort_values(["participant_code", "scenario_id"]).itertuples(), 1):
        notas = []
        n = normalize(r.text, jerga)
        if n in por.index:
            notas.append(f"Idéntica (tras normalizar) a la frase de entrenamiento {por[n]}: por la regla fijada se excluye de la medición (coincidencia exacta, no el modelo); revisa su etiqueta igual.")
        if r.real_id in otro:
            notas.append("Posible otro idioma (marcada; no se descarta).")
        ws.append([i, r.real_id, r.participant_code, r.scenario_id, cat.loc[r.scenario_id, "situacion"], r.intent_esperada, r.text, None, None, None, " ".join(notas) or None])
        ws.cell(i + 1, 8).fill = amarillo
        ws.cell(i + 1, 9).fill = amarillo
    n = len(val)
    for col, w in zip("ABCDEFGHIJK", (5, 9, 14, 11, 50, 34, 60, 13, 34, 30, 60)):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2, max_row=n + 1):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    wi = wb.create_sheet("Intenciones")
    wi.append(["Intención"])
    for i in intents:
        wi.append([i])
    wi.column_dimensions["A"].width = 40
    dv1 = DataValidation(type="list", formula1='"OK,CAMBIAR,DESCARTAR"', allow_blank=True)
    dv2 = DataValidation(type="list", formula1=f"=Intenciones!$A$2:$A${len(intents) + 1}", allow_blank=True)
    ws.add_data_validation(dv1)
    ws.add_data_validation(dv2)
    dv1.add(f"H2:H{n + 1}")
    dv2.add(f"I2:I{n + 1}")
    ws0.cell(len(NOTAS) + 2, 1, "Pendientes (sin decisión):")
    ws0.cell(len(NOTAS) + 2, 2, f"=COUNTBLANK(Revision!H2:H{n + 1})")
    wb.calculation.fullCalcOnLoad = True
    salida.parent.mkdir(parents=True, exist_ok=True)
    wb.save(salida)
    print(f"Hoja de revisión creada: {salida} ({n} frases; sin sugerencias ni predicciones; no se imprimió ninguna frase). Intenciones: {len(intents)}.")


if __name__ == "__main__":
    main()
