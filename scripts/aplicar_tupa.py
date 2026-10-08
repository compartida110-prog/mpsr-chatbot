"""Aplica a domain.yml las respuestas corregidas de la hoja de verificación del TUPA (docs/tupa/Verificacion_TUPA_v*.xlsx, la versión más alta).

Entran SOLO las filas con Resultado = «Corregir», texto corregido escrito, «Confirmado por el tesista» = «Sí» y sin alerta. Las demás (Pendiente, «No figura en la fuente», Corregir sin confirmar) no se tocan.

  (por defecto)  DRY-RUN: no escribe domain.yml. Lista qué filas entrarían (intención, respuesta actual y nueva) y hace comprobaciones; escribe el informe logs/avance/aplicar_tupa_dryrun.md.
  --aplicar      Reemplaza el texto de utter_<intención> en domain.yml editando solo esas líneas (el resto del archivo queda byte a byte igual) y deja una copia domain.ANTES_DE_TUPA.yml.
                 Se NIEGA mientras no existan la partición v3 (corpus/v3_real/dataset_split_v3.csv) y la evaluación del test real (logs/v3_real/eval_real_resumen.json y test_registro.json):
                 las respuestas se aplican después de medir, no antes. Reentrenar y medir después de aplicar es un paso aparte.

Solo LEE el libro (nunca lo guarda).

Uso:
    python scripts/aplicar_tupa.py                      # dry-run
    python scripts/aplicar_tupa.py --libro docs/tupa/Verificacion_TUPA_v10.xlsx
    python scripts/aplicar_tupa.py --aplicar            # bloqueado hasta terminar partición y evaluación
"""
import argparse
import re
import shutil
import sys
import warnings
from datetime import datetime
from pathlib import Path

import yaml

from common import ROOT

warnings.filterwarnings("ignore")
from openpyxl import load_workbook  # noqa: E402

MARCA = re.compile(r"[Verificar[^]]*]")
DOMAIN = ROOT / "domain.yml"
INFORME = ROOT / "logs" / "avance" / "aplicar_tupa_dryrun.md"
REQUISITOS_APLICAR = [("la partición v3", "corpus/v3_real/dataset_split_v3.csv"), ("la evaluación F1 del test real", "logs/v3_real/eval_real_resumen.json"),
                      ("el registro de la evaluación del test", "logs/v3_real/test_registro.json")]


def libro_mas_alto():
    c = sorted((ROOT / "docs" / "tupa").glob("Verificacion_TUPA_v*.xlsx"), key=lambda f: tuple(int(x) for x in re.search(r"_v(\d+(?:_\d+)*)", f.name).group(1).split("_")))
    if not c:
        sys.exit("ERROR: no hay docs/tupa/Verificacion_TUPA_v*.xlsx")
    return c[-1]


def leer_hoja(ruta):
    wb = load_workbook(ruta, read_only=True, data_only=True)
    try:
        filas = list(wb["Verificacion"].iter_rows(values_only=True))
        resumen = {r[0]: (r[1] if len(r) > 1 else None) for r in wb["Resumen"].iter_rows(values_only=True) if r and isinstance(r[0], str)}
    finally:
        wb.close()
    fe = next((i for i, f in enumerate(filas[:8]) if f and len(f) > 1 and f[1] == "Intención"), None)
    if fe is None:
        sys.exit("ERROR: no encuentro el encabezado «Intención» en la hoja Verificacion.")
    enc = [str(c).strip() if c else "" for c in filas[fe]]

    def j(prefijo):
        k = next((i for i, c in enumerate(enc) if c.startswith(prefijo)), None)
        if k is None:
            sys.exit(f"ERROR: falta la columna «{prefijo}…».")
        return k
    ix = {n: j(p) for n, p in (("n", "#"), ("intent", "Intención"), ("prio", "Prioridad"), ("res", "Resultado"), ("texto", "Texto corregido"), ("alerta", "Alerta"), ("conf", "Confirmado"))}
    out = []
    for f in filas[fe + 1:]:
        if f and len(f) > max(ix.values()) and f[ix["intent"]]:
            out.append({k: (f[i] if f[i] is not None else "") for k, i in ix.items()})
    return out, resumen


def respuestas_domain(ruta):
    d = yaml.safe_load(open(ruta, encoding="utf-8"))
    return {k: v[0].get("text", "") for k, v in d.get("responses", {}).items() if v}


def un_renglon(t):
    return " ".join(str(t).split())


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--libro", default="")
    ap.add_argument("--domain", default=str(DOMAIN))
    ap.add_argument("--aplicar", action="store_true", help="escribe domain.yml (bloqueado hasta terminar la partición y la evaluación F1)")
    ap.add_argument("--informe", default=str(INFORME))
    ap.add_argument("--raiz", default=str(ROOT), help=argparse.SUPPRESS)  # solo para las pruebas: dónde buscar la partición y la evaluación
    a = ap.parse_args()
    libro = Path(a.libro) if a.libro else libro_mas_alto()
    filas, resumen = leer_hoja(libro)
    resp = respuestas_domain(a.domain)
    entran, fuera = [], {"Pendiente": 0, "No figura en la fuente": 0, "Corregir sin confirmar": 0, "Corregir con alerta o sin texto": 0, "Coincide": 0}
    problemas = []
    for r in filas:
        res = str(r["res"]).strip()
        if res == "Corregir":
            if str(r["conf"]).strip() != "Sí":
                fuera["Corregir sin confirmar"] += 1
            elif r["alerta"] or not str(r["texto"]).strip():
                fuera["Corregir con alerta o sin texto"] += 1
            else:
                entran.append(r)
        elif res in fuera:
            fuera[res] += 1
        else:
            fuera["Pendiente"] += 1
    # ---- comprobaciones sobre las filas que entrarían
    intents = [r["intent"] for r in entran]
    if len(set(intents)) != len(intents):
        problemas.append(f"intenciones repetidas entre las filas que entran: {sorted({i for i in intents if intents.count(i) > 1})}")
    filas_md = []
    for r in entran:
        clave = f"utter_{r['intent']}"
        nuevo, viejo = un_renglon(r["texto"]), un_renglon(resp.get(clave, ""))
        avisos = []
        if clave not in resp:
            problemas.append(f"{clave}: no existe en domain.yml")
            avisos.append("NO EXISTE en domain.yml")
        if nuevo == viejo:
            avisos.append("igual a la actual")
        if MARCA.search(nuevo):
            avisos.append("conserva una marca [Verificar …]")
        if re.search(r"\{[^}]*\}", nuevo):
            avisos.append("tiene llaves {…}")
        if len(nuevo) < 40:
            avisos.append("muy corto")
        if len(nuevo) > 900:
            avisos.append("muy largo")
        filas_md.append((r["n"], r["intent"], r["prio"], MARCA.search(viejo), len(viejo), len(nuevo), viejo, nuevo, avisos))
    # ---- salida
    L = [f"# Dry-run de aplicar_tupa.py — {libro.name}", "",
         f"Generado el {datetime.now():%Y-%m-%d %H:%M}. **No se escribió domain.yml.** Entran {len(entran)} de {len(filas)} filas (Resultado = Corregir, texto corregido, confirmada por el tesista, sin alerta).", "",
         "No entran: " + "; ".join(f"{k}: {v}" for k, v in fuera.items() if v) + ".", "",
         "Hoja Resumen guardada: " + "; ".join(f"{k} = {resumen.get(k)}" for k in ("Confirmadas por el tesista", "Prioridad Alta pendientes", "Filas con alerta", "Corregir o Coincide sin confirmar por el tesista")) +
         ". (Si difiere del recuento, la hoja se guardó antes de la última edición: ábrela y guárdala en Excel.)", "",
         "| # | Intención | Prioridad | Tenía [Verificar] | Largo actual → nuevo | Avisos |", "|---|---|---|---|---|---|"]
    for n, i, p, tv, lv, ln, v, nu, av in filas_md:
        L.append(f"| {n} | `{i}` | {p} | {'sí' if tv else 'no'} | {lv} → {ln} | {', '.join(av) or '—'} |")
    L += ["", "## Texto actual y nuevo de cada fila", ""]
    for n, i, p, tv, lv, ln, v, nu, av in filas_md:
        L += [f"### {n}. `utter_{i}`", "", f"- **Actual:** {v}", f"- **Nuevo:** {nu}", ""]
    if problemas:
        L += ["## Problemas (bloquean --aplicar)", ""] + [f"- {p}" for p in problemas]
    Path(a.informe).parent.mkdir(parents=True, exist_ok=True)
    Path(a.informe).write_text("\n".join(L) + "\n", encoding="utf-8")
    print(f"Libro: {libro.name} | filas: {len(filas)} | ENTRARÍAN: {len(entran)} | no entran: " + ", ".join(f"{k} {v}" for k, v in fuera.items() if v))
    for n, i, p, tv, lv, ln, v, nu, av in filas_md:
        print(f"  {n:>2} {i:<38} {p:<6} [Verificar] {'sí' if tv else 'no'}  {lv:>3}→{ln:<3} {', '.join(av)}")
    print(f"Problemas: {len(problemas)} {problemas[:3] if problemas else ''}\nInforme: {a.informe}")
    if not a.aplicar:
        print("DRY-RUN: domain.yml no se modificó.")
        return
    # ---------------------------------------------------------------- aplicar (bloqueado hasta medir)
    faltan = [f"{n} ({r})" for n, r in REQUISITOS_APLICAR if not (Path(a.raiz) / r).exists()]
    if faltan:
        sys.exit("ME NIEGO a aplicar: primero deben terminar la partición y la evaluación F1 del test real. Faltan: " + "; ".join(faltan) + ". Las respuestas se aplican después de medir.")
    if problemas:
        sys.exit("ME NIEGO a aplicar: hay problemas en las filas (ver el informe).")
    ruta = Path(a.domain)
    texto = ruta.read_text(encoding="utf-8")
    copia = ruta.with_name(ruta.stem + ".ANTES_DE_TUPA.yml")
    if not copia.exists():
        shutil.copyfile(ruta, copia)
    for r in entran:
        clave = f"utter_{r['intent']}"
        m = re.search(r"(^  %s:\n  - text: )(.*?)(?=\n  (?:utter_|-)|\n[a-z_]+:|\Z)" % re.escape(clave), texto, re.S | re.M)
        if not m:
            sys.exit(f"ERROR: no encuentro el bloque de {clave} (no se escribió nada).")
        nuevo = yaml.safe_dump(un_renglon(r["texto"]), allow_unicode=True, default_flow_style=True, width=10**6).rstrip("\n").removesuffix("\n...").removesuffix("...").rstrip()
        texto = texto[:m.start(2)] + nuevo + texto[m.end(2):]
    yaml.safe_load(texto)  # debe seguir siendo YAML válido
    ruta.write_text(texto, encoding="utf-8")
    print(f"domain.yml actualizado con {len(entran)} respuestas (copia previa: {copia.name}). Reentrenar y medir es el paso siguiente.")


if __name__ == "__main__":
    main()
