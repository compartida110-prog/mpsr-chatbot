# %% [markdown]
# ## Etapa 13 — Bitácora de incidencias (P16)
# **Qué se hace:** se lee `incident_log.csv` y se muestra como tabla: fecha, incidencia, decisión y justificación, más dos columnas **orientativas** (`tipo` técnica/metodológica y `afecta validez`).
# **Por qué:** el protocolo manda registrar **toda** desviación; ninguna se resuelve en silencio. La bitácora es la evidencia de que los cambios y errores quedaron declarados.
# **Aviso:** el archivo original **no** tiene las columnas «tipo» ni «afecta validez»; aquí se *infieren por palabras clave* solo para orientar la lectura. Son una heurística, no una clasificación del tesista.

# %%
inc = pd.read_csv(BASE / "incident_log.csv", dtype=str, keep_default_na=False, encoding="utf-8")
inc.columns = ["fecha", "incidencia", "descripción", "decisión", "justificación"]
PAL_TECNICA = r"bloque|smart app|dll|\.pyd|versi[oó]n de|dependencia|entorno|sklearn|scikit|tensorflow|rasa\.exe|windows|encoding|heredoc|script|error de|fix|ejecutable"
PAL_METODOL = r"fuga|test|evaluaci|umbral|lote|compuerta|congel|etiqueta|revisi|kappa|g[1-7]\b|muestra|dise[ñn]o|frase|protocolo|corpus|partici|desviaci"
PAL_VALIDEZ = r"fuga|evaluaci[oó]n|test|frases? (real|del lote)|umbral|congel|duplicad|g3|kappa|contamin|sesgo|sin revisar|no comparable|artefacto"
txt = (inc["incidencia"] + " " + inc["descripción"]).str.lower()
tec, met = txt.str.contains(PAL_TECNICA, regex=True), txt.str.contains(PAL_METODOL, regex=True)
inc["tipo (heurística)"] = np.where(tec & met, "mixta", np.where(tec, "técnica", np.where(met, "metodológica", "sin clasificar")))
inc["afecta validez (heurística)"] = np.where(txt.str.contains(PAL_VALIDEZ, regex=True), "posible: revisar", "no declarado")
mostrar = inc.assign(descripción=inc["descripción"].str.slice(0, 170), decisión=inc["decisión"].str.slice(0, 120))[["fecha", "tipo (heurística)", "afecta validez (heurística)", "incidencia", "descripción", "decisión"]]
rotulo("Bitácora del proyecto (registro de decisiones, no datos del estudio)", f"{len(inc)} incidencias")
print("Incidencias:", len(inc), "| por tipo (heurística):", inc["tipo (heurística)"].value_counts().to_dict(), "| con posible efecto en la validez:", int((inc["afecta validez (heurística)"] != "no declarado").sum()))
display(mostrar)
rotulo("Últimas 5 incidencias, con su justificación completa")
for _, r in inc.tail(5).iterrows():
    print(f"- {r['fecha']} · {r['incidencia']}\n    Decisión: {r['decisión'][:400]}\n    Justificación: {r['justificación'][:400]}")

# %% [markdown]
# ### Cómo leer el resultado (etapa 13)
# - Cada fila es una desviación o un error **declarado**: por ejemplo, que se mostró por descuido una frase del lote 2 en una salida (y no se usó para ajustar nada), que la evaluación del lote 2 tiene un artefacto en la sección «umbral», o que el registro del pre-piloto se rehízo porque se había mezclado con filas generadas por un asistente.
# - «posible: revisar» no quiere decir que la validez esté comprometida; marca las incidencias que tocan evaluación, umbral, congelamiento o fuga, para que el docente las lea con más cuidado.

# %% [markdown]
# ## Etapa 14 — Matriz de trazabilidad
# **Qué se hace:** una tabla única que enlaza cada **paso (P01–P16)** y cada **compuerta (G1–G7)** con el script que lo ejecuta, el archivo de evidencia, el resultado y su estado (*ejecutado / simulado / planificado / pendiente*).
# **Por qué:** permite al docente seguir un hilo desde el protocolo hasta el dato. El estado de las compuertas se toma **del tablero**, no se escribe a mano.

# %%
tp = pd.DataFrame(meta["traza_pasos"], columns=["id", "paso", "script(s)", "archivo de evidencia", "resultado", "estado"])
tc = []
for gid, nombre, scr, evid in meta["traza_compuertas"]:
    c = next(x for x in tb["compuertas"] if x["id"] == gid)
    estado = {"Cumplida": "ejecutado", "No cumplida": "ejecutado (no cumplida)"}.get(c["estado"], "pendiente")
    if gid == "G7":
        estado = "pendiente (solo demostración simulada)"
    tc.append([gid, f"Compuerta: {nombre}", scr, evid, c["nota"][:230], estado])
tc = pd.DataFrame(tc, columns=tp.columns)


def en_paquete(ev):
    ruta = ev.split(",")[0].split(" + ")[0].strip()
    ruta = ruta.replace("…", "").replace("*", "")
    return "sí" if (BASE / ruta).exists() else "no (privado o ignorado por Git)"


matriz = pd.concat([tp, tc], ignore_index=True)
matriz["en este paquete"] = matriz["archivo de evidencia"].map(en_paquete)
rotulo("Trazabilidad del proyecto (el rótulo de origen de cada dato está en las etapas 3 a 12)")
print("Estados:", matriz["estado"].value_counts().to_dict())
display(matriz)

# %% [markdown]
# ### Cómo leer el resultado (etapa 14)
# - **ejecutado:** el paso se hizo con datos del proyecto. **simulado:** solo se demostró con datos inventados (línea base P01, análisis P14). **pendiente:** depende de datos que aún no existen (piloto de 60 sesiones, G7).
# - «en este paquete = no» significa que la evidencia vive en una carpeta ignorada por Git por traer datos de personas; los números agregados de esa evidencia están en las etapas 4, 7 y 10.

# %% [markdown]
# ## Etapa 15 — Demostración del flujo completo con datos SIMULADOS
# **Qué se hace:** se recorre de punta a punta la demostración simulada (`demostracion_simulada.py`, ejecutada el 5 de octubre de 2026): ingesta → partición → evaluación → congelamiento → sesiones → tablero. Todo con libros inventados: `Lote1_Transcripcion_SIMULADO_v2.xlsx` y `Registro_Sesiones_Piloto_SIMULADO_v4.xlsx`.
# **Por qué:** muestra que el *procedimiento* funciona (y que el código detecta lo que debe) sin gastar ni tocar datos reales. **Nada de esto es un hallazgo.**
# **Rótulo:** todo lo de esta etapa es **Simulado (demostración)** y nunca se suma a las cifras reales.

# %%
DEMO = BASE / "evidencias/simulado_demostracion/20261005_demostracion"
rotulo(ORIGEN_SIMULADO, "1. Ingesta (libro de transcripción simulado)")
lf = pd.read_csv(DEMO / "01_ingesta/lote1_real_final_SIMULADO.csv", dtype=str, keep_default_na=False, encoding="utf-8")
print("Frases del lote simulado tras la revisión automática:", len(lf), "| intenciones:", lf["intent"].nunique() if "intent" in lf else "—")

rotulo(ORIGEN_SIMULADO, "2. Partición v3 con el lote simulado")
sp = pd.read_csv(DEMO / "02_particion/dataset_split_v3_SIMULADO.csv", dtype=str, keep_default_na=False, encoding="utf-8")
print("Partición:", sp["split"].value_counts().to_dict())
fuga = int((sp[sp["split"].isin(["train", "validation", "test"])].groupby("base_phrase_id")["split"].nunique() > 1).sum())
print("Grupos de paráfrasis en más de una partición:", fuga)

rotulo(ORIGEN_SIMULADO, "3. Evaluación (mejor configuración ya elegida; una evaluación del test)")
er = json.loads((DEMO / "03_evaluacion/eval_real_resumen_SIMULADO.json").read_text(encoding="utf-8"))
display(pd.DataFrame([{"método": m, "F1 macro (promedio de semillas)": round(v["f1_macro"][0], 4), "IC inferior": round(v["f1_macro"][1], 4), "IC superior": round(v["f1_macro"][2], 4)} for m, v in er["metodos"].items()]))

rotulo(ORIGEN_SIMULADO, "4. «Congelamiento» de demostración (un marcador con otro nombre, no el modelo real)")
fzs = json.loads((DEMO / "04_congelado/modelo_congelado_SIMULADO.json").read_text(encoding="utf-8"))
mod_sim = DEMO / "04_congelado/modelo_demostracion_SIMULADO.tar.gz"
print("Estado declarado:", fzs["estado"])
print("sha256 del marcador coincide con el congelamiento simulado:", sha256(mod_sim) == fzs["sha256"]["modelo"])

rotulo(ORIGEN_SIMULADO, "5. Sesiones simuladas (60) y alfa de Cronbach")
an = json.loads((DEMO / "05_sesiones/analisis_piloto_SIMULADO.json").read_text(encoding="utf-8"))
print("Sesiones simuladas elegibles:", an["n"], "| alfa guardado de la demostración:", round(an["satisfaccion"]["alfa_cronbach"], 4))
from openpyxl import load_workbook
wbs = load_workbook(BASE / "docs/piloto/ejemplos_simulados/Registro_Sesiones_Piloto_SIMULADO_v4.xlsx", read_only=True, data_only=True)
filas_s = list(wbs["Sesiones"].iter_rows(values_only=True))
wbs.close()
enc_s = next(i for i, r in enumerate(filas_s[:10]) if r and str(r[0]).startswith("C") and "sesi" in str(r[0]))
cab = [str(c) if c else "" for c in filas_s[enc_s]]
j_el = cab.index("Elegible y completa")
j_it = [next(j for j, c in enumerate(cab) if c.startswith(f"Post ítem {i}:")) for i in range(1, 9)]
datos_s = [r for r in filas_s[enc_s + 1:] if r and r[0] and re.match(r"^(SA|PP)\d", str(r[0])) and str(r[j_el]) == "Sí"]
Xs = np.array([[float(r[j]) for j in j_it] for r in datos_s])
a_sim, _, _ = alfa_cronbach(Xs)
print(f"Alfa recalculado desde el registro simulado: {a_sim:.4f} (n = {len(Xs)}); la demostración NO es un resultado del pre-piloto ni del piloto.")
comparar("15", "alfa de la demostración simulada (Simulado)", a_sim, an["satisfaccion"]["alfa_cronbach"], "evidencias/simulado_demostracion/…/05_sesiones", tol=1e-6)

rotulo(ORIGEN_SIMULADO, "6. Tablero de la demostración")
print("Compuertas «Cumplida (Simulado)» en la demostración:", [c["id"] for c in ts["compuertas"] if "Simulado" in c["estado"]], "| ninguna cuenta como real")
inf = (DEMO / "INFORME_DEMOSTRACION_SIMULADA.md").read_text(encoding="utf-8")
print("\nAdvertencias del informe de la demostración (resumen):")
for linea in re.findall(r"^\d\. \*\*.*$", inf, re.M)[:4]:
    print(" -", linea[:260])
mostrar_comparaciones(etapas=["15"], titulo="Etapa 15 — recalculado frente a reportado (Simulado)")

# %% [markdown]
# ### Cómo leer el resultado (etapa 15)
# - La demostración **no mide nada del estudio**: las frases las escribió quien armó el libro de prueba (varias son idénticas al corpus sintético), las «consultas» de las sesiones son el texto de las tarjetas, y la prueba final usa un **predictor simulado**.
# - Sirve para comprobar que cada eslabón (ingesta → partición → evaluación → congelamiento → sesiones → tablero) corre y deja huellas. El alfa de la demostración (≈ 0,55) **no** tiene relación con el 0,366 del pre-piloto real.

# %% [markdown]
# ## Etapa 16 — Limitaciones y decisiones del protocolo (V1.6 a V1.8)
# **Qué se hace:** se listan las limitaciones declaradas y las decisiones metodológicas vigentes.
# **Nota sobre versiones:** el repositorio contiene hasta la **V1.8** (Nota v11). No hay una V1.9 instalada; si el docente trabaja con una V1.9, deberá contrastar estas cifras con ella (ver «Discrepancias» en la etapa 17).

# %%
rotulo("Decisiones y limitaciones (texto del protocolo y de la Nota; no son datos del estudio)")
decisiones = [
    ("V1.6 (sección 5.8)", "G3 se mide con una prueba INDEPENDIENTE (lote 2: solo test, participantes P33–P57), con el procedimiento cerrado ANTES de recoger los datos. La medición del lote 1 queda como «No cumplida, lote 1» (F1 0,687) y no se borra."),
    ("V1.6", "Refinamiento previo al lote 2: máximo 2 ciclos, solo con datos del lote 1, validación cruzada por participante, parar si la mejora es < 0,02. Se hizo 1 de 2 ciclos. Entrenamiento: 707 sintéticas + 185 reales activas del lote 1 + 51 sintéticas nuevas = 943."),
    ("V1.6", "Las frases del lote 2 idénticas (tras normalizar) a frases del entrenamiento se EXCLUYEN del test por coincidencia exacta, nunca por lo que el modelo prediga: 13 de 280."),
    ("V1.6", "Umbral de confianza t = 0,50 (y ambigüedad 0,1), elegido con la validación cruzada por participante del lote 1 y congelado ANTES de abrir el lote 2."),
    ("V1.7", "La evaluación del test del lote 2 es ÚNICA (el registro se escribe antes de predecir). Se congela el modelo medido (LOTE2-FINAL v1) sin reentrenar: G5 se evalúa por coincidencia de huellas y fecha posterior a G3 y G4."),
    ("V1.8", "Compuertas reales: 5 de 7 (G1–G5 cumplidas). G6 y G7 siguen pendientes. Las compuertas simuladas nunca cuentan como reales."),
    ("Pre-piloto", "El pre-piloto usa el asistente local con el modelo congelado; su registro NO se pasa por analizar_piloto.py (gastaría la prueba final única del modelo). G6 lee solo Registro_Sesiones_Prepiloto.xlsx."),
]
display(pd.DataFrame(decisiones, columns=["versión", "decisión"]))
limitaciones = [
    "Las 267 frases del test vienen de las MISMAS 56 situaciones del lote 1: el F1 0,907 es una medición independiente en participantes, no evidencia de generalización a situaciones nuevas.",
    "Una sola revisora de etiquetas: el «kappa» (0,935–0,938 en el lote 1; 1,000 en el lote 2) compara la etiqueta esperada con la revisión de UNA persona; no es acuerdo entre revisores.",
    "Solo 4–5 frases por intención en el test (mínimo de textos distintos por intención: 4): los IC por intención son muy inestables.",
    "El F1 oficial (0,9072) promedia 55 etiquetas y cuenta las abstenciones como error; sobre las 54 intenciones sería 0,9240. La sección «umbral» impresa por eval_lote2.py (98,9 % / 89,8 %) es un artefacto declarado; las cifras oficiales son cobertura 92,9 % y precisión 95,6 %.",
    "El 0,907 es más alto que la validación cruzada por participante del lote 1 (0,787–0,799); el 0,687 del lote 1 y estas cifras no son comparables entre sí (otros datos y otro procedimiento).",
    "Muestra de conveniencia: la MPSR no autorizó el despliegue ni el acceso al local; el piloto exploratorio (n = 60) no es confirmatorio ni generalizable.",
    "Pre-piloto: n = 5 elegibles; el alfa de Cronbach (0,366) es muy inestable.",
    "Punto abierto del protocolo: la sección 2.12 aún pide «F1 ≥ 0,75 sobre el conjunto real retenido del lote 1», que ya no es retenido (el modelo se entrenó con esas frases); la sección 5.8 debe actualizarse.",
    "Incidente: se imprimió por descuido una sola frase del lote 2 (P33, S01) en una salida de consola; está registrada en incident_log.csv y no se usó para ningún ajuste.",
]
rotulo("Limitaciones declaradas")
for i, l in enumerate(limitaciones, 1):
    print(f"{i}. {l}")

# %% [markdown]
# ### Cómo leer el resultado (etapa 16)
# Las decisiones explican *por qué* el diseño es creíble (procedimiento cerrado antes de medir, evaluación única, congelamiento por huellas). Las limitaciones explican *hasta dónde* se puede afirmar: el chatbot funciona con las personas del lote 2 en las 56 situaciones probadas, pero eso no demuestra que funcione igual con situaciones nuevas.

# %% [markdown]
# ## Etapa 17 — Cierre
# **Qué se hace:** se resume qué se ejecutó *en vivo*, qué usó *resultados guardados* y qué *requiere el paquete privado*; se imprime la tabla final «recalculado frente a reportado» y la lista de discrepancias.
# **Regla:** si algo recalculado no coincide con lo reportado, aparece aquí como discrepancia. No se corrige en silencio.

# %%
estado_etapas = [
    ("1 Entorno", "en vivo", "—"), ("2 Inventario del código", "en vivo", "—"), ("3 Diccionario de datos", "en vivo (público) + guardado (Real)", "estructura de archivos reales: guardada"),
    ("4 Partición", "en vivo (sintético)" + (" + recálculo real" if HAY_PRIVADO else ""), "recálculo real: " + ("hecho" if HAY_PRIVADO else "requiere paquete privado")),
    ("5 Línea base / entrenamiento", "en vivo (sintético; DIET solo con Rasa)", "—"),
    ("6 Congelamiento", "en vivo", "modelo y entrenamiento: " + ("verificados" if HAY_PRIVADO else "requieren paquete privado")),
    ("7 Resultados del lote 2", "recalculado desde predicciones guardadas" if HAY_PRIVADO else "valores guardados", "predicciones: " + ("usadas" if HAY_PRIVADO else "requiere paquete privado")),
    ("8 Glosario", "texto", "—"), ("9 Respuestas / TUPA", "en vivo", "—"),
    ("10 Pre-piloto", "recalculado desde registro anonimizado" if HAY_PRIVADO else "alfa desde varianzas agregadas guardadas", "registro: " + ("usado" if HAY_PRIVADO else "requiere paquete privado")),
    ("11 Pruebas automáticas", "resultados guardados", "—"), ("12 Tablero", "resultado guardado", "—"), ("13 Bitácora", "en vivo", "—"), ("14 Trazabilidad", "en vivo", "—"),
    ("15 Demostración simulada", "en vivo + resultados guardados (Simulado)", "—"), ("16 Limitaciones", "texto", "—"),
]
display(pd.DataFrame(estado_etapas, columns=["etapa", "cómo se obtiene", "dependencia"]))
print(f"Rasa disponible: {HAY_RASA} · paquete privado: {HAY_PRIVADO}")
mostrar_comparaciones(titulo="TABLA FINAL — recalculado frente a reportado (todas las etapas)")
print()
print("DISCREPANCIAS (recalculado ≠ reportado):", DISCREPANCIAS or "ninguna")
print("Diferencias CONOCIDAS y ya declaradas:")
for d in DECLARADAS:
    print(" -", d)
Path(BASE / "resultados_cuaderno").mkdir(exist_ok=True)
(BASE / "resultados_cuaderno/comparaciones.json").write_text(json.dumps({"comparaciones": COMPARACIONES, "discrepancias": DISCREPANCIAS, "declaradas": DECLARADAS, "hay_privado": HAY_PRIVADO, "hay_rasa": HAY_RASA}, ensure_ascii=False, indent=1, default=str), encoding="utf-8")

# %% [markdown]
# ### Cómo leer el resultado (etapa 17)
# - La tabla final reúne todas las cifras que el cuaderno recalcula: F1 0,9072, cobertura 92,9 %, precisión 95,6 %, 19 abstenciones, 943 frases de entrenamiento, alfa 0,366, 5 de 7 compuertas y las huellas del congelamiento.
# - **Sin Rasa y con solo el paquete público** el cuaderno corre completo; lo que depende del privado muestra el valor guardado y lo dice.
# - **Discrepancias:** si aparece alguna, es un hallazgo a investigar, no un error del cuaderno que haya que silenciar.
# - **Lo que este cuaderno NO hace:** no reevalúa el test del lote 2, no ejecuta `analizar_piloto.py`, no entrena ni reemplaza el modelo congelado, no imprime frases ni respuestas individuales.
