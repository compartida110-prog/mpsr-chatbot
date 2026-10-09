# %% [markdown]
# ## Etapa 7 — Resultados guardados del lote 2 (G3)
# **Qué se hace:** a partir de las **predicciones ya guardadas** del modelo congelado (y del SVM) sobre las 267 frases del test del lote 2, se **recalculan** F1 macro, exactitud, cobertura, precisión de lo respondido, abstenciones, los IC bootstrap (por frases y por participantes), el F1 por intención, las intenciones más débiles, la confusión agregada y la prueba de McNemar frente al SVM. Cada cifra se contrasta con la reportada.
# **Por qué:** la evaluación única del lote 2 **ya se gastó** (la sección 5.8 del protocolo prohíbe repetirla). Recalcular *desde las predicciones guardadas* permite que cualquiera verifique las cifras **sin volver a evaluar** y sin tocar el modelo.
# **Paso del protocolo:** P11 y compuerta **G3** (criterio F1 macro ≥ 0,75, sin cambios).
# **Qué NO hace:** no carga ningún modelo, no vuelve a predecir, no cambia el umbral.

# %%
def mostrar_comparaciones(etapas=None, titulo="Recalculado frente a reportado"):
    filas = [c for c in COMPARACIONES if etapas is None or c[0] in etapas]
    t = pd.DataFrame(filas, columns=["etapa", "cifra", "recalculado", "reportado", "fuente de lo reportado", "coincide"])
    print(f"{titulo}: {len(t)} cifras · coinciden: {(t['coincide'] == 'Sí').sum()} · NO coinciden: {(t['coincide'] == 'NO').sum()}")
    display(t)


res = leer_json("logs/avance/eval_lote2_resumen.json")
rasa_rep = res["metodos"]["rasa"]
svm_rep = res["metodos"]["svm"]
REP_PROTOCOLO = {"f1_macro": 0.9072, "cobertura": 0.929, "precision_respondida": 0.956, "abstenciones": 19, "n": 267, "participantes": 25, "exactitud": 0.8876}
DECLARADAS = []   # diferencias CONOCIDAS y ya declaradas en los informes; se muestran, no se corrigen

# cifras que el propio script de evaluación imprimió con el umbral (artefacto declarado)
u = res["umbral_congelado"]
rotulo(ORIGEN_REAL, "lo que imprimió eval_lote2.py en la sección «umbral» (artefacto declarado, NO es la cifra oficial)")
print(f"Sección «umbral» del script: respondidas {u['respondidas']}, abstenciones {u['abstenciones']}, cobertura {u['cobertura']:.1%}, precisión de lo respondido {u['precision_respondida']:.1%}.")
print("Esas cifras cuentan como 'respondidas' las 19 frases que el FallbackClassifier del pipeline ya había convertido en «nlu_fallback»: están declaradas como artefacto en logs/avance/eval_lote2_notas.md.")
DECLARADAS.append("eval_lote2_resumen.json: cobertura 98,9 % y precisión 89,8 % son un artefacto declarado del script (las 19 abstenciones 'nlu_fallback' se contaron como respuestas); las cifras oficiales son 92,9 % y 95,6 %.")

# %%
from scipy.stats import binomtest
from sklearn.metrics import f1_score, precision_recall_fscore_support

N_BOOT, SEMILLA = 1000, 42


def f1m(y, p):
    return float(f1_score(np.asarray(y), np.asarray(p), average="macro", zero_division=0))


def ic_bootstrap(y, p, grupos, rng, n=N_BOOT, por_participante=False):
    """(punto, límite inferior, límite superior, media del bootstrap) — misma lógica que scripts/eval_lote2.py"""
    y, p = np.asarray(y), np.asarray(p)
    if por_participante:
        idx = {g: np.where(grupos == g)[0] for g in np.unique(grupos)}
        gs = list(idx)
        vals = [f1m(y[ix], p[ix]) for ix in (np.concatenate([idx[g] for g in rng.choice(gs, len(gs), replace=True)]) for _ in range(n))]
    else:
        vals = [f1m(y[ix], p[ix]) for ix in (rng.integers(0, len(y), len(y)) for _ in range(n))]
    lo, hi = np.percentile(vals, [2.5, 97.5])
    return f1m(y, p), float(lo), float(hi), float(np.mean(vals))


if HAY_PRIVADO:
    pred = pd.read_csv(PRIV / "predicciones_lote2_sin_texto.csv", dtype={"confidence": float, "confidence_2": float}, keep_default_na=False, encoding="utf-8")
    y, p, g, ps = pred["intent"].values, pred["predicted"].values, pred["participant_code"].values, pred["predicted_svm"].values
    n = len(y)
    abst = int((p == "nlu_fallback").sum())
    aciertos = int((y == p).sum())
    metricas = {"n": n, "participantes": len(set(g)), "etiquetas en el promedio (esperadas ∪ predichas)": len(set(y) | set(p)),
                "F1 macro": f1m(y, p), "exactitud": aciertos / n, "abstenciones (nlu_fallback)": abst, "cobertura": (n - abst) / n, "precisión de lo respondido": aciertos / (n - abst)}
    f1_54 = f1_score(y, p, labels=sorted(set(y)), average="macro", zero_division=0)
    rotulo(ORIGEN_REAL, f"recalculado desde predicciones guardadas de {n} frases (sin texto)")
    display(pd.DataFrame({"métrica": list(metricas), "recalculado": [round(v, 4) if isinstance(v, float) else v for v in metricas.values()]}))
    print(f"F1 macro promediando solo las 54 intenciones esperadas (sin la etiqueta de abstención): {f1_54:.4f} — declarado en la Nota; el oficial (0,9072) promedia 55 etiquetas y cuenta las abstenciones como error.")
    comparar("7", "F1 macro (DIET)", metricas["F1 macro"], REP_PROTOCOLO["f1_macro"], "Protocolo V1.8 / eval_lote2_resumen.json")
    comparar("7", "F1 macro (DIET) frente al JSON del script", metricas["F1 macro"], rasa_rep["f1_macro"][0], "logs/avance/eval_lote2_resumen.json", tol=1e-9)
    comparar("7", "exactitud", metricas["exactitud"], rasa_rep["accuracy"], "logs/avance/eval_lote2_resumen.json", tol=1e-9)
    comparar("7", "frases de test", n, REP_PROTOCOLO["n"], "Protocolo V1.8", tol=0)
    comparar("7", "participantes", metricas["participantes"], REP_PROTOCOLO["participantes"], "Protocolo V1.8", tol=0)
    comparar("7", "abstenciones (nlu_fallback)", abst, REP_PROTOCOLO["abstenciones"], "Protocolo V1.8 / Nota v11", tol=0)
    comparar("7", "cobertura", metricas["cobertura"], REP_PROTOCOLO["cobertura"], "Protocolo V1.8 (92,9 %)", tol=5e-4)
    comparar("7", "precisión de lo respondido", metricas["precisión de lo respondido"], REP_PROTOCOLO["precision_respondida"], "Protocolo V1.8 (95,6 %)", tol=5e-4)
else:
    requiere_privado("predicciones del lote 2 sin texto")
    rotulo(ORIGEN_REAL, "valores GUARDADOS (no recalculados)")
    display(pd.DataFrame({"métrica": ["F1 macro", "exactitud", "abstenciones (nlu_fallback)", "cobertura", "precisión de lo respondido"],
                          "valor guardado": [rasa_rep["f1_macro"][0], rasa_rep["accuracy"], REP_PROTOCOLO["abstenciones"], REP_PROTOCOLO["cobertura"], REP_PROTOCOLO["precision_respondida"]]}))

# %%
# ----- Intervalos de confianza bootstrap (1000 remuestreos, semilla 42) -----
if HAY_PRIVADO:
    f1_f = ic_bootstrap(y, p, g, np.random.default_rng(SEMILLA))
    f1_p = ic_bootstrap(y, p, g, np.random.default_rng(SEMILLA), por_participante=True)
    rotulo(ORIGEN_REAL, "IC95 % del F1 macro (DIET) recalculados")
    display(pd.DataFrame([{"remuestreo": "frases", "F1": f1_f[0], "IC95 % inferior": f1_f[1], "IC95 % superior": f1_f[2], "media del bootstrap": f1_f[3]},
                          {"remuestreo": "participantes", "F1": f1_p[0], "IC95 % inferior": f1_p[1], "IC95 % superior": f1_p[2], "media del bootstrap": f1_p[3]}]).round(4))
    comparar("7", "IC95 % por frases (inferior)", f1_f[1], rasa_rep["f1_macro"][1], "logs/avance/eval_lote2_resumen.json")
    comparar("7", "IC95 % por frases (superior)", f1_f[2], rasa_rep["f1_macro"][2], "logs/avance/eval_lote2_resumen.json")
    comparar("7", "IC95 % por participantes (inferior)", f1_p[1], rasa_rep["f1_macro_ic_participantes"][1], "logs/avance/eval_lote2_resumen.json")
    comparar("7", "IC95 % por participantes (superior)", f1_p[2], rasa_rep["f1_macro_ic_participantes"][2], "logs/avance/eval_lote2_resumen.json")
    comparar("7", "media del bootstrap por frases", f1_f[3], res["bootstrap"]["media_por_frases"], "logs/avance/eval_lote2_resumen.json")
    comparar("7", "media del bootstrap por participantes", f1_p[3], res["bootstrap"]["media_por_participantes"], "logs/avance/eval_lote2_resumen.json")
    sobre_umbral_frases = f1_f[1] >= 0.75
    sobre_umbral_part = f1_p[1] >= 0.75
    print(f"¿F1 macro (punto) ≥ 0,75? {f1_f[0] >= 0.75} | ¿límite inferior IC por frases ≥ 0,75? {sobre_umbral_frases} | ¿límite inferior IC por participantes ≥ 0,75? {sobre_umbral_part}")
else:
    requiere_privado("remuestreo bootstrap")
    rotulo(ORIGEN_REAL, "IC95 % GUARDADOS")
    display(pd.DataFrame([{"remuestreo": "frases", "F1": rasa_rep["f1_macro"][0], "inferior": rasa_rep["f1_macro"][1], "superior": rasa_rep["f1_macro"][2]},
                          {"remuestreo": "participantes", "F1": rasa_rep["f1_macro_ic_participantes"][0], "inferior": rasa_rep["f1_macro_ic_participantes"][1], "superior": rasa_rep["f1_macro_ic_participantes"][2]}]).round(4))

# %%
# ----- F1 por intención, las tres más débiles y confusiones agregadas -----
f1_guardado = pd.read_csv(BASE / "logs/avance/eval_lote2_f1_por_intencion.csv")
if HAY_PRIVADO:
    intenciones = sorted(set(y))
    pre, rec, f1i, sup = precision_recall_fscore_support(y, p, labels=intenciones, zero_division=0)
    por_int = pd.DataFrame({"intent": intenciones, "n": sup, "F1": f1i.round(4)})
    fusion = por_int.merge(f1_guardado[["intent", "f1"]].rename(columns={"f1": "F1 guardado"}), on="intent")
    dif_max = float((fusion["F1"] - fusion["F1 guardado"]).abs().max())
    comparar("7", "F1 por intención (máxima diferencia absoluta, 54 intenciones)", dif_max, 0.0, "logs/avance/eval_lote2_f1_por_intencion.csv", tol=5e-4)
    rotulo(ORIGEN_REAL, "F1 por intención recalculado (n = frases de test de esa intención)")
    display(por_int.sort_values("F1").reset_index(drop=True))
    debiles = por_int.sort_values(["F1", "intent"]).head(3)
    debiles_guard = f1_guardado.sort_values(["f1", "intent"]).head(3)
    print("Las tres intenciones más débiles:", list(debiles["intent"]), "| F1:", list(debiles["F1"]))
    comparar("7", "tres intenciones más débiles", ",".join(debiles["intent"]), ",".join(debiles_guard["intent"]), "logs/avance/eval_lote2_f1_por_intencion.csv", tipo="txt")
    # confusiones agregadas (solo etiquetas)
    err = pd.DataFrame({"esperada": y, "predicha": p})
    err = err[err["esperada"] != err["predicha"]]
    top = err.groupby(["esperada", "predicha"]).size().rename("veces").reset_index().sort_values("veces", ascending=False).head(12)
    rotulo(ORIGEN_REAL, "confusión agregada: pares (esperada → predicha) más frecuentes; incluye «nlu_fallback» = abstención")
    display(top.reset_index(drop=True))
    print("Errores totales:", len(err), "de los cuales abstenciones (nlu_fallback):", int((err["predicha"] == "nlu_fallback").sum()), "| confusiones entre intenciones reales:", int((err["predicha"] != "nlu_fallback").sum()))
    cd = res["confusion_par_despedida_agradecimiento"]
    m_dep = {"n": int((y == "despedida").sum()), "predicha_despedida": int(((y == "despedida") & (p == "despedida")).sum()), "predicha_agradecimiento": int(((y == "despedida") & (p == "agradecimiento")).sum())}
    m_agr = {"n": int((y == "agradecimiento").sum()), "predicha_despedida": int(((y == "agradecimiento") & (p == "despedida")).sum()), "predicha_agradecimiento": int(((y == "agradecimiento") & (p == "agradecimiento")).sum())}
    comparar("7", "par despedida↔agradecimiento", json.dumps([m_dep, m_agr], sort_keys=True), json.dumps([{k: cd["despedida"][k] for k in m_dep}, {k: cd["agradecimiento"][k] for k in m_agr}], sort_keys=True), "logs/avance/eval_lote2_resumen.json", tipo="txt")
    print("Despedida:", m_dep, "| Agradecimiento:", m_agr, "(confusión esperable: «gracias» se usa también para despedirse)")
else:
    requiere_privado("F1 por intención y confusiones")
    rotulo(ORIGEN_REAL, "F1 por intención GUARDADO (las 10 más bajas)")
    display(f1_guardado.sort_values("f1").head(10)[["intent", "n", "f1"]])

# %%
# ----- SVM como referencia y McNemar exacto (informativo: G3 se mide solo con DIET) -----
if HAY_PRIVADO:
    f1_svm_f = ic_bootstrap(y, ps, g, np.random.default_rng(SEMILLA))
    f1_svm_p = ic_bootstrap(y, ps, g, np.random.default_rng(SEMILLA), por_participante=True)
    ok_svm, ok_dt = (ps == y), (p == y)
    solo_svm, solo_diet = int((ok_svm & ~ok_dt).sum()), int((~ok_svm & ok_dt).sum())
    p_mc = float(binomtest(min(solo_svm, solo_diet), solo_svm + solo_diet, 0.5).pvalue) if solo_svm + solo_diet else 1.0
    rotulo(ORIGEN_REAL, "DIET frente a SVM (mismas 267 frases)")
    display(pd.DataFrame([{"modelo": "DIET (congelado)", "F1 macro": f1m(y, p), "exactitud": float((p == y).mean())},
                          {"modelo": "SVM (baseline informativo)", "F1 macro": f1_svm_f[0], "exactitud": float((ps == y).mean())}]).round(4))
    print(f"Solo SVM acierta: {solo_svm} · solo DIET acierta: {solo_diet} · p de McNemar exacto = {p_mc:.3f} → no hay evidencia de que uno sea mejor (p > 0,05).")
    comparar("7", "F1 macro del SVM", f1_svm_f[0], svm_rep["f1_macro"][0], "logs/avance/eval_lote2_resumen.json", tol=1e-9)
    comparar("7", "McNemar: solo SVM acierta", solo_svm, svm_rep["comparacion_con_rasa"]["solo_svm_acierta"], "logs/avance/eval_lote2_resumen.json", tol=0)
    comparar("7", "McNemar: solo DIET acierta", solo_diet, svm_rep["comparacion_con_rasa"]["solo_rasa_acierta"], "logs/avance/eval_lote2_resumen.json", tol=0)
    comparar("7", "McNemar: p exacto", p_mc, svm_rep["comparacion_con_rasa"]["p_mcnemar_exacto"], "logs/avance/eval_lote2_resumen.json", tol=1e-9)
else:
    requiere_privado("McNemar frente al SVM")
    cr = svm_rep["comparacion_con_rasa"]
    rotulo(ORIGEN_REAL, "valores GUARDADOS")
    print(f"SVM F1 macro {svm_rep['f1_macro'][0]:.4f} · solo SVM acierta {cr['solo_svm_acierta']} · solo DIET acierta {cr['solo_rasa_acierta']} · p de McNemar exacto {cr['p_mcnemar_exacto']:.3f}")

# %%
mostrar_comparaciones(etapas=["7"], titulo="Etapa 7 — recalculado frente a reportado")
print()
print("Diferencias CONOCIDAS y ya declaradas (se muestran, no se corrigen):")
for d in DECLARADAS:
    print(" -", d)

# %% [markdown]
# ### Cómo leer el resultado (etapa 7)
# - **F1 macro = 0,9072** con 267 frases de 25 participantes: supera el criterio 0,75; los límites inferiores de los IC95 % (por frases ≈ 0,86 y por participantes ≈ 0,85) también lo superan. **Pero** el lote 2 usa las mismas 56 situaciones del lote 1, una sola revisora de etiquetas y 4–5 frases por intención: es una medición *independiente en participantes*, **no** prueba de generalización a situaciones nuevas. La validación cruzada del lote 1 dio 0,787–0,799, y el 0,907 es más alto que eso.
# - **Cobertura 92,9 % y precisión 95,6 %:** de 267 frases, 19 son abstenciones («nlu_fallback») y 248 son respuestas; 237 de ellas son correctas (237/248 = 95,6 %). La exactitud 0,8876 = 237/267 cuenta las abstenciones como error.
# - **Dos advertencias de lectura:** (1) el F1 oficial promedia **55 etiquetas** (las 54 intenciones + «nlu_fallback» como predicción, que vale F1 = 0); sobre las 54 intenciones sería 0,9240. (2) La sección «umbral» impresa por el script (98,9 % / 89,8 %) es un **artefacto declarado**, no la cifra oficial.
# - **Intenciones débiles** (`licencia_funcionamiento_plazo`, `fuera_de_alcance`, `denuncia_seguridad_ciudadana`): son las que más se confunden entre sí o se abstienen; `fuera_de_alcance` es una clase «cajón» con 15 frases.
# - **DIET vs SVM:** diferencia de 0,015 en F1 y p de McNemar = 0,541: *no* hay evidencia de que DIET supere al SVM en este test.
# - En la tabla «recalculado frente a reportado», cualquier «NO» sería una **discrepancia** que se reporta tal cual (no se corrige en silencio).

# %% [markdown]
# ## Etapa 8 — Métricas: glosario
# **Qué se hace:** se definen, con fórmula, interpretación y limitación, todas las métricas usadas en el cuaderno y en el protocolo.
# **Por qué:** una cifra sin su limitación se malinterpreta (por ejemplo, un alfa alto con n = 5, o un F1 sin decir que cuenta abstenciones como error).

# %%
rotulo("Definiciones (no son datos del estudio)")
display(pd.DataFrame(meta["glosario"]))

# %% [markdown]
# ### Cómo leer el resultado (etapa 8)
# Úsalo como consulta rápida. Las limitaciones de la última columna son las que se declaran en la Nota de Desviación y en el protocolo: pocas frases por intención, una sola revisora de etiquetas, y n pequeño en el pre-piloto.

# %% [markdown]
# ## Etapa 9 — Respuestas y TUPA (G4)
# **Qué se hace:** se cuentan las respuestas del dominio congelado (`domain_v3.yml`), cuántas siguen marcadas `[Verificar]`, y se lee el **Resumen** de la hoja de verificación contra el TUPA (`Verificacion_TUPA_v10_4.xlsx`): 44 respuestas de trámites, prioridad Alta pendiente, filas con alerta y confirmaciones del tesista.
# **Por qué:** las respuestas forman parte del modelo (cambiar un texto cambia el dominio y, por tanto, el congelamiento). Por eso **G4 va antes de G5**. El criterio de G4 es: Alta pendientes = 0, alertas = 0 y Corregir/Coincide sin confirmar = 0.
# **Paso del protocolo:** P09 y compuerta **G4**.

# %%
dom = yaml.safe_load((BASE / "domain_v3.yml").read_text(encoding="utf-8"))
resp = dom["responses"]
con_marca = [k for k, v in resp.items() if "[Verificar" in json.dumps(v, ensure_ascii=False)]
sin_texto = [k for k, v in resp.items() if not v or not str(v[0].get("text", "")).strip()]
intents = [i if isinstance(i, str) else list(i)[0] for i in dom["intents"]]
faltan = [i for i in intents if f"utter_{i}" not in resp and i != "nlu_fallback"]
rotulo(ORIGEN_SINTETICO, "texto de las respuestas del chatbot (código/dominio, no datos de personas)")
print("Respuestas en domain_v3.yml:", len(resp), "| intenciones:", len(intents), "| intenciones sin utter_<intención>:", faltan or "ninguna", "| respuestas vacías:", sin_texto or "ninguna")
print("Respuestas con marcador [Verificar]:", len(con_marca), con_marca)

from openpyxl import load_workbook
wbt = load_workbook(BASE / "docs/tupa/Verificacion_TUPA_v10_4.xlsx", read_only=True, data_only=True)
filas_t = list(wbt["Verificacion"].iter_rows(values_only=True))
resumen_t = {r[0]: (r[1] if len(r) > 1 else None) for r in wbt["Resumen"].iter_rows(values_only=True) if r and isinstance(r[0], str)}
wbt.close()
enc = next(i for i, r in enumerate(filas_t[:6]) if r and r[0] == "#")
col = {str(c): j for j, c in enumerate(filas_t[enc]) if c}
datos_t = [r for r in filas_t[enc + 1:] if r and r[col["Intención"]]]
resultados = pd.Series([str(r[col["Resultado"]]) for r in datos_t]).value_counts().to_dict()
intenciones_tupa = [r[col["Intención"]] for r in datos_t]
con_marca_tupa = [i for i in intenciones_tupa if f"utter_{i}" in con_marca]
rotulo(ORIGEN_REAL, "hoja de verificación del TUPA (valores guardados por Excel); el contenido del TUPA es público, no son datos de personas")
display(pd.DataFrame({"indicador (hoja Resumen)": [k for k in ("Respuestas a verificar", "Pendiente", "Coincide", "Corregir", "No figura en la fuente", "Con resultado (propuesto o confirmado)", "Confirmadas por el tesista", "Prioridad Alta pendientes", "Prioridad Media pendientes", "Filas con alerta", "Corregir o Coincide sin confirmar por el tesista")],
                      "valor": [resumen_t.get(k) for k in ("Respuestas a verificar", "Pendiente", "Coincide", "Corregir", "No figura en la fuente", "Con resultado (propuesto o confirmado)", "Confirmadas por el tesista", "Prioridad Alta pendientes", "Prioridad Media pendientes", "Filas con alerta", "Corregir o Coincide sin confirmar por el tesista")]}))
comparar("9", "respuestas de trámites a verificar", len(intenciones_tupa), 44, "Nota v11 / hoja Resumen", tol=0)
comparar("9", "respuestas con [Verificar] entre las 44", len(con_marca_tupa), 0, "Nota v11 («0 de 44»)", tol=0)
comparar("9", "respuestas con [Verificar] en todo el dominio", len(con_marca), 0, "Nota v11", tol=0)
comparar("9", "Prioridad Alta pendientes", resumen_t.get("Prioridad Alta pendientes"), 0, "criterio G4", tol=0)
comparar("9", "filas con alerta", resumen_t.get("Filas con alerta"), 0, "criterio G4", tol=0)
comparar("9", "Corregir o Coincide sin confirmar", resumen_t.get("Corregir o Coincide sin confirmar por el tesista"), 0, "criterio G4", tol=0)
comparar("9", "resultado «Corregir» contado en la hoja", resultados.get("Corregir", 0), resumen_t.get("Corregir"), "hoja Resumen", tol=0)
if resumen_t.get("Marcadas [Verificar] en el repositorio") not in (None, 0):
    DECLARADAS.append(f"Verificacion_TUPA_v10_4.xlsx, hoja Resumen: «Marcadas [Verificar] en el repositorio» = {resumen_t['Marcadas [Verificar] en el repositorio']} es un campo manual que NO se actualizó tras aplicar las respuestas; el dominio actual tiene {len(con_marca)} marcadores (comprobado arriba). No se modificó el libro.")
    print("Observación (no se corrige):", DECLARADAS[-1])
mostrar_comparaciones(etapas=["9"], titulo="Etapa 9 — recalculado frente a reportado")

# %% [markdown]
# ### Cómo leer el resultado (etapa 9)
# - **0 respuestas con `[Verificar]`** en el dominio congelado: cada texto con costo, plazo o requisito ya pasó por la verificación contra el TUPA (o por el texto confirmado por el tesista). Hay 55 respuestas porque el dominio incluye las 54 intenciones más `utter_no_entendi`.
# - De las **44** respuestas de trámites, **35 tienen resultado** (propuesto o confirmado) y **28 están confirmadas por el tesista**; solo las confirmadas se aplican al dominio. Quedan 9 pendientes de prioridad Media (canales, horarios, teléfonos); no bloquean G4 porque el criterio exige Alta pendientes = 0.
# - Si la hoja Resumen conserva un campo manual desactualizado (ver «Observación»), se reporta pero no se edita el libro: es evidencia.
