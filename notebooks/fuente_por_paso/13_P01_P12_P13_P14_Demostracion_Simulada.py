# %% [markdown]
# # P01, P12–P14 — Demostración con datos SIMULADOS
# ## (a) Paso del protocolo y objetivo
# **Todo lo de este cuaderno es «Simulado (demostración)»: no son hallazgos de campo y nada se suma a las cifras reales.**
# - **P01** (diagnóstico, **OE1**): línea base de tiempos **simulada** (120 respuestas inventadas del diseño V1.2).
# - **P12–P13** (piloto exploratorio de 60 sesiones asistidas; **OE3 y OE4**): hoy **planificado**; sin sesiones reales. La demostración usa un registro de 60 sesiones **inventadas**.
# - **P14** (análisis estadístico): demuestra que el *pipeline* corre (normalidad de Shapiro-Wilk → t pareada o Wilcoxon, tamaño del efecto, alfa de Cronbach).
# - La **demostración completa** (`demostracion_simulada.py`, 5 de octubre) encadena ingesta → partición → evaluación → congelamiento → sesiones → tablero.
#
# ## (b) Código del repositorio que lo implementa
# `demostracion_simulada.py`, `deteccion_simulado.py` (impide que lo simulado cuente como real), `simular_P14_n120.py`, `simular_analisis_p14.py`, `stats_analysis.py` y `analizar_piloto.py` (**modo `--demo-simulada` solo**). Entradas: `Lote1_Transcripcion_SIMULADO_v2.xlsx` y `Registro_Sesiones_Piloto_SIMULADO_v4.xlsx`. Salidas: `evidencias/simulado_demostracion/<ejecución>/`.

# %%
# %% prep

# %%
tabla_scripts(scripts_de("13"))

# %% [markdown]
# ## (c) Celdas de código
# ### 1. P01 — Línea base de tiempos SIMULADA

# %%
from openpyxl import load_workbook
wbq = load_workbook(BASE / "corpus/Encuestas_simuladas_TramiFacil_MPSR_120_v2.xlsx", read_only=True, data_only=True)
filas_q = list(wbq.worksheets[0].iter_rows(values_only=True))
wbq.close()
cab = [str(c) for c in filas_q[0]]
q = pd.DataFrame(filas_q[1:], columns=cab)
col_t = next(c for c in cab if c.startswith("Tiempo presencial"))
rotulo(ORIGEN_SIMULADO, "línea base inventada del diseño V1.2 (n = 120): NO es recolección de campo")
print("Respuestas simuladas:", len(q))
display(q[col_t].value_counts().rename("respuestas simuladas").to_frame())

# %% [markdown]
# ### 2. P14 — Análisis estadístico sobre datos simulados (resultado guardado)

# %%
for archivo in ("SIMULACION_resultado_P14_n120.json",):
    sim14 = leer_json("logs/simulaciones_P14/" + archivo)
    rotulo(ORIGEN_SIMULADO, "resultado guardado de simular_P14_n120.py — " + sim14["ADVERTENCIA"][:150] + "…")
    t = sim14["tiempo_atencion"]
    display(pd.DataFrame([{"variable": t["variable"], "n": t["n"], "media pre": t["media_pre"], "media post (simulada)": t["media_post"], "prueba": t["prueba_usada"], "p": t["p_value"], "decisión H0": t["decision_H0"]}]))
    print("Ninguna de estas cifras debe reportarse como hallazgo real de la tesis.")

# %% [markdown]
# ### 3. Demostración completa de punta a punta
# La celda siguiente necesita el tablero simulado.

# %%
ts = leer_json("evidencias/simulado_demostracion/20261005_demostracion/06_tablero/estado_compuertas_SIMULADO.json")

# %% reuse: e15_md, e15_c1, e15_leer

# %%
cierre("P01 y P12–P14 (Simulado)")

# %% [markdown]
# ## (d) Cómo leer el resultado
# - La línea base simulada es una tabla de rangos de tiempo inventados; la reducción de tiempo de P14 es **ilustrativa** (ADVERTENCIA del propio archivo). Sirve para probar el *pipeline*, no para afirmar nada sobre la MPSR.
# - En la demostración completa, G7 aparece «Cumplida (Simulado)» porque se inventaron 60 sesiones; **no cuenta** para el tablero real (5 de 7, con G7 planificada). El alfa de la demostración (≈ 0,55) no tiene relación con el 0,366 del pre-piloto real.
#
# ## Limitaciones
# - Las frases del lote simulado las escribió quien armó el libro de prueba (varias son idénticas al corpus sintético) y las «consultas» de las sesiones son el texto de las tarjetas; la prueba final usa un **predictor simulado**.
# - Los resultados de este cuaderno nunca deben citarse como evidencia empírica de la tesis: solo como prueba de que el procedimiento se puede ejecutar.
